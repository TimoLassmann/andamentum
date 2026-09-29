"""Tests for the slide checks and the ``--slides`` CLI mode. All slide text is invented."""

from __future__ import annotations

import json
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.util import Inches, Pt

from andamentum.proofread import Slide, analyze_slides
from andamentum.proofread.cli import main
from andamentum.proofread.slides import content_stems, stem


def _checks(slides: list[Slide], check: str) -> list[int]:
    return [f.slide for f in analyze_slides(slides).findings if f.check == check]


def test_stem_matches_word_forms() -> None:
    assert stem("summarises") == stem("summariser")
    assert stem("reducing") == stem("Reduce")
    assert stem("children’s") == stem("child")


def test_content_stems_keep_sentence_initial_word_and_drop_entities() -> None:
    stems = content_stems("Reducing pain in Apollo cohorts")
    assert stem("reducing") in stems
    assert "apollo" not in stems


def test_restatement_flags_body_that_repeats_headline() -> None:
    slides = [
        Slide(
            number=1,
            headline="SaltLog summarises the evidence for each variant",
            body=["one agent", "summariser", "evidence beneath"],
        ),
        Slide(
            number=2,
            headline="SaltLog halves the time to a report",
            body=["hours per report", "before: 40", "after: 20"],
        ),
    ]
    assert _checks(slides, "restatement") == [1]


def test_restatement_ignores_entity_labels() -> None:
    slides = [
        Slide(
            number=1,
            headline="Apollo Kinase blockers reach three syndromes",
            body=["Apollo Kinase", "syndrome A", "syndrome B"],
        )
    ]
    assert _checks(slides, "restatement") == []


def test_headline_length_and_deixis() -> None:
    slides = [
        Slide(number=1, headline="This halves the current"),
        Slide(number=2, headline="x" * 70),
        Slide(number=3, headline="x" * 95),
        Slide(number=4, headline="SaltLog halves the current"),
    ]
    assert _checks(slides, "deixis") == [1]
    assert _checks(slides, "headline_length") == [2, 3]


def test_acronym_needs_expansion_on_the_same_slide() -> None:
    slides = [
        Slide(
            number=1,
            headline="The Saltmarsh Genome Trust sends 40 questions a month (SGT)",
        ),
        Slide(number=2, headline="SGT questions repeat within a week"),
        Slide(number=3, headline="DNA from 12 donors"),
    ]
    assert _checks(slides, "acronym") == [2]
    allowed = analyze_slides(slides, allowed_acronyms=frozenset({"DNA", "SGT"}))
    assert not [f for f in allowed.findings if f.check == "acronym"]


def test_two_names_for_one_entity() -> None:
    slides = [
        Slide(number=1, headline="The Apollo Rare Centre sees 300 families a year"),
        Slide(
            number=2,
            headline="ARCC runs the family clinic",
            body=["Apollo Rare Clinic Centre (ARCC)"],
        ),
    ]
    found = [f for f in analyze_slides(slides).findings if f.check == "two_names"]
    assert found and "ARCC" in found[0].detail


def test_banned_word() -> None:
    slides = [Slide(number=1, headline="SaltLog lets us leverage every run")]
    assert _checks(slides, "banned_word") == [1]


def _deck(path: Path) -> None:
    prs = Presentation()
    blank = prs.slide_layouts[6]
    s = prs.slides.add_slide(blank)
    for text, top, size in [
        ("This deck shows SGT results", 0.5, 28),
        ("deck results for SGT", 3.0, 18),
        ("Source: invented, 2026", 6.3, 14),
        ("2", 6.8, 14),
    ]:
        tb = s.shapes.add_textbox(Inches(1), Inches(top), Inches(8), Inches(0.5))
        run = tb.text_frame.paragraphs[0].add_run()
        run.text = text
        run.font.size = Pt(size)
    notes = s.notes_slide.notes_text_frame
    assert notes is not None
    notes.text = "HOLD: the invented cohort"
    prs.save(str(path))


def test_cli_slides_json(tmp_path: Path, capsys: pytest.CaptureFixture[str]) -> None:
    deck = tmp_path / "deck.pptx"
    _deck(deck)
    assert main([str(deck), "--slides", "--format", "json"]) == 0
    payload = json.loads(capsys.readouterr().out)
    checks = {f["check"] for f in payload["findings"]}
    assert {"deixis", "acronym", "restatement"} <= checks
    assert payload["slide_count"] == 1


def test_cli_slides_allow(tmp_path: Path, capsys: pytest.CaptureFixture[str]) -> None:
    deck = tmp_path / "deck.pptx"
    _deck(deck)
    assert main([str(deck), "--slides", "--allow", "SGT"]) == 0
    assert "acronym" not in capsys.readouterr().out


def test_cli_slides_rejects_non_pptx(tmp_path: Path) -> None:
    f = tmp_path / "notes.txt"
    f.write_text("hello")
    with pytest.raises(SystemExit) as exc:
        main([str(f), "--slides"])
    assert exc.value.code == 1


def test_placeholders_are_not_acronyms() -> None:
    slides = [
        Slide(
            number=1,
            headline="SaltLog halves the time to a report",
            body=["[SOURCE NEEDED: SGT count for 2026]"],
        )
    ]
    assert _checks(slides, "acronym") == []


def test_statement_slide_takes_three_lines() -> None:
    text = "SaltLog now answers every question from the Saltmarsh group within a day"
    assert _checks([Slide(number=1, headline=text)], "headline_length") == [1]
    assert (
        _checks([Slide(number=1, headline=text, headline_pt=44)], "headline_length")
        == []
    )
    assert _checks(
        [Slide(number=1, headline="x" * 105, headline_pt=44)], "headline_length"
    ) == [1]
