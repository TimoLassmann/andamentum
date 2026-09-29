"""Read a ``.pptx`` file into ``Slide`` objects for ``analyze_slides``.

Used only by the CLI. The library (``slides.py``) stays free of I/O.

Which text frame is the headline: the top-most frame set at 24 pt or larger. A
statement slide has no separate headline, so its one large sentence is taken as
the headline, which is what the checks should see. When no frame carries a size,
the top-most frame is the headline. Frames at 14 pt or smaller are the source
line and the slide number; a frame holding only digits is dropped.
"""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

from pptx import Presentation
from pptx.util import Emu

from .slides import Slide

_HEADLINE_MIN_PT = 24
_SOURCE_MAX_PT = 14


@dataclass
class _Frame:
    top: int
    left: int
    text: str
    max_pt: float | None


def _frames(shapes) -> list[_Frame]:  # noqa: ANN001 — python-pptx shape collections are untyped
    out: list[_Frame] = []
    for sh in shapes:
        if sh.shape_type == 6:  # MSO_SHAPE_TYPE.GROUP
            out.extend(_frames(sh.shapes))
            continue
        top = int(sh.top or 0)
        left = int(sh.left or 0)
        if getattr(sh, "has_table", False) and sh.has_table:
            cells = [
                c.text for row in sh.table.rows for c in row.cells if c.text.strip()
            ]
            if cells:
                out.append(_Frame(top, left, " ".join(cells), None))
            continue
        if not getattr(sh, "has_text_frame", False) or not sh.has_text_frame:
            continue
        text = "\n".join(p.text for p in sh.text_frame.paragraphs).strip()
        if not text:
            continue
        sizes = [
            r.font.size
            for p in sh.text_frame.paragraphs
            for r in p.runs
            if r.font.size is not None
        ]
        max_pt = max(Emu(s).pt for s in sizes) if sizes else None
        out.append(_Frame(top, left, text, max_pt))
    return out


def read_pptx(path: str | Path) -> list[Slide]:
    """Return one ``Slide`` per slide in ``path``, in deck order."""
    prs = Presentation(str(path))
    slides: list[Slide] = []
    for number, sl in enumerate(prs.slides, start=1):
        frames = sorted(_frames(sl.shapes), key=lambda f: (f.top, f.left))
        frames = [f for f in frames if not f.text.strip().isdigit()]

        source = [
            f for f in frames if f.max_pt is not None and f.max_pt <= _SOURCE_MAX_PT
        ]
        rest = [f for f in frames if f not in source]
        big = [f for f in rest if f.max_pt is not None and f.max_pt >= _HEADLINE_MIN_PT]
        head = big[0] if big else (rest[0] if rest else None)

        notes = ""
        if sl.has_notes_slide and sl.notes_slide.notes_text_frame is not None:
            notes = sl.notes_slide.notes_text_frame.text

        slides.append(
            Slide(
                number=number,
                headline=head.text if head else "",
                headline_pt=head.max_pt if head else None,
                body=[f.text for f in rest if f is not head],
                source=" ".join(f.text for f in source),
                notes=notes,
            )
        )
    return slides
