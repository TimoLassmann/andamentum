"""Drafted sections into a Word document: the transactions of one build.

A build starts from the base document (the user's file for ``--into``, a clean
template for a new document) and never edits it. It writes three copies, each
through the one safe writer (``DocxAdapter.apply``):

1. ``1-text.docx``: every heading and paragraph as a tracked inserted
   paragraph (``InsertParagraph``). Paragraph ids of new text are not known
   before this step, so everything that needs one comes after it.
2. ``2-marks.docx``: bound numbers (``Bind``) and the author's questions as
   comments on ``[TO FILL: ...]`` markers and interpretation sentences.
3. ``3-cites.docx``: citations (``InsertCitation``) and, in the same
   transaction because the writer lists the works of that transaction, the
   reference list (``UpdateBibliography``).

Text is inserted clean (markers cut out, see ``body``) so that the offsets of
steps 2 and 3 are exact; the citation text itself is the writer's business.

A build is a pure function of the drafted bodies and the base document, so the
repair loop rebuilds from the base each round instead of editing its own
output: what was revised is simply drafted again.
"""

from __future__ import annotations

import re
from collections import Counter
from collections.abc import Mapping, Sequence
from dataclasses import dataclass, field
from pathlib import Path
from typing import TYPE_CHECKING

import docx
from docx.enum.style import WD_STYLE_TYPE

from andamentum.compose.schemas import Fact, SectionPlan

from ..bindings.ledger import NumberBinding
from ..docx.writer import DocxAdapter
from ..kernel import (
    Bind,
    CommentOp,
    DocumentView,
    InsertParagraph,
    Transaction,
)
from ..kernel.csl import work_key
from ..kernel.ops import CitationItem, InsertCitation, Operation, UpdateBibliography
from .body import BodyParagraph
from .claims import SectionAnalysis, SourceInfo

if TYPE_CHECKING:
    from .cite import Rewrite

AUTHOR = "andamentum compose"
HEADING = "heading"
BODY = "body"
_INTERPRETATION_NOTE = (
    "Interpretation: no source carries this judgement. Confirm it, rewrite it "
    "or remove it."
)
TYPED_CITATION_NOTE = (
    "[error] Typed citation, not verified: the harness did not insert this "
    "citation and could not check the work it names. Remove it, or cite the "
    "work with find-citation and insert-citation."
)


class DraftError(RuntimeError):
    """A build step did not do what was asked of it (never a quiet degradation)."""


@dataclass(frozen=True)
class Placement:
    """Where a section goes: after this paragraph of the base document."""

    anchor: str
    heading: (
        bool  # insert the section heading first (False: append under an existing one)
    )


@dataclass
class SectionDraft:
    key: str
    title: str
    body: str
    plan: SectionPlan
    placement: Placement
    paragraphs: list[BodyParagraph] = field(default_factory=list)
    analysis: SectionAnalysis = field(default_factory=SectionAnalysis)


@dataclass(frozen=True)
class Styles:
    """Paragraph style names to use; ``None`` means the document default."""

    heading: str
    body: str | None
    bullet: str | None
    number: str | None


@dataclass(frozen=True)
class PlacedParagraph:
    """A paragraph as inserted: the drafted text with a list marker in front and a
    space opened before each citation, and how drafted offsets map into it."""

    section: str
    role: str  # HEADING or BODY
    index: int  # into the section's body paragraphs (BODY only)
    text: str
    style: str | None
    shift: int = 0  # characters put in front of the drafted text (a list marker)
    gaps: tuple[int, ...] = ()  # drafted offsets where a space was opened

    def start(self, offset: int) -> int:
        """Where a drafted offset that begins something (a range start, a citation) lands."""
        return offset + self.shift + sum(1 for g in self.gaps if g <= offset)

    def end(self, offset: int) -> int:
        """Where a drafted offset that ends a range lands (a space just after it stays outside)."""
        return offset + self.shift + sum(1 for g in self.gaps if g < offset)


@dataclass
class BuiltRound:
    docx: Path
    view: DocumentView
    owner: dict[str, str]  # paragraph id of ``view`` -> section key
    bindings: list[NumberBinding]
    citations: int
    works: list[str]  # keys of the works cited, in order of first use
    notes: list[str]
    refused: list[str]


# ----------------------------------------------------------------- styles


def paragraph_style_names(path: Path) -> set[str]:
    document = docx.Document(str(path))
    return {
        s.name.lower()
        for s in document.styles
        if s.type == WD_STYLE_TYPE.PARAGRAPH and s.name
    }


def resolve_styles(path: Path, view: DocumentView, *, heading_level: int) -> Styles:
    """The styles a build uses in *path*.

    Headings use ``Heading N`` when the document has it, else the style of an
    existing heading of that level. Body text copies the commonest style of the
    document's own body paragraphs (so a house style survives); bullets and
    numbered items use the ``List Bullet`` / ``List Number`` styles when present
    and otherwise carry their marker in the text.
    """
    names = paragraph_style_names(path)
    heading = f"heading {heading_level}"
    if heading not in names:
        existing = next(
            (p.style for p in view.paragraphs if p.heading_level == heading_level),
            None,
        )
        if existing is None:
            raise DraftError(
                f"the document has no 'Heading {heading_level}' style and no heading "
                "of that level to copy; add one to the template"
            )
        heading = existing
    styles = Counter(
        p.style
        for p in view.paragraphs
        if p.heading_level is None and len(p.text) >= 40 and p.style.lower() != "title"
    )
    body = styles.most_common(1)[0][0] if styles else None
    if body is not None and body.lower() == "normal":
        body = None
    return Styles(
        heading=heading,
        body=body,
        bullet="List Bullet" if "list bullet" in names else None,
        number="List Number" if "list number" in names else None,
    )


# --------------------------------------------------------------- placement


def _normal(text: str) -> str:
    stripped = re.sub(r"^\s*\d+(?:\.\d+)*[.)]?\s+", "", text)
    return " ".join(stripped.casefold().split())


def normalised_title(text: str) -> str:
    return _normal(text)


def section_spans(view: DocumentView) -> dict[str, tuple[int, int]]:
    """Normalised heading text -> (heading index, index of the section's last paragraph)."""
    headings = [p for p in view.paragraphs if p.heading_level is not None]
    spans: dict[str, tuple[int, int]] = {}
    for position, heading in enumerate(headings):
        nxt = (
            headings[position + 1].index
            if position + 1 < len(headings)
            else len(view.paragraphs)
        )
        spans.setdefault(_normal(heading.text), (heading.index, nxt - 1))
    return spans


def place_sections(
    titles: Sequence[str],
    targets: set[str],
    view: DocumentView,
) -> dict[str, Placement]:
    """A placement for each target title, following the outline order.

    A target that exists in the document is appended under its heading. A new
    one goes after the last paragraph of the nearest earlier outline section
    that is in the document (or just before the nearest later one), chained
    behind sections placed at the same spot; with neither, at the document's end.
    """
    spans = section_spans(view)
    last = view.paragraphs[-1].id
    placements: dict[str, Placement] = {}
    previous: str | None = None
    for position, title in enumerate(titles):
        norm = _normal(title)
        if norm in spans:
            previous = view.paragraphs[spans[norm][1]].id
            if title in targets:
                placements[title] = Placement(anchor=previous, heading=False)
            continue
        if title not in targets:
            continue
        anchor = previous
        if anchor is None:
            later = next(
                (
                    spans[_normal(t)][0]
                    for t in titles[position + 1 :]
                    if _normal(t) in spans
                ),
                None,
            )
            anchor = view.paragraphs[later - 1].id if later else last
        placements[title] = Placement(anchor=anchor, heading=True)
    return placements


# -------------------------------------------------------------- the build


def cite_items(
    fact_ids: Sequence[str],
    *,
    facts: Mapping[str, Fact],
    infos: Mapping[str, SourceInfo],
) -> list[CitationItem]:
    """The works behind a marker group, each once; empty for facts of the user's own sources."""
    items: list[CitationItem] = []
    for fact_id in fact_ids:
        fact = facts.get(fact_id)
        info = infos.get(fact.source_id) if fact else None
        if info is None or info.work is None:
            continue
        if any(work_key(i.csl) == info.work.key for i in items):
            continue
        items.append(CitationItem(csl=info.work.csl, uri=info.work.uri))
    return items


def _gaps(
    paragraph: BodyParagraph,
    *,
    facts: Mapping[str, Fact],
    infos: Mapping[str, SourceInfo],
) -> tuple[int, ...]:
    """Offsets where a space must be opened so a citation does not touch the word before it.

    The writer inserts a citation verbatim, so ``fidelity.`` would read
    ``fidelity(Okafor 2020).``; the space is part of the inserted paragraph.
    """
    found = []
    for cite in paragraph.cites:
        before = paragraph.text[cite.at - 1] if cite.at > 0 else " "
        if cite_items(cite.fact_ids, facts=facts, infos=infos) and not (
            before.isspace() or before in "(["
        ):
            found.append(cite.at)
    return tuple(found)


def placed_paragraphs(
    sections: Sequence[SectionDraft],
    *,
    styles: Styles,
    facts: Mapping[str, Fact],
    infos: Mapping[str, SourceInfo],
) -> list[PlacedParagraph]:
    placed: list[PlacedParagraph] = []
    for section in sections:
        if section.placement.heading:
            placed.append(
                PlacedParagraph(
                    section=section.key,
                    role=HEADING,
                    index=0,
                    text=section.title,
                    style=styles.heading,
                    shift=0,
                )
            )
        for index, paragraph in enumerate(section.paragraphs):
            style = styles.body
            if paragraph.list_kind == "bullet" and styles.bullet:
                style = styles.bullet
            elif paragraph.list_kind == "number" and styles.number:
                style = styles.number
            prefix = ""
            if paragraph.list_kind and style not in (styles.bullet, styles.number):
                prefix = f"{paragraph.list_marker} "
            gaps = _gaps(paragraph, facts=facts, infos=infos)
            text = paragraph.text
            for gap in sorted(gaps, reverse=True):
                text = text[:gap] + " " + text[gap:]
            placed.append(
                PlacedParagraph(
                    section=section.key,
                    role=BODY,
                    index=index,
                    text=prefix + text,
                    style=style,
                    shift=len(prefix),
                    gaps=gaps,
                )
            )
    return placed


def carry_ids(before: DocumentView, after: DocumentView) -> dict[str, str]:
    """Paragraph id in *before* -> id in *after*, when *after* only grew.

    Paragraphs keep their order; an existing paragraph's text may have gained
    inserted text (a citation), never lost any. Each is matched, in order, to
    the first later paragraph of *after* that contains its text as a
    subsequence of characters.
    """

    def contains(whole: str, part: str) -> bool:
        chars = iter(whole)
        return all(ch in chars for ch in part)

    mapping: dict[str, str] = {}
    cursor = 0
    for paragraph in before.paragraphs:
        for j in range(cursor, len(after.paragraphs)):
            if contains(after.paragraphs[j].text, paragraph.text):
                mapping[paragraph.id] = after.paragraphs[j].id
                cursor = j + 1
                break
        else:
            raise DraftError(
                f"paragraph {paragraph.index} of the intermediate document is missing "
                "after the next step"
            )
    return mapping


def _inserted_ids(
    base: DocumentView,
    after_text: DocumentView,
    placed: Sequence[tuple[str, PlacedParagraph]],
) -> list[str]:
    """Ids (in *after_text*) of the inserted paragraphs, in *placed* order.

    The writer puts each new paragraph directly after its anchor, in operation
    order, so the position of every inserted paragraph follows from the counts.
    """
    index_of = {p.id: p.index for p in base.paragraphs}
    per_anchor = Counter(anchor for anchor, _ in placed)
    inserted_before = {
        paragraph.index: sum(
            count
            for anchor, count in per_anchor.items()
            if index_of[anchor] < paragraph.index
        )
        for paragraph in base.paragraphs
    }
    seen: Counter[str] = Counter()
    ids: list[str] = []
    for anchor, item in placed:
        seen[anchor] += 1
        position = index_of[anchor] + inserted_before[index_of[anchor]] + seen[anchor]
        paragraph = after_text.paragraphs[position]
        if paragraph.text != item.text:
            raise DraftError(
                f"inserted paragraph {position} reads {paragraph.text[:50]!r}, "
                f"expected {item.text[:50]!r}"
            )
        ids.append(paragraph.id)
    return ids


def _apply(
    adapter: DocxAdapter,
    view: DocumentView,
    ops: list[Operation],
    *,
    output: Path,
    reason: str,
) -> tuple[DocumentView, list[str]]:
    """Apply *ops* to a copy at *output*; no operations means no new file."""
    if not ops:
        return view, []
    result = adapter.apply(
        view,
        Transaction(ops=ops, author=AUTHOR, reason=reason, source="write"),
        output=output,
    )
    return adapter.open(result.output), result.refused


def build_round(
    *,
    base: Path,
    round_dir: Path,
    sections: Sequence[SectionDraft],
    styles: Styles,
    facts: Mapping[str, Fact],
    infos: Mapping[str, SourceInfo],
    document_notes: Sequence[str],
    reason: str,
    rewrites: Sequence[Rewrite] = (),
) -> BuiltRound:
    """Write the sections into copies of *base* (see the module docstring)."""
    adapter = DocxAdapter()
    round_dir.mkdir(parents=True, exist_ok=True)
    notes: list[str] = []

    # 1. the text
    base_view = adapter.open(base)
    placed = placed_paragraphs(sections, styles=styles, facts=facts, infos=infos)
    by_key = {s.key: s for s in sections}
    anchored = [(by_key[p.section].placement.anchor, p) for p in placed]
    view1, refused_text = _apply(
        adapter,
        base_view,
        [
            InsertParagraph(after_paragraph_id=anchor, text=p.text, style=p.style)
            for anchor, p in anchored
        ],
        output=round_dir / "1-text.docx",
        reason=reason,
    )
    if refused_text:
        raise DraftError("the writer refused drafted text: " + "; ".join(refused_text))
    new_ids = _inserted_ids(base_view, view1, anchored)
    where = {
        (p.section, p.role, p.index): (pid, p)
        for p, pid in zip(placed, new_ids, strict=True)
    }

    # 2. numbers, questions
    mark_ops: list[Operation] = []
    wanted: dict[str, NumberBinding] = {}
    for section in sections:
        for number, paragraph in enumerate(section.paragraphs):
            pid, at = where[(section.key, BODY, number)]
            for start, end in paragraph.to_fill():
                label = paragraph.text[start:end].removeprefix("[TO FILL:").rstrip("]")
                mark_ops.append(
                    CommentOp(
                        paragraph_id=pid,
                        start=at.start(start),
                        end=at.end(end),
                        text=f"Needs the author: {label.strip() or 'something is missing here'}.",
                    )
                )
        for typed in section.analysis.typed:
            pid, at = where[(section.key, BODY, typed.paragraph)]
            mark_ops.append(
                CommentOp(
                    paragraph_id=pid,
                    start=at.start(typed.start),
                    end=at.end(typed.end),
                    text=TYPED_CITATION_NOTE,
                )
            )
        for sentence in section.analysis.sentences:
            pid, at = where[(section.key, BODY, sentence.paragraph)]
            if sentence.kind == "interpretation":
                mark_ops.append(
                    CommentOp(
                        paragraph_id=pid,
                        start=at.start(sentence.start),
                        end=at.end(sentence.end),
                        text=_INTERPRETATION_NOTE,
                    )
                )
            for spot in sentence.numbers:
                wanted[spot.binding.anchor_id] = spot.binding
                mark_ops.append(
                    Bind(
                        paragraph_id=pid,
                        start=at.start(spot.start),
                        end=at.end(spot.end),
                        anchor_id=spot.binding.anchor_id,
                        kind="number",
                    )
                )
    for rewrite in rewrites:
        section = by_key.get(rewrite.section_key)
        if section is None:
            continue
        for number, paragraph in enumerate(section.paragraphs):
            found = paragraph.text.find(rewrite.new)
            if found < 0 or not rewrite.new:
                continue
            pid, at = where[(section.key, BODY, number)]
            mark_ops.append(
                CommentOp(
                    paragraph_id=pid,
                    start=at.start(found),
                    end=at.end(found + len(rewrite.new)),
                    text=f"Changed to match the cited paper. Drafted text was: {rewrite.old}",
                )
            )
            break
        else:
            notes.append(
                f"a sentence changed to match its cited paper could not be marked "
                f"with a comment (a revise round changed it): {rewrite.new}"
            )
    for note in document_notes:
        mark_ops.append(
            CommentOp(
                paragraph_id=new_ids[0], start=0, end=len(placed[0].text), text=note
            )
        )
    view2, refused = _apply(
        adapter, view1, mark_ops, output=round_dir / "2-marks.docx", reason=reason
    )
    bound = {a.id for a in view2.anchors}
    bindings = [b for anchor, b in wanted.items() if anchor in bound]
    notes += [
        f"number {b.written!r} was not bound (the writer refused it)"
        for anchor, b in wanted.items()
        if anchor not in bound
    ]
    index_of = {p.id: p.index for p in view1.paragraphs}
    id2 = {p.index: p.id for p in view2.paragraphs}

    # 3. citations and the reference list
    cite_ops: list[Operation] = []
    works: list[str] = []
    for section in sections:
        for number, paragraph in enumerate(section.paragraphs):
            pid, at = where[(section.key, BODY, number)]
            for cite in paragraph.cites:
                items = cite_items(cite.fact_ids, facts=facts, infos=infos)
                if not items:
                    continue
                works += [
                    work_key(i.csl) for i in items if work_key(i.csl) not in works
                ]
                cite_ops.append(
                    InsertCitation(
                        paragraph_id=id2[index_of[pid]],
                        at=at.start(cite.at),
                        items=items,
                    )
                )
    asked = len(cite_ops)
    if cite_ops:
        cite_ops.append(UpdateBibliography())
    view3, refused_cites = _apply(
        adapter, view2, cite_ops, output=round_dir / "3-cites.docx", reason=reason
    )
    carried = carry_ids(view2, view3)
    owner = {
        carried[id2[index_of[pid]]]: p.section
        for p, pid in zip(placed, new_ids, strict=True)
    }
    return BuiltRound(
        docx=Path(view3.path),
        view=view3,
        owner=owner,
        bindings=bindings,
        citations=asked - len([r for r in refused_cites if r.startswith("insert")]),
        works=works,
        notes=notes,
        refused=refused + refused_cites,
    )
