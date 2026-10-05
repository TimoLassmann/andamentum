"""The docx writer: apply a transaction, reopen with the reader, assert what Word would show."""

from __future__ import annotations

import zipfile
from pathlib import Path

import docx
import pytest
from lxml import etree  # type: ignore

from andamentum.compose.harness.docx.anchors import read_project_id
from andamentum.compose.harness.docx.ooxml import w, w15
from andamentum.compose.harness.docx.reader import open_docx
from andamentum.compose.harness.docx import writer as writer_module
from andamentum.compose.harness.docx.writer import DocxAdapter, DocxApplyError
from andamentum.compose.harness.kernel import (
    Bind,
    CommentOp,
    Delete,
    DocumentView,
    Insert,
    InsertParagraph,
    Move,
    Replace,
    Reply,
    Transaction,
    Unbind,
)
from andamentum.compose.harness.spike import make_spike_docx
from andamentum.compose.harness.spike.make import (
    build_docx,
    comment_part,
    paragraph,
    run,
)
from andamentum.compose.harness.spike.survival import SOFFICE, libreoffice_roundtrip

AUTHOR = "andamentum check"
METHODS = "p6-10000007"
ABSTRACT = "p2-10000003"
INTRO = "p4-10000005"


@pytest.fixture
def spike(tmp_path: Path) -> Path:
    return make_spike_docx(tmp_path / "spike.docx")


@pytest.fixture
def adapter() -> DocxAdapter:
    return DocxAdapter()


def _txn(*ops) -> Transaction:
    return Transaction(ops=list(ops), author=AUTHOR, reason="test", source="test")


def _document_root(path: Path):
    with zipfile.ZipFile(path) as archive:
        return etree.fromstring(archive.read("word/document.xml"))


def _text(view: DocumentView, paragraph_id: str) -> str:
    return next(p.text for p in view.paragraphs if p.id == paragraph_id)


def _span(view: DocumentView, paragraph_id: str, needle: str) -> tuple[int, int]:
    start = _text(view, paragraph_id).index(needle)
    return start, start + len(needle)


def test_replace_is_a_tracked_change_with_changes_accepted(
    spike: Path, adapter: DocxAdapter, tmp_path: Path
):
    before = spike.read_bytes()
    view = adapter.open(spike)
    start, end = _span(view, METHODS, "every meal")
    result = adapter.apply(
        view,
        _txn(
            Replace(
                paragraph_id=METHODS,
                start=start,
                end=end,
                text="each meal",
                expect="every meal",
            )
        ),
        output=tmp_path / "out.docx",
    )
    assert result.applied == 1 and result.refused == []
    assert result.untouched.ok
    assert spike.read_bytes() == before  # the input is never edited

    out = adapter.open(result.output)
    assert (
        _text(out, METHODS)
        == "Participants logged each meal using the SaltLog app. Intake is summarised in Table 1."
    )
    root = _document_root(result.output)
    mine = [e for e in root.iter(w("ins"), w("del")) if e.get(w("author")) == AUTHOR]
    assert {e.tag for e in mine} == {w("ins"), w("del")}
    # The co-author's earlier revisions are intact and the ids stay unique.
    ours = {e.get(w("id")) for e in mine}
    theirs = {
        e.get(w("id"))
        for e in root.iter(w("ins"), w("del"))
        if e.get(w("author")) == "Co-author"
    }
    assert ours.isdisjoint(theirs) and len(theirs) == 2


def test_untouched_parts_and_paragraphs_are_byte_identical(
    spike: Path, adapter: DocxAdapter, tmp_path: Path
):
    view = adapter.open(spike)
    start, end = _span(view, METHODS, "every")
    adapter.apply(
        view,
        _txn(
            Replace(
                paragraph_id=METHODS, start=start, end=end, text="each", expect="every"
            )
        ),
        output=tmp_path / "out.docx",
    )
    with zipfile.ZipFile(spike) as a, zipfile.ZipFile(tmp_path / "out.docx") as b:
        assert a.namelist() == b.namelist()
        for name in a.namelist():
            if name != "word/document.xml":
                assert a.read(name) == b.read(name), name
    before = _document_root(spike)
    after = _document_root(tmp_path / "out.docx")
    paragraphs_before = list(before.iter(w("p")))
    paragraphs_after = list(after.iter(w("p")))
    assert len(paragraphs_before) == len(paragraphs_after)
    changed = [
        i
        for i, (x, y) in enumerate(zip(paragraphs_before, paragraphs_after))
        if etree.tostring(x) != etree.tostring(y)
    ]
    assert changed == [6]


def test_expect_mismatch_refuses_only_that_operation(
    spike: Path, adapter: DocxAdapter, tmp_path: Path
):
    view = adapter.open(spike)
    start, end = _span(view, METHODS, "every meal")
    good = Replace(
        paragraph_id=METHODS,
        start=start,
        end=end,
        text="each meal",
        expect="every meal",
    )
    bad = Replace(
        paragraph_id=METHODS, start=0, end=12, text="X", expect="Participante"
    )
    result = adapter.apply(view, _txn(bad, good), output=tmp_path / "out.docx")
    assert result.applied == 1
    assert len(result.refused) == 1 and "Participante" in result.refused[0]
    assert "each meal" in _text(adapter.open(result.output), METHODS)
    assert _text(adapter.open(result.output), METHODS).startswith("Participants")


def test_insert_and_delete(spike: Path, adapter: DocxAdapter, tmp_path: Path):
    view = adapter.open(spike)
    at = _text(view, METHODS).index(". Intake")
    start, end = _span(view, INTRO, " and meal diaries are common")
    result = adapter.apply(
        view,
        _txn(
            Insert(paragraph_id=METHODS, at=at, text=" (n = 48)"),
            Delete(
                paragraph_id=INTRO,
                start=start,
                end=end,
                expect=" and meal diaries are common",
            ),
        ),
        output=tmp_path / "out.docx",
    )
    assert result.applied == 2
    out = adapter.open(result.output)
    assert _text(out, METHODS).endswith(
        "SaltLog app (n = 48). Intake is summarised in Table 1."
    )
    assert "meal diaries" not in _text(out, INTRO)
    assert _text(out, INTRO).startswith("Sodium intake exceeds guidelines")


def test_overlapping_edits_refuse_the_later_one(
    spike: Path, adapter: DocxAdapter, tmp_path: Path
):
    view = adapter.open(spike)
    start, _ = _span(view, METHODS, "every meal")
    first = Replace(
        paragraph_id=METHODS,
        start=start,
        end=start + 10,
        text="each meal",
        expect="every meal",
    )
    second = Delete(
        paragraph_id=METHODS, start=start + 3, end=start + 8, expect="ry me"
    )
    result = adapter.apply(view, _txn(first, second), output=tmp_path / "out.docx")
    assert result.applied == 1 and len(result.refused) == 1
    assert "overlaps" in result.refused[0]
    assert "each meal" in _text(adapter.open(result.output), METHODS)


def test_edit_inside_a_co_authors_insertion_nests(
    spike: Path, adapter: DocxAdapter, tmp_path: Path
):
    view = adapter.open(spike)
    start, end = _span(view, METHODS, "SaltLog")
    result = adapter.apply(
        view,
        _txn(
            Replace(
                paragraph_id=METHODS,
                start=start,
                end=end,
                text="DietLog",
                expect="SaltLog",
            )
        ),
        output=tmp_path / "out.docx",
    )
    assert "using the DietLog app" in _text(adapter.open(result.output), METHODS)
    root = _document_root(result.output)
    co_author = [e for e in root.iter(w("ins")) if e.get(w("author")) == "Co-author"]
    assert len(co_author) == 1  # still there, with our revision inside or beside it


def test_insert_paragraph_is_a_tracked_inserted_paragraph(
    spike: Path, adapter: DocxAdapter, tmp_path: Path
):
    view = adapter.open(spike)
    result = adapter.apply(
        view,
        _txn(
            InsertParagraph(
                after_paragraph_id=METHODS, text="First added.", style="heading 2"
            ),
            InsertParagraph(after_paragraph_id=METHODS, text="Second added."),
        ),
        output=tmp_path / "out.docx",
    )
    assert result.applied == 2
    out = adapter.open(result.output)
    texts = [p.text for p in out.paragraphs]
    at = texts.index("First added.")
    assert texts[at + 1] == "Second added."
    assert (
        out.paragraphs[at].style == "heading 2"
        and out.paragraphs[at].heading_level == 2
    )
    assert out.paragraphs[at + 1].style == "Normal"
    assert texts[at - 1].startswith("Participants logged")
    assert texts[at + 2].startswith("Table 1:")
    root = _document_root(result.output)
    new = [p for p in root.iter(w("p")) if "".join(p.itertext()) == "First added."]
    assert len(new) == 1
    mark = new[0].find(f"{w('pPr')}/{w('rPr')}/{w('ins')}")
    assert mark is not None and mark.get(w("author")) == AUTHOR
    assert new[0].find(w("ins")) is not None  # the content is a tracked insertion too


def test_insert_paragraph_unknown_style_is_refused(
    spike: Path, adapter: DocxAdapter, tmp_path: Path
):
    view = adapter.open(spike)
    result = adapter.apply(
        view,
        _txn(
            InsertParagraph(
                after_paragraph_id=METHODS, text="X.", style="No Such Style"
            )
        ),
        output=tmp_path / "out.docx",
    )
    assert result.applied == 0 and "No Such Style" in result.refused[0]
    assert [p.text for p in adapter.open(result.output).paragraphs] == [
        p.text for p in view.paragraphs
    ]


def test_comment_is_anchored_where_asked(
    spike: Path, adapter: DocxAdapter, tmp_path: Path
):
    view = adapter.open(spike)
    start, end = _span(view, INTRO, "Sodium intake exceeds guidelines")
    result = adapter.apply(
        view,
        _txn(
            CommentOp(
                paragraph_id=INTRO,
                start=start,
                end=end,
                text="Cite a source.\nSecond line.",
            )
        ),
        output=tmp_path / "out.docx",
    )
    out = adapter.open(result.output)
    new = [c for c in out.comments if c.author == AUTHOR]
    assert len(new) == 1
    comment = new[0]
    assert comment.text == "Cite a source.\nSecond line."
    assert comment.parent_id is None and not comment.resolved
    assert comment.paragraph_id == INTRO
    assert (
        _text(out, INTRO)[comment.start : comment.end]
        == "Sodium intake exceeds guidelines"
    )
    assert comment.id not in {c.id for c in view.comments}
    # The text itself is unchanged and no comments part was missing.
    assert _text(out, INTRO) == _text(view, INTRO)


def test_comment_range_inside_a_run_splits_it_without_changing_the_text(
    spike: Path, adapter: DocxAdapter, tmp_path: Path
):
    view = adapter.open(spike)
    start, end = _span(view, METHODS, "logged every")
    result = adapter.apply(
        view,
        _txn(
            CommentOp(paragraph_id=METHODS, start=start, end=end, text="Which meals?")
        ),
        output=tmp_path / "out.docx",
    )
    out = adapter.open(result.output)
    comment = next(c for c in out.comments if c.author == AUTHOR)
    assert _text(out, METHODS)[comment.start : comment.end] == "logged every"
    assert _text(out, METHODS) == _text(view, METHODS)


def test_comment_on_text_rewritten_in_the_same_transaction_covers_the_new_text(
    spike: Path, adapter: DocxAdapter, tmp_path: Path
):
    view = adapter.open(spike)
    start, end = _span(view, METHODS, "every meal")
    result = adapter.apply(
        view,
        _txn(
            Replace(
                paragraph_id=METHODS,
                start=start,
                end=end,
                text="each meal",
                expect="every meal",
            ),
            CommentOp(paragraph_id=METHODS, start=start, end=end, text="Reworded."),
        ),
        output=tmp_path / "out.docx",
    )
    assert result.applied == 2
    out = adapter.open(result.output)
    comment = next(c for c in out.comments if c.author == AUTHOR)
    assert _text(out, METHODS)[comment.start : comment.end] == "each meal"


def test_comment_on_a_deleted_range_is_refused(
    spike: Path, adapter: DocxAdapter, tmp_path: Path
):
    view = adapter.open(spike)
    start, end = _span(view, INTRO, " and meal diaries are common")
    result = adapter.apply(
        view,
        _txn(
            Delete(
                paragraph_id=INTRO,
                start=start,
                end=end,
                expect=" and meal diaries are common",
            ),
            CommentOp(paragraph_id=INTRO, start=start + 1, end=end, text="Why cut?"),
        ),
        output=tmp_path / "out.docx",
    )
    assert result.applied == 1
    assert "nothing is left" in result.refused[0]


def test_reply_joins_the_thread_of_the_root_comment(
    spike: Path, adapter: DocxAdapter, tmp_path: Path
):
    view = adapter.open(spike)
    # "1" is itself a reply; replying to it still threads to the root "0".
    result = adapter.apply(
        view,
        _txn(
            Reply(parent_comment_id="1", text="Reworded in the next draft."),
            Reply(parent_comment_id="0", text="One more point."),
        ),
        output=tmp_path / "out.docx",
    )
    assert result.applied == 2
    out = adapter.open(result.output)
    replies = [c for c in out.comments if c.author == AUTHOR]
    assert [c.text for c in replies] == [
        "Reworded in the next draft.",
        "One more point.",
    ]
    root = next(c for c in out.comments if c.id == "0")
    for reply in replies:
        assert reply.parent_id == "0"
        assert (reply.paragraph_id, reply.start, reply.end) == (
            root.paragraph_id,
            root.start,
            root.end,
        )
    # Replies keep their order in the document, after the existing ones.
    marks = [
        e.get(w("id"))
        for e in _document_root(result.output).iter(w("commentReference"))
    ]
    assert marks == ["0", "1", "2", "3"]


def test_reply_to_unknown_comment_is_refused(
    spike: Path, adapter: DocxAdapter, tmp_path: Path
):
    result = adapter.apply(
        adapter.open(spike),
        _txn(Reply(parent_comment_id="99", text="Hello?")),
        output=tmp_path / "out.docx",
    )
    assert result.applied == 0 and "no such comment" in result.refused[0]


def test_bind_wraps_the_range_in_a_hidden_tagged_control(
    spike: Path, adapter: DocxAdapter, tmp_path: Path
):
    view = adapter.open(spike)
    start, end = _span(view, INTRO, "exceeds guidelines")
    result = adapter.apply(
        view,
        _txn(
            Bind(
                paragraph_id=INTRO,
                start=start,
                end=end,
                anchor_id="claim-sodium",
                kind="claim",
            )
        ),
        output=tmp_path / "out.docx",
    )
    out = adapter.open(result.output)
    anchor = next(a for a in out.anchors if a.id == "claim-sodium")
    assert anchor.kind == "claim" and anchor.paragraph_id == INTRO
    assert _text(out, INTRO)[anchor.start : anchor.end] == "exceeds guidelines"
    assert _text(out, INTRO) == _text(view, INTRO)
    sdt = next(
        s
        for s in _document_root(result.output).iter(w("sdt"))
        if s.find(f"{w('sdtPr')}/{w('tag')}").get(w("val")) == "am:claim-sodium"
    )
    properties = sdt.find(w("sdtPr"))
    assert properties.find(w("alias")).get(w("val")) == "claim"
    assert properties.find(w15("appearance")).get(w15("val")) == "hidden"
    # The two existing anchors are untouched.
    assert {"spike-number", "spike-claim"} <= {a.id for a in out.anchors}


def test_bind_refusals(spike: Path, adapter: DocxAdapter, tmp_path: Path):
    view = adapter.open(spike)
    text = _text(view, INTRO)
    cuts_field = text.index("adults (Do")
    start, _ = _span(view, METHODS, "every meal")
    result = adapter.apply(
        view,
        _txn(
            Bind(
                paragraph_id=INTRO,
                start=cuts_field,
                end=cuts_field + 10,
                anchor_id="x1",
                kind="span",
            ),
            Bind(
                paragraph_id=INTRO, start=0, end=6, anchor_id="spike-claim", kind="span"
            ),
            Bind(
                paragraph_id=METHODS,
                start=start,
                end=start + 10,
                anchor_id="x2",
                kind="span",
            ),
            Replace(
                paragraph_id=METHODS,
                start=start,
                end=start + 10,
                text="each meal",
                expect="every meal",
            ),
        ),
        output=tmp_path / "out.docx",
    )
    assert result.applied == 1
    reasons = " | ".join(result.refused)
    assert "cuts a field" in reasons
    assert "already exists" in reasons
    assert "overlaps an edit" in reasons
    out = adapter.open(result.output)
    assert {a.id for a in out.anchors} == {"spike-number", "spike-claim"}


def test_comment_and_bind_on_the_same_range_keep_the_comment_outside_the_control(
    spike: Path, adapter: DocxAdapter, tmp_path: Path
):
    view = adapter.open(spike)
    start, end = _span(view, INTRO, "Sodium intake")
    result = adapter.apply(
        view,
        _txn(
            Bind(paragraph_id=INTRO, start=start, end=end, anchor_id="b", kind="span"),
            CommentOp(paragraph_id=INTRO, start=start, end=end, text="Check."),
        ),
        output=tmp_path / "out.docx",
    )
    assert result.applied == 2
    root = _document_root(result.output)
    sdt = next(
        s
        for s in root.iter(w("sdt"))
        if s.find(f"{w('sdtPr')}/{w('tag')}").get(w("val")) == "am:b"
    )
    assert list(sdt.iter(w("commentRangeStart"))) == []
    out = adapter.open(result.output)
    comment = next(c for c in out.comments if c.author == AUTHOR)
    assert _text(out, INTRO)[comment.start : comment.end] == "Sodium intake"


def test_comments_part_is_created_when_the_document_has_none(
    adapter: DocxAdapter, tmp_path: Path
):
    source = build_docx(
        tmp_path / "plain.docx",
        paragraph(run("Alpha beta gamma."), para_id="20000001")
        + paragraph(run("Delta epsilon."), para_id="20000002"),
    )
    view = adapter.open(source)
    result = adapter.apply(
        view,
        _txn(
            CommentOp(
                paragraph_id=view.paragraphs[0].id, start=6, end=10, text="beta?"
            ),
            Reply(parent_comment_id="1", text="not possible: it was just created"),
        ),
        output=tmp_path / "out.docx",
    )
    assert result.applied == 1  # the reply to a comment that does not exist is refused
    out = adapter.open(result.output)
    comment = out.comments[0]
    assert (
        comment.text == "beta?"
        and out.paragraphs[0].text[comment.start : comment.end] == "beta"
    )
    with zipfile.ZipFile(result.output) as archive:
        assert "word/comments.xml" in archive.namelist()
        assert b"comments+xml" in archive.read("[Content_Types].xml")
        assert b"relationships/comments" in archive.read("word/_rels/document.xml.rels")
    assert docx.Document(str(result.output)).paragraphs[0].text.startswith("Alpha")


def test_reply_creates_the_extended_part_when_missing(
    adapter: DocxAdapter, tmp_path: Path
):
    source = build_docx(
        tmp_path / "threadless.docx",
        paragraph(
            '<w:commentRangeStart w:id="5"/>'
            + run("Alpha beta.")
            + '<w:commentRangeEnd w:id="5"/><w:r><w:commentReference w:id="5"/></w:r>',
            para_id="20000001",
        ),
        comments_xml=comment_part(
            [(5, "Reviewer", "2026-10-01T09:00:00Z", "Why?", "30000001")]
        ),
    )
    view = adapter.open(source)
    result = adapter.apply(
        view,
        _txn(Reply(parent_comment_id="5", text="Because.")),
        output=tmp_path / "out.docx",
    )
    assert result.applied == 1
    out = adapter.open(result.output)
    reply = next(c for c in out.comments if c.author == AUTHOR)
    assert reply.parent_id == "5" and reply.id == "6"
    with zipfile.ZipFile(result.output) as archive:
        assert "word/commentsExtended.xml" in archive.namelist()


def test_a_comment_without_an_anchor_cannot_be_replied_to(
    adapter: DocxAdapter, tmp_path: Path
):
    source = build_docx(
        tmp_path / "orphan.docx",
        paragraph(run("Alpha."), para_id="20000001"),
        comments_xml=comment_part(
            [(5, "Reviewer", "2026-10-01T09:00:00Z", "Lost", "30000001")]
        ),
    )
    view = adapter.open(source)
    result = adapter.apply(
        view,
        _txn(Reply(parent_comment_id="5", text="Hi")),
        output=tmp_path / "out.docx",
    )
    assert result.applied == 0 and "no anchor" in result.refused[0]


def test_output_opens_in_python_docx_and_keeps_existing_structures(
    spike: Path, adapter: DocxAdapter, tmp_path: Path
):
    view = adapter.open(spike)
    start, end = _span(view, METHODS, "every meal")
    result = adapter.apply(
        view,
        _txn(
            Replace(
                paragraph_id=METHODS,
                start=start,
                end=end,
                text="each meal",
                expect="every meal",
            ),
            CommentOp(paragraph_id=INTRO, start=0, end=6, text="Sodium?"),
            Reply(parent_comment_id="0", text="Noted."),
            Bind(paragraph_id=ABSTRACT, start=0, end=2, anchor_id="in", kind="span"),
            InsertParagraph(after_paragraph_id=METHODS, text="Added."),
        ),
        output=tmp_path / "out.docx",
    )
    assert result.applied == 5
    docx.Document(str(result.output))
    out = adapter.open(result.output)
    assert out.project_id == view.project_id
    assert [f.kind for f in out.fields] == [f.kind for f in view.fields]
    assert {a.id for a in view.anchors} <= {a.id for a in out.anchors}


def test_move_and_unbind_are_not_implemented(
    spike: Path, adapter: DocxAdapter, tmp_path: Path
):
    view = adapter.open(spike)
    with pytest.raises(NotImplementedError):
        adapter.apply(
            view,
            _txn(Move(paragraph_id=METHODS, after_paragraph_id=INTRO)),
            output=tmp_path / "o.docx",
        )
    with pytest.raises(NotImplementedError):
        adapter.apply(
            view, _txn(Unbind(anchor_id="spike-claim")), output=tmp_path / "o.docx"
        )
    assert not (tmp_path / "o.docx").exists()


def test_apply_refuses_in_place_output_and_a_stale_view(
    spike: Path, adapter: DocxAdapter, tmp_path: Path
):
    view = adapter.open(spike)
    with pytest.raises(ValueError, match="never edits in place"):
        adapter.apply(view, _txn(), output=spike)
    changed = tmp_path / "changed.docx"
    start, end = _span(view, METHODS, "every")
    adapter.apply(
        view,
        _txn(
            Replace(
                paragraph_id=METHODS, start=start, end=end, text="each", expect="every"
            )
        ),
        output=changed,
    )
    stale = view.model_copy(update={"path": changed})
    with pytest.raises(ValueError, match="stale"):
        adapter.apply(stale, _txn(), output=tmp_path / "x.docx")


def test_a_range_outside_the_paragraph_or_an_unknown_paragraph_is_refused(
    spike: Path, adapter: DocxAdapter, tmp_path: Path
):
    view = adapter.open(spike)
    result = adapter.apply(
        view,
        _txn(
            CommentOp(paragraph_id=INTRO, start=0, end=10_000, text="Too far."),
            Insert(paragraph_id="p99", at=0, text="X"),
            Insert(paragraph_id=METHODS, at=-0, text="Participants "),
        ),
        output=tmp_path / "out.docx",
    )
    assert result.applied == 1
    assert len(result.refused) == 2


def test_verification_failure_raises_and_removes_the_output(
    spike: Path, adapter: DocxAdapter, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
):
    """A surgery that claims success but changes nothing must not pass silently."""
    real = writer_module.apply_edits

    def lying(textmap, edits, **kwargs):
        report = real(textmap, [], **kwargs)
        report.applied_edits = list(edits)
        return report

    monkeypatch.setattr(writer_module, "apply_edits", lying)
    view = adapter.open(spike)
    start, end = _span(view, METHODS, "every")
    output = tmp_path / "out.docx"
    with pytest.raises(DocxApplyError, match="do not read as intended"):
        adapter.apply(
            view,
            _txn(
                Replace(
                    paragraph_id=METHODS,
                    start=start,
                    end=end,
                    text="each",
                    expect="every",
                )
            ),
            output=output,
        )
    assert not output.exists()


def test_write_project_id_round_trips_and_refuses_another_id(
    adapter: DocxAdapter, tmp_path: Path
):
    source = build_docx(tmp_path / "plain.docx", paragraph(run("Alpha.")))
    tagged = adapter.write_project_id(
        source, project_id="proj-123", output=tmp_path / "tagged.docx"
    )
    assert adapter.open(tagged).project_id == "proj-123"
    with zipfile.ZipFile(tagged) as archive:
        assert read_project_id(archive) == "proj-123"
        assert "customXml/item1.xml" in archive.namelist()
    docx.Document(str(tagged))
    # Writing the same id again is a copy; a different id is refused.
    again = adapter.write_project_id(
        tagged, project_id="proj-123", output=tmp_path / "again.docx"
    )
    assert adapter.open(again).project_id == "proj-123"
    with pytest.raises(ValueError, match="already belongs"):
        adapter.write_project_id(
            tagged, project_id="other", output=tmp_path / "bad.docx"
        )
    with pytest.raises(ValueError):
        adapter.write_project_id(
            source, project_id="a<b", output=tmp_path / "bad2.docx"
        )


def test_apply_keeps_the_project_id_and_works_after_write_project_id(
    adapter: DocxAdapter, tmp_path: Path
):
    source = build_docx(tmp_path / "plain.docx", paragraph(run("Alpha beta.")))
    tagged = adapter.write_project_id(
        source, project_id="proj-9", output=tmp_path / "t.docx"
    )
    view = adapter.open(tagged)
    result = adapter.apply(
        view,
        _txn(
            CommentOp(paragraph_id=view.paragraphs[0].id, start=0, end=5, text="Alpha?")
        ),
        output=tmp_path / "out.docx",
    )
    assert adapter.open(result.output).project_id == "proj-9"


@pytest.mark.skipif(not SOFFICE, reason="LibreOffice not installed")
def test_libreoffice_opens_the_output_and_keeps_text_comments_and_anchors(
    spike: Path, adapter: DocxAdapter, tmp_path: Path
):
    view = adapter.open(spike)
    start, end = _span(view, METHODS, "every meal")
    result = adapter.apply(
        view,
        _txn(
            Replace(
                paragraph_id=METHODS,
                start=start,
                end=end,
                text="each meal",
                expect="every meal",
            ),
            CommentOp(paragraph_id=INTRO, start=0, end=13, text="Cite a source."),
            Reply(parent_comment_id="0", text="Noted."),
            Bind(
                paragraph_id=INTRO,
                start=0,
                end=13,
                anchor_id="claim-sodium",
                kind="claim",
            ),
            InsertParagraph(after_paragraph_id=METHODS, text="Added by the harness."),
        ),
        output=tmp_path / "out.docx",
    )
    assert result.applied == 5
    converted = libreoffice_roundtrip(result.output, workdir=tmp_path / "lo")
    after = open_docx(converted)
    texts = [p.text for p in after.paragraphs]
    assert any("each meal" in t for t in texts)
    assert "Added by the harness." in texts
    assert {"Cite a source.", "Noted."} <= {c.text for c in after.comments}
    assert "claim-sodium" in {a.id for a in after.anchors}


def _corrupt_other_paragraph(monkeypatch: pytest.MonkeyPatch, how: str) -> None:
    """Make the surgery also change a paragraph no operation targets."""
    real = writer_module._Writer.apply_edits

    def wrapper(self, origin):
        real(self, origin)
        for paragraph_id, element in zip(
            [p.id for p in self.view.paragraphs], self.elements, strict=True
        ):
            if paragraph_id == METHODS:
                continue
            node = next(element.iter(w("t")), None)
            if node is None:
                continue
            if how == "text":
                node.text = (node.text or "") + "!"
            else:
                element.set(w("rsidRDefault"), "0000BEEF")
            return
        raise AssertionError("no untargeted paragraph with text")

    monkeypatch.setattr(writer_module._Writer, "apply_edits", wrapper)


@pytest.mark.parametrize(
    ("how", "message"),
    [
        ("text", "paragraphs outside the operations changed"),
        ("bytes", "not byte-identical"),
    ],
)
def test_a_change_to_an_untargeted_paragraph_fails_verification(
    how: str,
    message: str,
    spike: Path,
    adapter: DocxAdapter,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
):
    _corrupt_other_paragraph(monkeypatch, how)
    view = adapter.open(spike)
    start, end = _span(view, METHODS, "every")
    output = tmp_path / "out.docx"
    with pytest.raises(DocxApplyError, match=message):
        adapter.apply(
            view,
            _txn(
                Replace(
                    paragraph_id=METHODS,
                    start=start,
                    end=end,
                    text="each",
                    expect="every",
                )
            ),
            output=output,
        )
    assert not output.exists()


# ------------------------------------------------- field results are not edited


def _field_doc(tmp_path: Path, content: str, name: str = "f.docx") -> Path:

    return build_docx(tmp_path / name, paragraph(content, para_id="10000001"))


def _in_field_result(root) -> list[str]:
    """Tags of revision, comment and control elements between a ``separate`` and
    its ``end`` fldChar (document order), which a refresh would destroy."""
    inside: list[str] = []
    depth = 0
    for element in root.iter():
        if element.tag == w("fldChar"):
            kind = element.get(w("fldCharType"))
            if kind == "separate":
                depth += 1
            elif kind == "end":
                depth = max(0, depth - 1)
        elif depth and element.tag in {
            w("ins"),
            w("del"),
            w("commentRangeStart"),
            w("commentRangeEnd"),
            w("commentReference"),
            w("sdt"),
        }:
            inside.append(element.tag.rpartition("}")[2])
    return inside


def _zotero_paragraph() -> str:
    from andamentum.compose.harness.spike.make import complex_field

    return (
        run("Sodium rose (")
        + complex_field("ADDIN ZOTERO_ITEM CSL_CITATION {}", "Doe 2020")
        + run(") overall")
    )


def _cited_at_end() -> str:
    from andamentum.compose.harness.spike.make import complex_field

    return run("Sodium rose ") + complex_field(
        "ADDIN ZOTERO_ITEM CSL_CITATION {}", "(Doe 2020)"
    )


def test_insert_after_a_citation_at_the_end_goes_outside_the_field(
    adapter: DocxAdapter, tmp_path: Path
):
    path = _field_doc(tmp_path, _cited_at_end())
    view = adapter.open(path)
    text = view.paragraphs[0].text
    result = adapter.apply(
        view,
        _txn(Insert(paragraph_id=view.paragraphs[0].id, at=len(text), text=" today")),
        output=tmp_path / "out.docx",
    )
    assert result.refused == []
    assert _in_field_result(_document_root(result.output)) == []
    assert adapter.open(result.output).paragraphs[0].text == text + " today"


def test_insert_at_the_start_of_a_field_at_paragraph_start_goes_outside(
    adapter: DocxAdapter, tmp_path: Path
):
    from andamentum.compose.harness.spike.make import complex_field

    path = _field_doc(
        tmp_path,
        complex_field("ADDIN ZOTERO_ITEM CSL_CITATION {}", "Doe 2020") + run(" says"),
    )
    view = adapter.open(path)
    result = adapter.apply(
        view,
        _txn(Insert(paragraph_id=view.paragraphs[0].id, at=0, text="See ")),
        output=tmp_path / "out.docx",
    )
    assert result.refused == []
    assert _in_field_result(_document_root(result.output)) == []
    assert adapter.open(result.output).paragraphs[0].text == "See Doe 2020 says"


def test_replace_across_a_field_boundary_is_refused(
    adapter: DocxAdapter, tmp_path: Path
):
    path = _field_doc(tmp_path, _zotero_paragraph())
    before = path.read_bytes()
    view = adapter.open(path)
    pid = view.paragraphs[0].id
    start, end = _span(view, pid, "rose (Doe")
    result = adapter.apply(
        view,
        _txn(
            Replace(
                paragraph_id=pid,
                start=start,
                end=end,
                text="fell (Roe",
                expect="rose (Doe",
            )
        ),
        output=tmp_path / "out.docx",
    )
    assert result.applied == 0 and "field result" in result.refused[0]
    assert adapter.open(result.output).paragraphs[0].text == view.paragraphs[0].text
    assert path.read_bytes() == before


def test_replace_and_delete_inside_a_cross_reference_result_are_refused(
    adapter: DocxAdapter, tmp_path: Path
):
    from andamentum.compose.harness.spike.make import complex_field

    path = _field_doc(
        tmp_path,
        run("See ") + complex_field(" REF tab1 \\h ", "Table 1") + run(" below."),
    )
    view = adapter.open(path)
    pid = view.paragraphs[0].id
    start, end = _span(view, pid, "Table 1")
    result = adapter.apply(
        view,
        _txn(
            Replace(
                paragraph_id=pid, start=start, end=end, text="Table 2", expect="Table 1"
            ),
            Delete(paragraph_id=pid, start=start, end=start + 1, expect="T"),
            Insert(paragraph_id=pid, at=start + 3, text="X"),
        ),
        output=tmp_path / "out.docx",
    )
    assert result.applied == 0 and len(result.refused) == 3
    assert _in_field_result(_document_root(result.output)) == []


def test_a_comment_on_a_field_puts_its_markers_around_the_field(
    adapter: DocxAdapter, tmp_path: Path
):
    path = _field_doc(tmp_path, _zotero_paragraph())
    view = adapter.open(path)
    pid = view.paragraphs[0].id
    start, end = _span(view, pid, "Doe")  # part of the result only
    result = adapter.apply(
        view,
        _txn(CommentOp(paragraph_id=pid, start=start, end=end, text="Check this")),
        output=tmp_path / "out.docx",
    )
    assert result.applied == 1 and result.refused == []
    assert _in_field_result(_document_root(result.output)) == []
    out = adapter.open(result.output)
    (comment,) = out.comments
    # The anchor covers the whole result, not the three characters asked for.
    assert out.paragraphs[0].text[comment.start : comment.end] == "Doe 2020"


def test_a_comment_over_a_whole_field_and_its_neighbours_is_unchanged(
    adapter: DocxAdapter, tmp_path: Path
):
    path = _field_doc(tmp_path, _zotero_paragraph())
    view = adapter.open(path)
    pid = view.paragraphs[0].id
    start, end = _span(view, pid, "rose (Doe 2020) overall")
    result = adapter.apply(
        view,
        _txn(CommentOp(paragraph_id=pid, start=start, end=end, text="Reword")),
        output=tmp_path / "out.docx",
    )
    assert result.refused == []
    assert _in_field_result(_document_root(result.output)) == []
    (comment,) = adapter.open(result.output).comments
    assert (comment.start, comment.end) == (start, end)


def test_bind_inside_a_field_is_refused_and_over_a_whole_result_wraps_the_field(
    adapter: DocxAdapter, tmp_path: Path
):
    path = _field_doc(tmp_path, _zotero_paragraph())
    view = adapter.open(path)
    pid = view.paragraphs[0].id
    start, end = _span(view, pid, "Doe")
    refused = adapter.apply(
        view,
        _txn(Bind(paragraph_id=pid, start=start, end=end, anchor_id="n1", kind="span")),
        output=tmp_path / "refused.docx",
    )
    assert refused.applied == 0 and "field result" in refused.refused[0]

    start, end = _span(view, pid, "Doe 2020")
    wrapped = adapter.apply(
        view,
        _txn(
            Bind(paragraph_id=pid, start=start, end=end, anchor_id="c1", kind="claim")
        ),
        output=tmp_path / "wrapped.docx",
    )
    assert wrapped.applied == 1
    assert _in_field_result(_document_root(wrapped.output)) == []
    (anchor,) = adapter.open(wrapped.output).anchors
    assert anchor.id == "c1"
    assert (
        adapter.open(wrapped.output).paragraphs[0].text[anchor.start : anchor.end]
        == "Doe 2020"
    )


def test_insert_at_the_edge_of_an_anchor_does_not_widen_it(
    adapter: DocxAdapter, tmp_path: Path
):
    from andamentum.compose.harness.spike.make import content_control

    path = _field_doc(
        tmp_path,
        run("The dose was ") + content_control("n1", "number", run("42")) + run(" mg."),
    )
    view = adapter.open(path)
    pid = view.paragraphs[0].id
    at = view.paragraphs[0].text.index("42") + 2
    result = adapter.apply(
        view,
        _txn(Insert(paragraph_id=pid, at=at, text="0")),
        output=tmp_path / "out.docx",
    )
    assert result.refused == []
    out = adapter.open(result.output)
    (anchor,) = out.anchors
    assert out.paragraphs[0].text[anchor.start : anchor.end] == "42"
    assert out.paragraphs[0].text == "The dose was 420 mg."


def test_inserted_text_beside_a_superscript_citation_is_not_superscript(
    adapter: DocxAdapter, tmp_path: Path
):

    superscript = (
        '<w:r><w:rPr><w:vertAlign w:val="superscript"/></w:rPr>'
        '<w:t xml:space="preserve">1</w:t></w:r>'
    )
    field = (
        '<w:r><w:fldChar w:fldCharType="begin"/></w:r>'
        '<w:r><w:instrText xml:space="preserve">ADDIN EN.CITE x</w:instrText></w:r>'
        '<w:r><w:fldChar w:fldCharType="separate"/></w:r>'
        + superscript
        + '<w:r><w:fldChar w:fldCharType="end"/></w:r>'
    )
    path = _field_doc(tmp_path, run("Sodium rose") + field)
    view = adapter.open(path)
    text = view.paragraphs[0].text
    result = adapter.apply(
        view,
        _txn(Insert(paragraph_id=view.paragraphs[0].id, at=len(text), text=" again")),
        output=tmp_path / "out.docx",
    )
    root = _document_root(result.output)
    (inserted,) = [e for e in root.iter(w("ins")) if e.get(w("author")) == AUTHOR]
    assert inserted.find(f".//{w('vertAlign')}") is None
