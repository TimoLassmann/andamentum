"""Placement, styles and the three-step build of drafted sections into a document."""

from __future__ import annotations

from pathlib import Path

import docx
import pytest

from andamentum.compose.harness.docx.reader import open_docx
from andamentum.compose.harness.draft.assemble import (
    DraftError,
    PlacedParagraph,
    Placement,
    SectionDraft,
    Styles,
    build_round,
    carry_ids,
    place_sections,
    resolve_styles,
    section_spans,
)
from andamentum.compose.harness.draft.body import parse_body
from andamentum.compose.harness.draft.claims import SourceInfo, analyse_section
from andamentum.compose.harness.draft.run import choose_targets, write_base_document
from andamentum.compose.harness.draft.works import CitableWork
from andamentum.compose.schemas import Fact, Outline, SectionPlan, Source

from ._draft_helpers import WORK_ABSTRACT, WORK_CSL, install_script
from ._helpers import make_view


def collaborator(path: Path) -> Path:
    document = docx.Document()
    document.add_heading("Collaborator draft", 0)
    document.add_heading("Introduction", 1)
    document.add_paragraph(
        "Reading support matters for early learners in every school."
    )
    document.add_heading("Methods", 1)
    document.add_paragraph(
        "We enrolled children from three schools and logged every session."
    )
    document.add_heading("Limitations", 1)
    document.save(str(path))
    return path


def plan(title: str) -> SectionPlan:
    return SectionPlan(
        title=title, purpose="", key_points=[], fact_ids=[], target_words=100
    )


def test_section_spans_use_normalised_headings() -> None:
    view = make_view(
        ("Title", None, "Title"),
        ("1. Introduction", 1),
        ("text", None),
        ("2.1 Methods", 1),
        ("more", None),
        ("last", None),
    )
    assert section_spans(view) == {"introduction": (1, 2), "methods": (3, 5)}


def test_new_sections_go_after_the_nearest_earlier_section() -> None:
    view = make_view(
        ("Intro", 1), ("a", None), ("Methods", 1), ("b", None), ("References", 1)
    )
    titles = ["Intro", "Methods", "Results", "Discussion"]
    placed = place_sections(titles, {"Results", "Discussion"}, view)
    assert placed["Results"] == Placement(anchor="p3", heading=True)
    assert placed["Discussion"] == Placement(anchor="p3", heading=True)


def test_a_new_first_section_goes_before_the_first_later_one() -> None:
    view = make_view(("Title", None, "Title"), ("Methods", 1), ("b", None))
    placed = place_sections(["Intro", "Methods"], {"Intro"}, view)
    assert placed["Intro"] == Placement(anchor="p0", heading=True)


def test_without_any_neighbour_a_section_goes_to_the_end() -> None:
    view = make_view(("Title", None, "Title"), ("text", None))
    placed = place_sections(["Intro"], {"Intro"}, view)
    assert placed["Intro"] == Placement(anchor="p1", heading=True)


def test_a_named_existing_section_is_appended_under_its_heading() -> None:
    view = make_view(
        ("Intro", 1), ("a", None), ("b", None), ("Methods", 1), ("c", None)
    )
    placed = place_sections(["Intro", "Methods"], {"Intro"}, view)
    assert placed == {"Intro": Placement(anchor="p2", heading=False)}


def test_carry_ids_follows_growing_paragraphs() -> None:
    before = make_view(("Heading", 1), ("Fidelity.", None), ("Tail", None))
    after = make_view(
        ("Heading", 1),
        ("Fidelity (Okafor 2020).", None),
        ("Tail", None),
        ("Entry", None),
    )
    assert carry_ids(before, after) == {"p0": "p0", "p1": "p1", "p2": "p2"}


def test_carry_ids_survives_an_inserted_paragraph_in_the_middle() -> None:
    before = make_view(("Heading", 1), ("A claim.", None), ("Tail", None))
    after = make_view(
        ("Heading", 1), ("A claim (X 2020).", None), ("New entry", None), ("Tail", None)
    )
    assert carry_ids(before, after) == {"p0": "p0", "p1": "p1", "p2": "p3"}


def test_carry_ids_raises_when_a_paragraph_vanished() -> None:
    before = make_view(("Heading", 1), ("Body text", None))
    after = make_view(("Heading", 1))
    with pytest.raises(DraftError, match="missing"):
        carry_ids(before, after)


def test_offsets_map_around_a_list_marker_and_opened_spaces() -> None:
    paragraph = PlacedParagraph(
        section="s1",
        role="body",
        index=0,
        text="- ab cd.",
        style=None,
        shift=2,
        gaps=(2,),
    )
    assert paragraph.start(0) == 2
    assert paragraph.start(2) == 5  # a citation goes after the opened space
    assert paragraph.end(2) == 4  # a range ending there leaves the space outside
    assert paragraph.end(4) == 7


def test_styles_come_from_the_document(tmp_path: Path) -> None:
    path = collaborator(tmp_path / "in.docx")
    styles = resolve_styles(path, open_docx(path), heading_level=1)
    assert styles == Styles(
        heading="heading 1", body=None, bullet="List Bullet", number="List Number"
    )


def test_a_template_without_the_heading_style_is_an_error(tmp_path: Path) -> None:
    path = collaborator(tmp_path / "in.docx")
    with pytest.raises(DraftError, match="Heading 10"):
        resolve_styles(path, open_docx(path), heading_level=10)


def test_new_base_document_drops_template_body(tmp_path: Path) -> None:
    template = collaborator(tmp_path / "template.docx")
    base = tmp_path / "base.docx"
    write_base_document(base, template=template, title="A title", scratch=tmp_path)
    view = open_docx(base)
    assert [p.text for p in view.paragraphs] == ["A title"]
    assert view.paragraphs[0].style.lower() == "title"


def test_a_dotx_template_is_read_as_a_template(tmp_path: Path) -> None:
    import zipfile

    source = collaborator(tmp_path / "plain.docx")
    template = tmp_path / "journal.dotx"
    with zipfile.ZipFile(source) as src, zipfile.ZipFile(template, "w") as dst:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename == "[Content_Types].xml":
                data = data.replace(
                    b"wordprocessingml.document.main+xml",
                    b"wordprocessingml.template.main+xml",
                )
            dst.writestr(item, data)
    base = tmp_path / "base.docx"
    write_base_document(base, template=template, title="T", scratch=tmp_path)
    assert [p.text for p in open_docx(base).paragraphs] == ["T"]


# ------------------------------------------------------------ choose_targets


def outline(*titles: str) -> Outline:
    return Outline(sections=[plan(t) for t in titles])


def test_a_new_document_writes_every_planned_section() -> None:
    chosen, targets = choose_targets(outline("A", "B"), task="anything", view=None)
    assert targets == {"A", "B"}
    assert chosen == outline("A", "B")


def test_named_sections_are_the_targets_into_a_document(tmp_path: Path) -> None:
    view = open_docx(collaborator(tmp_path / "in.docx"))
    _, targets = choose_targets(
        outline("Introduction", "Methods", "Results", "Discussion"),
        task="Write the discussion and the results into this paper",
        view=view,
    )
    assert targets == {"Results", "Discussion"}


def test_without_names_the_missing_or_empty_sections_are_targets(
    tmp_path: Path,
) -> None:
    view = open_docx(collaborator(tmp_path / "in.docx"))
    _, targets = choose_targets(
        outline("Introduction", "Methods", "Limitations", "Discussion"),
        task="Finish the paper",
        view=view,
    )
    assert targets == {"Limitations", "Discussion"}  # Limitations exists but is empty


def test_a_named_heading_the_outline_lacks_joins_the_outline(tmp_path: Path) -> None:
    view = open_docx(collaborator(tmp_path / "in.docx"))
    chosen, targets = choose_targets(
        outline("Introduction", "Results"),
        task="Expand the limitations",
        view=view,
    )
    assert targets == {"Limitations"}
    assert [p.title for p in chosen.sections] == [
        "Introduction",
        "Results",
        "Limitations",
    ]


def test_nothing_to_draft_is_a_usage_error(tmp_path: Path) -> None:
    from andamentum.compose.harness.draft import DraftUsageError

    view = open_docx(collaborator(tmp_path / "in.docx"))
    with pytest.raises(DraftUsageError, match="nothing to draft"):
        choose_targets(outline("Introduction", "Methods"), task="Polish it", view=view)


# ----------------------------------------------------------------- the build


WORK = CitableWork(csl=WORK_CSL, abstract=WORK_ABSTRACT)
FACTS = {
    "f1": Fact(id="f1", source_id="src1", statement="s", quote="q", start=0, end=1),
    "f2": Fact(id="f2", source_id="src2", statement="s", quote="q", start=0, end=1),
}
SOURCE = Source(
    id="src1",
    path="p",
    title="t",
    sha256="0" * 64,
    chars=1,
    kind="file",
    label="notes.md",
)
INFOS = {
    "src1": SourceInfo(source=SOURCE, work=WORK),
    "src2": SourceInfo(source=SOURCE.model_copy(update={"id": "src2"})),
}


async def section(
    key: str, title: str, body: str, placement: Placement
) -> SectionDraft:
    draft = SectionDraft(
        key=key, title=title, body=body, plan=plan(title), placement=placement
    )
    draft.paragraphs = parse_body(body)
    draft.analysis = await analyse_section(
        draft.paragraphs,
        facts=FACTS,
        infos=INFOS,
        catalogue=None,
        taken_anchors=set(),
        model="stub",
        cache={},
        log=None,
    )
    return draft


async def test_a_build_inserts_text_then_marks_then_citations(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    install_script(monkeypatch)
    base = collaborator(tmp_path / "in.docx")
    view = open_docx(base)
    styles = resolve_styles(base, view, heading_level=1)
    anchor = view.paragraphs[-2].id  # the Methods paragraph
    draft = await section(
        "s1",
        "Results",
        "Logging improves fidelity [[f1]].\n\n- first point\n- second point",
        Placement(anchor=anchor, heading=True),
    )
    built = build_round(
        base=base,
        round_dir=tmp_path / "round",
        sections=[draft],
        styles=styles,
        facts=FACTS,
        infos=INFOS,
        document_notes=["A question about the whole document."],
        reason="test",
    )
    texts = [p.text for p in built.view.paragraphs]
    assert texts[:5] == [
        "Collaborator draft",
        "Introduction",
        "Reading support matters for early learners in every school.",
        "Methods",
        "We enrolled children from three schools and logged every session.",
    ]
    assert texts[5] == "Results"
    assert texts[6] == "Logging improves fidelity (Okafor 2020)."
    assert texts[7:9] == ["first point", "second point"]
    assert texts[9] == "Limitations"
    assert "References" in texts
    assert built.citations == 1 and built.works == ["doi:10.5555/saltlog.2020.1"]
    # the heading and body of the new section are owned by it, nothing else is
    owned = [p.text for p in built.view.paragraphs if p.id in built.owner]
    assert owned[0] == "Results" and "first point" in owned
    assert all(built.owner[i] == "s1" for i in built.owner)
    assert [p.style for p in built.view.paragraphs][7:9] == [
        "List Bullet",
        "List Bullet",
    ]
    assert [c.text for c in built.view.comments] == [
        "A question about the whole document."
    ]
    assert not built.refused and not built.notes
    assert (tmp_path / "round" / "3-cites.docx").is_file()
    # the base is untouched
    assert open_docx(base).text() == view.text()


async def test_the_build_reads_back_what_it_wrote(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    install_script(monkeypatch)
    base = collaborator(tmp_path / "in.docx")
    view = open_docx(base)
    styles = resolve_styles(base, view, heading_level=1)
    draft = await section(
        "s1",
        "Results",
        "A plain claim.",
        Placement(anchor=view.paragraphs[-1].id, heading=True),
    )
    built = build_round(
        base=base,
        round_dir=tmp_path / "round",
        sections=[draft],
        styles=styles,
        facts=FACTS,
        infos=INFOS,
        document_notes=[],
        reason="test",
    )
    assert built.citations == 0 and built.works == []
    assert "References" not in [p.text for p in built.view.paragraphs]
