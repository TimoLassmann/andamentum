"""The planted-defect generator: files open, ground truth matches the files, fixtures cover them."""

from __future__ import annotations

import hashlib
import json
import re
from pathlib import Path
from typing import Any

import pytest
from docx import Document
from docx.oxml.ns import qn

from andamentum.compose.harness.bench import planted
from andamentum.compose.harness.bench.planted import (
    FIXTURES_DIR,
    FixtureResponse,
    KINDS,
    Stat,
    build_fixtures,
    generate,
    load_fixtures,
    pools,
    stat_p_value,
)


@pytest.fixture(scope="module")
def out_dir(tmp_path_factory: pytest.TempPathFactory) -> Path:
    out = tmp_path_factory.mktemp("planted")
    generate(out, seed=0)
    return out


def _message(response: FixtureResponse) -> dict[str, Any]:
    assert isinstance(response.body, dict)
    return response.body["message"]


def _items(response: FixtureResponse) -> list[dict[str, Any]]:
    return _message(response)["items"]


def _truths(out: Path) -> list[dict[str, Any]]:
    return [json.loads(p.read_text()) for p in sorted(out.glob("*.truth.json"))]


def _paragraph_texts(path: Path) -> list[str]:
    doc = Document(str(path))
    return ["".join(t.text or "" for t in p._p.iter(qn("w:t"))) for p in doc.paragraphs]


def _field_instructions(path: Path) -> list[str]:
    doc = Document(str(path))
    return [i.text or "" for i in doc.element.body.iter(qn("w:instrText"))]


def _bookmarks(path: Path) -> set[str]:
    doc = Document(str(path))
    return {
        b.get(qn("w:name")) or "" for b in doc.element.body.iter(qn("w:bookmarkStart"))
    }


def test_fixtures_are_committed_and_current() -> None:
    assert load_fixtures() == build_fixtures()
    assert {p.name for p in FIXTURES_DIR.glob("*.json")} == {
        "crossref.json",
        "urls.json",
    }


def test_totals(out_dir: Path) -> None:
    manifest = json.loads((out_dir / "manifest.json").read_text())
    totals = manifest["totals"]
    assert totals["files"] >= 12
    assert totals["defects"] >= 40
    assert totals["clean"] >= 4
    assert set(totals["kinds"]) == set(KINDS)
    assert all(n >= 1 for n in totals["kinds"].values())
    for entry in manifest["files"]:
        assert (
            hashlib.sha256((out_dir / entry["docx"]).read_bytes()).hexdigest()
            == entry["sha256"]
        )


def test_clean_files_have_no_defects(out_dir: Path) -> None:
    clean = [t for t in _truths(out_dir) if t["clean"]]
    assert len(clean) >= 4
    assert all(t["defects"] == [] and t["valid_items"] for t in clean)


def test_files_open_and_locations_match(out_dir: Path) -> None:
    for truth in _truths(out_dir):
        path = out_dir / truth["file"]
        texts = _paragraph_texts(path)
        assert len(texts) == truth["paragraph_count"]
        for rec in [*truth["defects"], *truth["valid_items"]]:
            assert rec["locations"]
            for loc in rec["locations"]:
                assert texts[loc["paragraph_index"]] == loc["paragraph_text"]
                assert loc["snippet"] in loc["paragraph_text"]
        assert len({d["id"] for d in truth["defects"]}) == len(truth["defects"])


def test_dois_and_urls_are_covered_by_fixtures(out_dir: Path) -> None:
    fx = load_fixtures()
    works, updates, urls = fx.crossref.works, fx.crossref.updates, fx.urls.urls
    for truth in _truths(out_dir):
        for doi in truth["dois"]:
            assert doi in works and doi in updates
        for url in truth["urls"]:
            assert url in urls
        for d in truth["defects"]:
            if d["kind"] == "fabricated_doi":
                assert works[d["doi"]].status == 404
                assert _items(updates[d["doi"]]) == []
            elif d["kind"] == "retracted_doi":
                assert works[d["doi"]].status == 200
                (item,) = _items(updates[d["doi"]])
                assert item["update-to"][0]["type"] == "retraction"
                assert item["update-to"][0]["DOI"] == d["doi"]
            elif d["kind"] == "broken_url":
                entry = urls[d["url"]]
                if d["variant"].startswith("http_"):
                    assert entry.status == int(d["variant"][-3:])
                else:
                    assert entry.status is None and entry.error == d["variant"]
        for v in truth["valid_items"]:
            if v["kind"] == "valid_doi":
                assert works[v["doi"]].status == 200
                assert not _items(updates[v["doi"]])
            if v["kind"] == "valid_url":
                assert urls[v["url"]].status == 200


def test_valid_doi_metadata_matches_the_bibliography(out_dir: Path) -> None:
    works = load_fixtures().crossref.works
    for truth in _truths(out_dir):
        texts = _paragraph_texts(out_dir / truth["file"])
        for v in truth["valid_items"]:
            if v["kind"] != "valid_doi":
                continue
            msg = _message(works[v["doi"]])
            bib = next(t for t in texts if v["doi"] in t)
            assert msg["title"][0] in bib
            assert msg["author"][0]["family"] in bib


def test_zotero_json_parses_unless_planted_malformed(out_dir: Path) -> None:
    malformed = good = 0
    for truth in _truths(out_dir):
        planted_bad = sum(
            1 for d in truth["defects"] if d["kind"] == "malformed_citation_json"
        )
        bad = 0
        for instr in _field_instructions(out_dir / truth["file"]):
            m = re.match(
                r"\s*ADDIN (?:ZOTERO_ITEM CSL_CITATION|CSL_CITATION) (.*?)\s*$",
                instr,
                re.S,
            )
            if not m:
                continue
            try:
                payload = json.loads(m.group(1))
                assert payload["citationItems"][0]["itemData"]["DOI"]
                good += 1
            except json.JSONDecodeError:
                bad += 1
        assert bad == planted_bad
        malformed += bad
    assert malformed >= 4 and good > malformed


def test_field_flavours_present(out_dir: Path) -> None:
    seen: set[str] = set()
    for truth in _truths(out_dir):
        for instr in _field_instructions(out_dir / truth["file"]):
            for marker in (
                "ZOTERO_ITEM",
                "ZOTERO_BIBL",
                "EN.CITE",
                "EN.REFLIST",
                "CSL_CITATION",
                " REF ",
                " SEQ ",
            ):
                if marker in instr:
                    seen.add(marker)
    assert seen == {
        "ZOTERO_ITEM",
        "ZOTERO_BIBL",
        "EN.CITE",
        "EN.REFLIST",
        "CSL_CITATION",
        " REF ",
        " SEQ ",
    }


def test_crossrefs_dangling_vs_valid(out_dir: Path) -> None:
    for truth in _truths(out_dir):
        path = out_dir / truth["file"]
        marks = _bookmarks(path)
        targets = []
        for instr in _field_instructions(path):
            m = re.match(r"\s*REF (\S+)", instr)
            if m:
                targets.append(m.group(1))
        dangling = [t for t in targets if t not in marks]
        assert len(dangling) == sum(
            1 for d in truth["defects"] if d["kind"] == "dangling_crossref"
        )


def test_missing_figures_have_no_caption(out_dir: Path) -> None:
    for truth in _truths(out_dir):
        texts = _paragraph_texts(out_dir / truth["file"])
        captioned = {
            int(m.group(1))
            for t in texts
            for m in [re.match(r"Figure (\d+)\. ", t)]
            if m
        }
        for d in truth["defects"]:
            if d["kind"] == "missing_figure":
                assert d["number"] not in captioned
        for v in truth["valid_items"]:
            if v["kind"] == "valid_figure_mention":
                assert 1 in captioned


_STAT = re.compile(
    r"(?P<k>t|F|χ²|r|z)(?:\((?P<df>[\d, ]+)\))? = (?P<v>-?\.?\d*\.?\d+), p (?P<op>[=<]) (?P<p>\.\d+)"
)


def _parse_stat(text: str) -> tuple[str, tuple[int, ...], float, str, float]:
    m = _STAT.fullmatch(text)
    assert m, text
    dfs = tuple(int(x) for x in m["df"].split(",")) if m["df"] else ()
    return m["k"], dfs, float(m["v"]), m["op"], float(m["p"])


def _recompute(kind: str, dfs: tuple[int, ...], value: float) -> float:
    key = {"χ²": "chi2"}.get(kind, kind)
    return stat_p_value(key, dfs, value)  # type: ignore[arg-type]


def test_statistics_recompute(out_dir: Path) -> None:
    flips = 0
    for truth in _truths(out_dir):
        for v in truth["valid_items"]:
            if v["kind"] != "valid_statistic":
                continue
            kind, dfs, value, op, reported = _parse_stat(v["locations"][0]["snippet"])
            true_p = _recompute(kind, dfs, value)
            if op == "<":
                assert true_p < 0.0005
            else:
                assert abs(true_p - reported) <= 0.0005 + 1e-9
        for d in truth["defects"]:
            if d["kind"] != "inconsistent_statistic":
                continue
            kind, dfs, value, op, reported = _parse_stat(d["locations"][0]["snippet"])
            true_p = _recompute(kind, dfs, value)
            assert abs(true_p - d["recomputed_p"]) < 1e-5
            assert abs(true_p - reported) > 0.002
            flipped = (true_p < 0.05) != (reported < 0.05)
            assert flipped == d["decision_flips"]
            assert d["expected_severity"] == ("error" if flipped else "warning")
            flips += flipped
    assert flips >= 3


def test_abstract_numbers(out_dir: Path) -> None:
    for truth in _truths(out_dir):
        texts = _paragraph_texts(out_dir / truth["file"])
        abstract = texts[truth["abstract_paragraph_index"]]
        refs_at = next(i for i, t in enumerate(texts) if t == "References")
        body = "\n".join(texts[truth["abstract_paragraph_index"] + 2 : refs_at])
        for d in truth["defects"]:
            if d["kind"] == "abstract_number_missing":
                assert d["number"] in abstract
                assert (
                    re.search(rf"(?<![\d.]){re.escape(d['number'])}(?!\d)", body)
                    is None
                )
        for v in truth["valid_items"]:
            if v["kind"] == "valid_abstract_number":
                assert re.search(rf"(?<![\d.]){re.escape(v['number'])}(?!\d)", body)


def test_deterministic_and_seed_sensitive(tmp_path: Path) -> None:
    a, b, c = tmp_path / "a", tmp_path / "b", tmp_path / "c"
    generate(a, seed=7)
    generate(b, seed=7)
    generate(c, seed=8)
    for f in sorted(a.iterdir()):
        assert f.read_bytes() == (b / f.name).read_bytes()
    assert (a / "d01_mixed_a.docx").read_bytes() != (
        c / "d01_mixed_a.docx"
    ).read_bytes()


@pytest.mark.parametrize("seed", [1, 2, 3, 12345])
def test_other_seeds_meet_the_floors_and_stay_covered(
    tmp_path: Path, seed: int
) -> None:
    manifest = generate(tmp_path, seed=seed)
    assert manifest.totals.defects >= 40
    fx = load_fixtures()
    for truth in _truths(tmp_path):
        assert all(d in fx.crossref.works for d in truth["dois"])
        assert all(u in fx.urls.urls for u in truth["urls"])
        _paragraph_texts(tmp_path / truth["file"])


def test_pools_cover_each_variant() -> None:
    p = pools()
    assert {u.variant for u in p.urls_broken} == {
        "http_404",
        "http_410",
        "dns_failure",
        "timeout",
    }
    assert any(u.variant == "redirect" for u in p.urls_ok)
    assert any(r.doi.startswith("10.1016/") for r in p.retracted)


def test_stat_helper_matches_known_values() -> None:
    assert round(stat_p_value("t", (28,), 2.05), 3) == pytest.approx(0.05, abs=0.001)
    assert Stat("z", (), 1.96, stat_p_value("z", (), 1.96)).p == pytest.approx(
        0.05, abs=0.001
    )


def test_cli(tmp_path: Path) -> None:
    out = tmp_path / "cli"
    assert planted.main(["--out", str(out), "--seed", "0"]) == 0
    assert (out / "manifest.json").exists()
    fx = tmp_path / "fx"
    assert planted.main(["--write-fixtures", str(fx)]) == 0
    assert (fx / "crossref.json").read_text() == (
        FIXTURES_DIR / "crossref.json"
    ).read_text()


@pytest.mark.parametrize("seed", [0, 1, 2, 3])
def test_only_planted_malformed_fields_fail_to_parse(tmp_path: Path, seed: int):
    """The reader finds a parse error in exactly the planted malformed citation fields."""
    from andamentum.compose.harness.docx.reader import open_docx

    generate(tmp_path, seed=seed)
    for truth in _truths(tmp_path):
        view = open_docx(tmp_path / truth["file"])
        broken = [f for f in view.fields if "parse_error" in f.data]
        planted_count = sum(
            1 for d in truth["defects"] if d["kind"] == "malformed_citation_json"
        )
        assert len(broken) == planted_count, truth["file"]


def test_reader_flags_exactly_the_planted_dangling_references(out_dir: Path):
    from andamentum.compose.harness.docx.reader import open_docx

    for truth in _truths(out_dir):
        view = open_docx(out_dir / truth["file"])
        refs = [
            f
            for f in view.fields
            if f.kind == "crossref"
            and f.data.get("field") in {"REF", "PAGEREF", "NOTEREF"}
        ]
        assert all(isinstance(f.data["target_exists"], bool) for f in refs)
        missing = [f for f in refs if f.data["target_exists"] is False]
        planted_count = sum(
            1 for d in truth["defects"] if d["kind"] == "dangling_crossref"
        )
        assert len(missing) == planted_count, truth["file"]
        assert all(
            "target_exists" not in f.data
            for f in view.fields
            if f.data.get("field") == "SEQ"
        )
