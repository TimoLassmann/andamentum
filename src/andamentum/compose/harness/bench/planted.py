"""Planted-defect generator: synthetic .docx manuscripts plus ground-truth JSON.

Every manuscript is invented (topics, authors, journals, DOIs, URLs). Defects
are planted into otherwise valid documents so a checker can be scored on what
it finds, what it misses and what it flags wrongly. The document XML is real
OOXML: Zotero / Mendeley / EndNote complex fields, REF fields with bookmarks,
SEQ captions, hyperlinks.

Defect kinds (``KINDS``): ``fabricated_doi``, ``retracted_doi``, ``broken_url``,
``dangling_crossref``, ``malformed_citation_json``, ``abstract_number_missing``,
``inconsistent_statistic``, ``missing_figure``.

Network is never touched. Every DOI and URL a document can contain is covered by
the recorded fixtures in ``bench/fixtures/`` (``crossref.json``, ``urls.json``).
The pools behind them are fixed and independent of the seed, so one committed
fixture set serves every seed. ``build_fixtures`` regenerates them and a test
pins the committed files to its output.

    uv run python -m andamentum.compose.harness.bench.planted --out DIR --seed 0

Output per file: ``<name>.docx``, ``<name>.truth.json``; plus ``manifest.json``.
The same seed gives byte-identical files (zip timestamps and core properties are
pinned).
"""

from __future__ import annotations

import argparse
import datetime as dt
import functools
import hashlib
import json
import math
import random
import re
import sys
import zipfile
from xml.sax.saxutils import escape as xml_escape
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Literal

from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pydantic import BaseModel, ConfigDict, Field
from scipy import stats as sps

KINDS: tuple[str, ...] = (
    "fabricated_doi",
    "retracted_doi",
    "broken_url",
    "dangling_crossref",
    "malformed_citation_json",
    "abstract_number_missing",
    "inconsistent_statistic",
    "missing_figure",
)

CitationStyle = Literal[
    "zotero", "mendeley", "endnote", "numbered", "author_year", "mixed"
]

FIXTURES_DIR = Path(__file__).parent / "fixtures"
SCHEMA_VERSION = 1
_FIXED_TIME = dt.datetime(2026, 1, 1, 0, 0, 0)
_MAX_ATTEMPTS = 40

# ---------------------------------------------------------------------------
# Fixed pools (seed-independent) and the recorded fixtures built from them
# ---------------------------------------------------------------------------

_FAMILIES = [
    "Quillfeather",
    "Marrowby",
    "Brindlemoor",
    "Ashgrove-Pell",
    "Fennimore",
    "Haldane-Wyck",
    "Pemberlow",
    "Corrigan-Tate",
    "Delacroft",
    "Eastmarch",
    "Garrowby",
    "Hollowell",
    "Ingleby-Rusk",
    "Jessamoor",
    "Kettlewick",
    "Lanternby",
    "Mossgrave",
    "Nettlefold",
    "Oakhollow",
    "Plumtree-Vane",
    "Quarrelby",
    "Ravenscroft",
    "Sallowfield",
    "Thistlewood",
    "Underhill-Cray",
    "Vexworth",
    "Wetherby-Lund",
    "Yarrowdale",
    "Zimmerlith",
    "Abernethy-Roke",
    "Bellweather",
    "Cranmoor",
    "Dunmarrow",
    "Elderthwaite",
    "Foxglove-Hart",
    "Greymantle",
    "Hartlebury-Nye",
    "Inkwell",
    "Juniperby",
    "Kestrelmoor",
]
_ADJ = [
    "Randomised",
    "Pragmatic",
    "Longitudinal",
    "Exploratory",
    "Cluster",
    "Open-label",
    "Blinded",
    "Crossover",
]
_NOUN = [
    "trial of timed hydration",
    "cohort study of shift patterns",
    "evaluation of a pacing aid",
    "survey of reporting habits",
    "analysis of cold-chain logging",
    "pilot of peer coaching",
    "audit of handover notes",
    "comparison of rest schedules",
]
_CONTEXT = [
    "community clinics",
    "regional depots",
    "coastal schools",
    "night-shift crews",
    "rural pharmacies",
    "small laboratories",
    "river-port workers",
    "orchard teams",
]
_JOURNALS = [
    "Journal of Applied Brinework",
    "Annals of Quillwork Medicine",
    "Harrow Journal of Field Physiology",
    "Meridian Reports in Behavioural Health",
    "Cinder Bay Clinical Letters",
    "Orrery Methods in Epidemiology",
]
_URL_A = [
    "saltlog",
    "apollo",
    "quill",
    "marrow",
    "tern",
    "lumen",
    "basalt",
    "cinder",
    "harrow",
    "pelican",
    "wren",
    "kestrel",
    "mistral",
    "orrery",
    "bramble",
]
_URL_B = ["lab", "study", "pilot", "group", "unit", "project"]
_URL_C = ["data", "code", "protocol", "supplement", "materials", "analysis"]

# A real, public, retracted work used as data (the Crossref records below are
# hand-built in the shape of the REST API output, not live recordings).
_REAL_RETRACTED = {
    "doi": "10.1016/S0140-6736(97)11096-0",
    "title": "Ileal-lymphoid-nodular hyperplasia, non-specific colitis, and "
    "pervasive developmental disorder in children",
    "journal": "The Lancet",
    "year": 1998,
    "volume": "351",
    "issue": "9103",
    "pages": "637-641",
    "authors": [
        ("Wakefield", "A. J."),
        ("Murch", "S. H."),
        ("Anthony", "A."),
        ("Linnell", "J."),
        ("Casson", "D. M."),
        ("Malik", "M."),
        ("Berelowitz", "M."),
        ("Dhillon", "A. P."),
        ("Thomson", "M. A."),
        ("Harvey", "P."),
        ("Valentine", "A."),
        ("Davies", "S. E."),
        ("Walker-Smith", "J. A."),
    ],
    "notice_doi": "10.1016/S0140-6736(10)60175-4",
    "notice_source": "publisher",
    "notice_date": [2010, 2, 6],
}


@dataclass(frozen=True)
class Ref:
    key: str
    status: Literal["valid", "fabricated", "retracted"]
    authors: tuple[tuple[str, str], ...]
    year: int
    title: str
    journal: str
    volume: str
    issue: str
    pages: str
    doi: str
    notice_doi: str | None = None
    notice_source: str | None = None
    notice_date: tuple[int, int, int] | None = None


@dataclass(frozen=True)
class UrlEntry:
    url: str
    variant: Literal["ok", "redirect", "http_404", "http_410", "dns_failure", "timeout"]
    final_url: str | None = None


@dataclass(frozen=True)
class Pools:
    valid: tuple[Ref, ...]
    fabricated: tuple[Ref, ...]
    retracted: tuple[Ref, ...]
    urls_ok: tuple[UrlEntry, ...]
    urls_broken: tuple[UrlEntry, ...]


def _make_ref(
    rng: random.Random,
    *,
    key: str,
    status: Literal["valid", "fabricated", "retracted"],
    doi: str,
    used: set[str],
) -> Ref:
    while True:
        title = f"{rng.choice(_ADJ)} {rng.choice(_NOUN)} in {rng.choice(_CONTEXT)}"
        if title not in used:
            used.add(title)
            break
    n_auth = rng.randint(1, 4)
    families = rng.sample(_FAMILIES, n_auth)
    authors = tuple((f, f"{rng.choice('ABCDEFGHJKLMNPRST')}.") for f in families)
    volume = str(rng.randint(3, 40))
    first = rng.randint(1, 900)
    return Ref(
        key=key,
        status=status,
        authors=authors,
        year=rng.randint(2009, 2025),
        title=title,
        journal=rng.choice(_JOURNALS),
        volume=volume,
        issue=str(rng.randint(1, 6)),
        pages=f"{first}-{first + rng.randint(6, 14)}",
        doi=doi,
    )


@functools.cache
def pools() -> Pools:
    rng = random.Random("andamentum-bench-pools-v1")
    used: set[str] = {str(_REAL_RETRACTED["title"])}
    valid = tuple(
        _make_ref(
            rng,
            key=f"v{i:03d}",
            status="valid",
            doi=f"10.5555/am.bench.v{i:03d}",
            used=used,
        )
        for i in range(60)
    )
    fabricated = tuple(
        _make_ref(
            rng,
            key=f"x{i:03d}",
            status="fabricated",
            doi=f"10.5555/am.bench.x{i:03d}",
            used=used,
        )
        for i in range(12)
    )
    retracted: list[Ref] = []
    r = _REAL_RETRACTED
    retracted.append(
        Ref(
            key="r000",
            status="retracted",
            authors=tuple(r["authors"]),  # type: ignore[arg-type]
            year=int(r["year"]),  # type: ignore[call-overload]
            title=str(r["title"]),
            journal=str(r["journal"]),
            volume=str(r["volume"]),
            issue=str(r["issue"]),
            pages=str(r["pages"]),
            doi=str(r["doi"]),
            notice_doi=str(r["notice_doi"]),
            notice_source=str(r["notice_source"]),
            notice_date=tuple(r["notice_date"]),  # type: ignore[arg-type]
        )
    )
    for i in range(1, 5):
        base = _make_ref(
            rng,
            key=f"r{i:03d}",
            status="retracted",
            doi=f"10.5555/am.bench.r{i:03d}",
            used=used,
        )
        retracted.append(
            Ref(
                **{
                    **base.__dict__,
                    "notice_doi": f"10.5555/am.bench.r{i:03d}.retraction",
                    "notice_source": "retraction-watch" if i % 2 else "publisher",
                    "notice_date": (
                        base.year + 2,
                        rng.randint(1, 12),
                        rng.randint(1, 28),
                    ),
                }
            )
        )

    hosts: set[str] = set()
    while len(hosts) < 44:
        hosts.add(f"{rng.choice(_URL_A)}-{rng.choice(_URL_B)}")
    host_list = sorted(hosts)
    rng.shuffle(host_list)

    def mk(host: str) -> str:
        return f"https://{rng.choice(_URL_C)}.{host}.example/{rng.choice(_URL_C)}/{rng.randint(10, 99)}"

    ok: list[UrlEntry] = []
    for i, h in enumerate(host_list[:30]):
        u = mk(h)
        if i % 10 == 9:
            ok.append(
                UrlEntry(
                    u, "redirect", final_url=u.replace("https://", "https://www.", 1)
                )
            )
        else:
            ok.append(UrlEntry(u, "ok"))
    variants: list[Literal["http_404", "http_410", "dns_failure", "timeout"]] = [
        "http_404", "http_410", "dns_failure", "timeout", "http_404", "dns_failure",
        "http_404", "http_410", "dns_failure", "timeout", "http_404", "dns_failure",
        "http_404", "timeout",
    ]  # fmt: skip
    broken = tuple(
        UrlEntry(mk(h), v) for h, v in zip(host_list[30:44], variants, strict=True)
    )
    return Pools(valid, fabricated, tuple(retracted), tuple(ok), broken)


class FixtureResponse(BaseModel):
    """One recorded HTTP response: the status and the body (JSON or text)."""

    status: int
    body: object


class UrlFixture(BaseModel):
    """What a URL probe saw: a status, or an error (``dns_failure`` / ``timeout``)."""

    status: int | None
    final_url: str | None
    error: str | None
    redirects: list[dict[str, object]]


class CrossrefFixtures(BaseModel):
    """``works`` (GET /works/<doi>), ``updates`` (GET /works?filter=updates:<doi>)
    and ``search`` (lower-cased title -> candidate work messages, best first)."""

    model_config = ConfigDict(populate_by_name=True)
    meta: dict[str, object] = Field(alias="_meta")
    works: dict[str, FixtureResponse]
    updates: dict[str, FixtureResponse]
    search: dict[str, list[dict[str, object]]]


class UrlFixtures(BaseModel):
    model_config = ConfigDict(populate_by_name=True)
    meta: dict[str, object] = Field(alias="_meta")
    urls: dict[str, UrlFixture]


class Fixtures(BaseModel):
    crossref: CrossrefFixtures
    urls: UrlFixtures


def _csl_authors(ref: Ref) -> list[dict[str, str]]:
    return [{"family": f, "given": g} for f, g in ref.authors]


def _work_message(ref: Ref) -> dict[str, object]:
    return {
        "DOI": ref.doi,
        "URL": f"https://doi.org/{ref.doi}",
        "type": "journal-article",
        "title": [ref.title],
        "container-title": [ref.journal],
        "author": [
            {"given": g, "family": f, "sequence": "first" if i == 0 else "additional"}
            for i, (f, g) in enumerate(ref.authors)
        ],
        "issued": {"date-parts": [[ref.year]]},
        "volume": ref.volume,
        "issue": ref.issue,
        "page": ref.pages,
    }


def _notice_item(ref: Ref) -> dict[str, object]:
    assert ref.notice_doi and ref.notice_date and ref.notice_source
    y, m, d = ref.notice_date
    return {
        "DOI": ref.notice_doi,
        "type": "journal-article",
        "title": [f"Retraction: {ref.title}"],
        "container-title": [ref.journal],
        "issued": {"date-parts": [[y, m, d]]},
        "update-to": [
            {
                "DOI": ref.doi,
                "type": "retraction",
                "label": "Retraction",
                "source": ref.notice_source,
                "updated": {"date-parts": [[y, m, d]]},
            }
        ],
    }


def build_fixtures() -> Fixtures:
    """The recorded-style Crossref and URL fixtures for every pool entry.

    Returns the contents of the two fixture files (``crossref.json`` and
    ``urls.json``).
    ``crossref`` holds ``works`` (DOI -> status + body), ``updates`` (DOI ->
    the ``filter=updates:<doi>`` listing, notices carry ``update-to``) and
    ``search`` (lower-cased title -> candidate items, best first).
    """
    p = pools()
    works: dict[str, object] = {}
    updates: dict[str, object] = {}
    search: dict[str, object] = {}
    all_real = [*p.valid, *p.retracted]
    for ref in all_real:
        msg = _work_message(ref)
        if ref.status == "retracted":
            msg["updated-by"] = [
                {
                    "DOI": ref.notice_doi,
                    "type": "retraction",
                    "label": "Retraction",
                    "source": ref.notice_source,
                }
            ]
        works[ref.doi] = {
            "status": 200,
            "body": {"status": "ok", "message-type": "work", "message": msg},
        }
        search[ref.title.lower()] = [_work_message(ref)]
    for ref in [*p.valid, *p.retracted, *p.fabricated]:
        if ref.status == "retracted":
            updates[ref.doi] = {
                "status": 200,
                "body": {
                    "status": "ok",
                    "message-type": "work-list",
                    "message": {"total-results": 1, "items": [_notice_item(ref)]},
                },
            }
        else:
            updates[ref.doi] = {
                "status": 200,
                "body": {
                    "status": "ok",
                    "message-type": "work-list",
                    "message": {"total-results": 0, "items": []},
                },
            }
    for i, ref in enumerate(p.fabricated):
        works[ref.doi] = {"status": 404, "body": "Resource not found."}
        decoys = [p.valid[(i * 3 + k) % len(p.valid)] for k in range(3)]
        search[ref.title.lower()] = [_work_message(d) for d in decoys]
    urls: dict[str, object] = {}
    for u in [*p.urls_ok, *p.urls_broken]:
        if u.variant == "ok":
            urls[u.url] = {
                "status": 200,
                "final_url": u.url,
                "error": None,
                "redirects": [],
            }
        elif u.variant == "redirect":
            urls[u.url] = {
                "status": 200,
                "final_url": u.final_url,
                "error": None,
                "redirects": [{"status": 301, "location": u.final_url}],
            }
        elif u.variant in ("http_404", "http_410"):
            urls[u.url] = {
                "status": int(u.variant[-3:]),
                "final_url": u.url,
                "error": None,
                "redirects": [],
            }
        else:
            urls[u.url] = {
                "status": None,
                "final_url": None,
                "error": u.variant,
                "redirects": [],
            }
    meta = {
        "provenance": "hand-built in the shape of Crossref REST API output; "
        "not a live recording. Regenerate with "
        "`python -m andamentum.compose.harness.bench.planted --write-fixtures DIR`.",
        "version": SCHEMA_VERSION,
    }
    return Fixtures(
        crossref=CrossrefFixtures.model_validate(
            {
                "_meta": {
                    **meta,
                    "keys": {
                        "works": "DOI -> {status, body} for GET /works/<doi>",
                        "updates": "DOI -> {status, body} for GET /works?filter=updates:<doi>; "
                        "notice items carry update-to (type retraction|correction|expression_of_concern)",
                        "search": "lower-cased title -> candidate work messages, best first "
                        "(fabricated titles get unrelated candidates)",
                    },
                },
                "works": works,
                "updates": updates,
                "search": search,
            }
        ),
        urls=UrlFixtures.model_validate(
            {
                "_meta": {
                    **meta,
                    "keys": {
                        "urls": "URL -> {status|None, final_url, error: dns_failure|timeout|None, redirects}"
                    },
                },
                "urls": urls,
            }
        ),
    )


def write_fixtures(directory: Path) -> list[Path]:
    directory.mkdir(parents=True, exist_ok=True)
    built = build_fixtures()
    out = []
    for name, model in (("crossref", built.crossref), ("urls", built.urls)):
        path = directory / f"{name}.json"
        content = model.model_dump(mode="json", by_alias=True)
        path.write_text(
            json.dumps(content, indent=2, sort_keys=True, ensure_ascii=False) + "\n",
            encoding="utf-8",
        )
        out.append(path)
    return out


def load_fixtures(directory: Path = FIXTURES_DIR) -> Fixtures:
    """Read the committed fixtures (``crossref.json`` and ``urls.json``)."""
    return Fixtures(
        crossref=CrossrefFixtures.model_validate_json(
            (directory / "crossref.json").read_text(encoding="utf-8")
        ),
        urls=UrlFixtures.model_validate_json(
            (directory / "urls.json").read_text(encoding="utf-8")
        ),
    )


# ---------------------------------------------------------------------------
# Statistics
# ---------------------------------------------------------------------------

StatKind = Literal["t", "F", "chi2", "r", "z"]
_STAT_KINDS: tuple[StatKind, ...] = ("t", "F", "chi2", "r", "z")


@dataclass(frozen=True)
class Stat:
    kind: StatKind
    dfs: tuple[int, ...]
    value: float  # rounded to two decimals, as printed
    p: float  # exact p for the printed value


def stat_p_value(kind: StatKind, dfs: tuple[int, ...], value: float) -> float:
    if kind == "t":
        return float(2 * sps.t.sf(abs(value), dfs[0]))
    if kind == "F":
        return float(sps.f.sf(value, dfs[0], dfs[1]))
    if kind == "chi2":
        return float(sps.chi2.sf(value, dfs[0]))
    if kind == "r":
        t = value * math.sqrt(dfs[0] / (1 - value * value))
        return float(2 * sps.t.sf(abs(t), dfs[0]))
    return float(2 * sps.norm.sf(abs(value)))


def _draw_stat(rng: random.Random, p_lo: float, p_hi: float) -> Stat:
    for _ in range(5000):
        kind = rng.choice(_STAT_KINDS)
        if kind == "t":
            dfs, v = (rng.randint(10, 150),), rng.uniform(0.5, 4.5)
        elif kind == "F":
            dfs, v = (rng.randint(1, 3), rng.randint(20, 150)), rng.uniform(0.5, 9.0)
        elif kind == "chi2":
            dfs, v = (rng.randint(1, 5),), rng.uniform(0.5, 20.0)
        elif kind == "r":
            dfs, v = (rng.randint(20, 200),), rng.uniform(0.05, 0.5)
        else:
            dfs, v = (), rng.uniform(0.5, 3.8)
        v = round(v, 2)
        if kind == "r" and not 0 < v < 1:
            continue
        p = stat_p_value(kind, dfs, v)
        if p_lo <= p <= p_hi:
            return Stat(kind, dfs, v, p)
    raise RuntimeError(f"no statistic found with p in [{p_lo}, {p_hi}]")


def _fmt_p(p: float) -> str:
    if p < 0.0005:
        return "p < .001"
    return "p = " + f"{p:.3f}".lstrip("0")


def format_stat(s: Stat, p_text: str) -> str:
    if s.kind == "t":
        head = f"t({s.dfs[0]}) = {s.value:.2f}"
    elif s.kind == "F":
        head = f"F({s.dfs[0]}, {s.dfs[1]}) = {s.value:.2f}"
    elif s.kind == "chi2":
        head = f"χ²({s.dfs[0]}) = {s.value:.2f}"
    elif s.kind == "r":
        head = f"r({s.dfs[0]}) = " + f"{s.value:.2f}".lstrip("0")
    else:
        head = f"z = {s.value:.2f}"
    return f"{head}, {p_text}"


# ---------------------------------------------------------------------------
# Topics
# ---------------------------------------------------------------------------


@dataclass(frozen=True)
class Topic:
    name: str
    intervention: str
    outcome: str
    unit: str
    population: str
    exposure: str


_TOPICS = [
    Topic(
        "SaltLog",
        "a salt-logging app",
        "systolic blood pressure",
        "mmHg",
        "adults in coastal towns",
        "sodium intake",
    ),
    Topic(
        "Apollo",
        "timed light exposure",
        "sleep latency",
        "min",
        "night-shift crews",
        "evening light",
    ),
    Topic(
        "Quillwork",
        "peer-led pacing",
        "walking cadence",
        "steps/min",
        "older community members",
        "daily pacing",
    ),
    Topic(
        "Marrowby",
        "a hydration reminder",
        "urine osmolality",
        "mOsm/kg",
        "orchard workers",
        "fluid intake",
    ),
    Topic(
        "Tern",
        "short breathing breaks",
        "resting heart rate",
        "bpm",
        "depot staff",
        "breaks at work",
    ),
    Topic(
        "Lumen",
        "screen dimming",
        "headache days",
        "days/month",
        "school teachers",
        "screen brightness",
    ),
    Topic(
        "Basalt",
        "standing desks",
        "sitting time",
        "min/day",
        "office clerks",
        "seated work",
    ),
    Topic(
        "Cinder",
        "cold-chain logging",
        "spoilage rate",
        "units/week",
        "pharmacy staff",
        "storage temperature",
    ),
    Topic(
        "Harrow",
        "a stretching routine",
        "lower-back pain score",
        "points",
        "river-port workers",
        "lifting load",
    ),
    Topic(
        "Pelican",
        "meal-timing advice",
        "fasting glucose",
        "mmol/L",
        "shift workers",
        "late meals",
    ),
    Topic(
        "Wren",
        "a quiet-hours policy",
        "noise annoyance score",
        "points",
        "ward nurses",
        "ambient noise",
    ),
    Topic(
        "Kestrel",
        "a step-count challenge",
        "daily steps",
        "steps/day",
        "library staff",
        "step goals",
    ),
    Topic(
        "Mistral",
        "ventilation upgrades",
        "cough days",
        "days/month",
        "classroom pupils",
        "indoor air changes",
    ),
    Topic(
        "Orrery",
        "a sleep diary",
        "sleep duration",
        "min/night",
        "university tutors",
        "bedtime regularity",
    ),
]

_CLAIMS = [
    "Earlier work linked {exposure} to {outcome}",
    "Observational data tie {exposure} to differences in {outcome}",
    "Interventions aimed at {exposure} have shown mixed effects on {outcome}",
    "Several groups have studied how {exposure} relates to {outcome}",
    "The link between {exposure} and {outcome} remains debated",
    "A methods comparison examined how {outcome} is best measured",
    "Reporting habits around {outcome} vary widely between sites",
    "Pilot programmes targeting {exposure} reported modest gains in {outcome}",
]
_FILLER = [
    "The programme was designed to be low cost and simple to run.",
    "Staff received a short briefing before the study began.",
    "Participants were free to withdraw at any time.",
    "Sites were chosen for variety in size and setting.",
    "All measurements followed a written protocol.",
]


# ---------------------------------------------------------------------------
# docx building blocks
# ---------------------------------------------------------------------------


@dataclass(frozen=True)
class Cite:
    ref: Ref
    malformed: Literal["truncated", "unbalanced"] | None = None


@dataclass(frozen=True)
class XRef:
    target: str  # "fig1", "tab1"
    broken: Literal["deleted", "error"] | None = None


@dataclass(frozen=True)
class Link:
    url: str


Segment = str | Cite | XRef | Link

_ZOTERO_SCHEMA = (
    "https://github.com/citation-style-language/schema/raw/master/csl-citation.json"
)


def _t(text: str) -> Any:
    el = OxmlElement("w:t")
    el.text = text
    el.set(qn("xml:space"), "preserve")
    return el


def _run(text: str | None = None, *, extra: Any = None) -> Any:
    r = OxmlElement("w:r")
    if extra is not None:
        r.append(extra)
    if text is not None:
        r.append(_t(text))
    return r


def _fld(kind: str) -> Any:
    el = OxmlElement("w:fldChar")
    el.set(qn("w:fldCharType"), kind)
    return el


def _instr(text: str) -> Any:
    el = OxmlElement("w:instrText")
    el.text = text
    el.set(qn("xml:space"), "preserve")
    return el


def _field_runs(instruction: str, display: str | None) -> list[Any]:
    runs = [_run(extra=_fld("begin")), _run(extra=_instr(instruction))]
    if display is not None:
        runs += [_run(extra=_fld("separate")), _run(display)]
    runs.append(_run(extra=_fld("end")))
    return runs


def _authors_display(ref: Ref) -> str:
    fam = [f for f, _ in ref.authors]
    if len(fam) == 1:
        return fam[0]
    if len(fam) == 2:
        return f"{fam[0]} & {fam[1]}"
    return f"{fam[0]} et al."


def _bib_entry(ref: Ref) -> str:
    shown = ref.authors[:3]
    names = ", ".join(f"{f}, {g}" for f, g in shown)
    if len(ref.authors) > 3:
        names += ", et al."
    return (
        f"{names} ({ref.year}). {ref.title}. {ref.journal}, {ref.volume}({ref.issue}), "
        f"{ref.pages}. https://doi.org/{ref.doi}"
    )


def _csl_item(ref: Ref, item_id: str) -> dict[str, object]:
    return {
        "id": item_id,
        "type": "article-journal",
        "title": ref.title,
        "container-title": ref.journal,
        "volume": ref.volume,
        "issue": ref.issue,
        "page": ref.pages,
        "DOI": ref.doi,
        "URL": f"https://doi.org/{ref.doi}",
        "author": _csl_authors(ref),
        "issued": {"date-parts": [[ref.year]]},
    }


def _zotero_json(ref: Ref, display: str, idx: int) -> str:
    uri = f"http://zotero.org/users/local/abc/items/BN{idx:05d}"
    payload = {
        "citationID": f"bn{idx:05d}",
        "properties": {
            "formattedCitation": display,
            "plainCitation": display,
            "noteIndex": 0,
        },
        "citationItems": [{"id": uri, "uris": [uri], "itemData": _csl_item(ref, uri)}],
        "schema": _ZOTERO_SCHEMA,
    }
    return json.dumps(payload, ensure_ascii=False, separators=(",", ":"))


def _mendeley_json(ref: Ref, display: str, idx: int) -> str:
    payload = {
        "citationItems": [
            {"id": f"ITEM-{idx}", "itemData": _csl_item(ref, f"ITEM-{idx}")}
        ],
        "mendeley": {
            "formattedCitation": display,
            "plainTextFormattedCitation": display,
            "previouslyFormattedCitation": display,
        },
        "properties": {"noteIndex": 0},
        "schema": _ZOTERO_SCHEMA,
    }
    return json.dumps(payload, ensure_ascii=False, separators=(",", ":"))


def _endnote_xml(ref: Ref, display: str, idx: int) -> str:
    e = xml_escape
    authors = "".join(f"<author>{e(f)}, {e(g)}</author>" for f, g in ref.authors)
    return (
        f"<EndNote><Cite><Author>{e(ref.authors[0][0])}</Author><Year>{ref.year}</Year>"
        f"<RecNum>{idx}</RecNum><DisplayText>{e(display)}</DisplayText><record>"
        f'<rec-number>{idx}</rec-number><ref-type name="Journal Article">17</ref-type>'
        f"<contributors><authors>{authors}</authors></contributors>"
        f"<titles><title>{e(ref.title)}</title><secondary-title>{e(ref.journal)}</secondary-title></titles>"
        f"<periodical><full-title>{e(ref.journal)}</full-title></periodical>"
        f"<pages>{e(ref.pages)}</pages><volume>{e(ref.volume)}</volume><number>{e(ref.issue)}</number>"
        f"<dates><year>{ref.year}</year></dates>"
        f"<electronic-resource-num>{e(ref.doi)}</electronic-resource-num></record></Cite></EndNote>"
    )


def _corrupt_json(text: str, how: Literal["truncated", "unbalanced"]) -> str:
    if how == "truncated":
        return text[: int(len(text) * 0.7)]
    return text[:-1]


# ---------------------------------------------------------------------------
# One manuscript
# ---------------------------------------------------------------------------


@dataclass
class FileSpec:
    name: str
    style: CitationStyle
    defects: tuple[str, ...]


@dataclass
class Counters:
    """Round-robin state shared across files so every variant occurs in a set."""

    url: int = 0
    stat: int = 0
    malformed: int = 0
    dangling: int = 0

    def copy(self) -> Counters:
        return Counters(self.url, self.stat, self.malformed, self.dangling)


_URL_ORDER = (
    "http_404",
    "dns_failure",
    "http_404",
    "http_410",
    "http_404",
    "timeout",
    "dns_failure",
)
_STAT_ORDER = ("flip_to_significant", "flip_to_nonsignificant", "magnitude_only")


@dataclass
class Rec:
    """One ground-truth record: a planted defect or a valid control item."""

    kind: str
    locations: list[tuple[int, str]]
    id: str | None = None
    expected_severity: str | None = None
    detail: str | None = None
    extra: dict[str, object] = field(default_factory=dict)

    def to_json(self, texts: list[str]) -> dict[str, object]:
        out: dict[str, object] = {}
        if self.id is not None:
            out["id"] = self.id
        out["kind"] = self.kind
        if self.expected_severity is not None:
            out["expected_severity"] = self.expected_severity
        out["locations"] = [
            {"paragraph_index": i, "snippet": snip, "paragraph_text": texts[i]}
            for i, snip in self.locations
        ]
        if self.detail is not None:
            out["detail"] = self.detail
        return {**out, **self.extra}


class Totals(BaseModel):
    files: int
    clean: int
    defects: int
    kinds: dict[str, int]


class ManifestEntry(BaseModel):
    docx: str
    truth: str
    sha256: str
    clean: bool
    defects: int
    kinds: dict[str, int]


class Manifest(BaseModel):
    schema_version: int
    seed: int
    files: list[ManifestEntry]
    totals: Totals


class _Manuscript:
    def __init__(
        self,
        spec: FileSpec,
        rng: random.Random,
        counters: Counters,
        used_dois: set[str],
    ) -> None:
        self.spec = spec
        self.rng = rng
        self.counters = counters
        self.doc = Document()
        self.cp = self.doc.core_properties
        self.cp.author = "Bench Generator"
        self.cp.last_modified_by = "Bench Generator"
        self.cp.created = self.cp.modified = _FIXED_TIME
        self.cp.title = ""
        self.paras: list[Any] = []
        self.bookmark_id = 0
        self.cite_count = 0
        self.cite_numbers: dict[str, int] = {}
        self.cite_locs: dict[str, list[tuple[int, str]]] = {}
        self.bib_idx: dict[str, int] = {}
        self.defects: list[Rec] = []
        self.valid: list[Rec] = []
        self.deleted_targets: set[str] = set()
        self.dois: set[str] = set()
        self.urls: set[str] = set()
        self.refs_order: list[Ref] = []
        self.abstract_idx = -1
        self.intro_idx = -1
        self.refs_heading_idx = -1
        self.pending_abstract: list[
            tuple[str, str, bool]
        ] = []  # (snippet, number text, is_defect)
        self.topic = rng.choice(_TOPICS)
        self.xref_display = ""

    # -- paragraph helpers ------------------------------------------------
    def para(self, style: str | None = None) -> tuple[int, Any]:
        p = self.doc.add_paragraph(style=style)
        self.paras.append(p)
        return len(self.paras) - 1, p

    def bookmark_name(self, target: str) -> str:
        number = int(re.sub(r"\D", "", target))
        return f"_Ref{410000000 + number + (0 if target.startswith('fig') else 100)}"

    def _cite_runs(self, c: Cite) -> tuple[list[Any], str]:
        ref = c.ref
        style = self.spec.style
        n = self.cite_numbers.setdefault(ref.key, len(self.cite_numbers) + 1)
        self.cite_count += 1
        if style == "mixed":
            style = "zotero" if (self.cite_count % 2 == 0 or c.malformed) else "endnote"
        display = (
            f"[{n}]"
            if style == "numbered"
            else f"({_authors_display(ref)}, {ref.year})"
        )
        if style == "numbered" or style == "author_year":
            return [_run(display)], display
        idx = self.cite_count
        if style == "zotero":
            body = _zotero_json(ref, display, idx)
            if c.malformed:
                body = _corrupt_json(body, c.malformed)
            return _field_runs(
                f" ADDIN ZOTERO_ITEM CSL_CITATION {body} ", display
            ), display
        if style == "mendeley":
            body = _mendeley_json(ref, display, idx)
            if c.malformed:
                body = _corrupt_json(body, c.malformed)
            return _field_runs(f" ADDIN CSL_CITATION {body} ", display), display
        return _field_runs(
            f" ADDIN EN.CITE {_endnote_xml(ref, display, idx)} ", display
        ), display

    def emit(self, idx: int, p: Any, segments: list[Segment]) -> None:
        for seg in segments:
            if isinstance(seg, str):
                p._p.append(_run(seg))
            elif isinstance(seg, Cite):
                runs, display = self._cite_runs(seg)
                for r in runs:
                    p._p.append(r)
                self.cite_locs.setdefault(seg.ref.key, []).append((idx, display))
                if seg.ref.doi:
                    self.dois.add(seg.ref.doi)
            elif isinstance(seg, XRef):
                if seg.broken == "error":
                    name, display = (
                        "_Ref999999001",
                        "Error! Reference source not found.",
                    )
                else:
                    name = self.bookmark_name(seg.target)
                    display = (
                        "Figure " if seg.target.startswith("fig") else "Table "
                    ) + re.sub(r"\D", "", seg.target)
                    if seg.broken == "deleted":
                        self.deleted_targets.add(seg.target)
                for r in _field_runs(f" REF {name} \\h ", display):
                    p._p.append(r)
                self.xref_display = display
            else:
                self.urls.add(seg.url)
                rid = self.doc.part.relate_to(seg.url, RT.HYPERLINK, is_external=True)
                h = OxmlElement("w:hyperlink")
                h.set(qn("r:id"), rid)
                h.set(qn("w:history"), "1")
                rpr = OxmlElement("w:rPr")
                color = OxmlElement("w:color")
                color.set(qn("w:val"), "0563C1")
                under = OxmlElement("w:u")
                under.set(qn("w:val"), "single")
                rpr.append(color)
                rpr.append(under)
                r = _run(seg.url)
                r.insert(0, rpr)
                h.append(r)
                p._p.append(h)

    def text_para(self, segments: list[Segment], style: str | None = None) -> int:
        idx, p = self.para(style)
        self.emit(idx, p, segments)
        return idx

    def caption(self, target: str, label: str, number: int, text: str) -> int:
        idx, p = self.para("Caption")
        name = self.bookmark_name(target)
        has_bookmark = target not in self.deleted_targets
        if has_bookmark:
            self.bookmark_id += 1
            start = OxmlElement("w:bookmarkStart")
            start.set(qn("w:id"), str(self.bookmark_id))
            start.set(qn("w:name"), name)
            p._p.append(start)
        p._p.append(_run(f"{label} "))
        for r in _field_runs(f" SEQ {label} \\* ARABIC ", str(number)):
            p._p.append(r)
        if has_bookmark:
            end = OxmlElement("w:bookmarkEnd")
            end.set(qn("w:id"), str(self.bookmark_id))
            p._p.append(end)
        p._p.append(_run(f". {text}"))
        return idx

    # -- truth bookkeeping ------------------------------------------------
    def add_defect(
        self,
        kind: str,
        *,
        severity: str,
        locations: list[tuple[int, str]],
        detail: str,
        **extra: object,
    ) -> None:
        self.defects.append(
            Rec(
                kind,
                locations,
                id=f"{self.spec.name.split('_')[0]}-{len(self.defects) + 1:02d}",
                expected_severity=severity,
                detail=detail,
                extra=dict(extra),
            )
        )

    def add_valid(
        self, kind: str, locations: list[tuple[int, str]], **extra: object
    ) -> None:
        self.valid.append(Rec(kind, locations, extra=dict(extra)))

    # -- build ------------------------------------------------------------
    def build(self) -> None:
        rng, spec, pl = self.rng, self.spec, pools()
        want = {k: spec.defects.count(k) for k in KINDS}
        topic = self.topic
        n = rng.randint(40, 400)
        pct = round(rng.uniform(4.0, 19.0), 1)
        diff = round(rng.uniform(1.2, 9.8), 1)
        n_figs = 3

        # references to cite
        base = rng.sample(pl.valid, 6 + want["malformed_citation_json"])
        fab = rng.sample(pl.fabricated, want["fabricated_doi"])
        ret = rng.sample(pl.retracted, want["retracted_doi"])
        malformed_refs = {r.key for r in base[: want["malformed_citation_json"]]}
        cited = [*base, *fab, *ret]
        rng.shuffle(cited)
        malformed_how: dict[str, Literal["truncated", "unbalanced"]] = {}
        for key in sorted(malformed_refs):
            malformed_how[key] = ("truncated", "unbalanced")[
                self.counters.malformed % 2
            ]
            self.counters.malformed += 1

        # number of abstract defects: replace pct, then diff, then n
        abstract_defects = want["abstract_number_missing"]
        alts = {
            "pct": f"{round(pct + rng.choice([-1, 1]) * rng.uniform(6.0, 9.0), 1):.1f}",
            "diff": f"{round(diff + rng.uniform(3.5, 6.0), 1):.1f}",
            "n": str(n + rng.randint(101, 350)),
        }
        replaced = ["pct", "diff", "n"][:abstract_defects]
        a_pct = alts["pct"] if "pct" in replaced else f"{pct:.1f}"
        a_diff = alts["diff"] if "diff" in replaced else f"{diff:.1f}"
        a_n = alts["n"] if "n" in replaced else str(n)

        title = f"{topic.name} pilot: effect of {topic.intervention} on {topic.outcome} in {topic.population}"
        self.cp.title = title
        self.para("Title")[1].add_run(title)
        self.para("Heading 1")[1].add_run("Abstract")
        self.abstract_idx = self.text_para(
            [
                f"Background: {topic.name} asked whether {topic.intervention} changes {topic.outcome} among "
                f"{topic.population}. Methods: we enrolled {a_n} participants and followed them for twelve weeks. "
                f"Results: {topic.outcome} fell by {a_pct}% in the intervention arm, {a_diff} {topic.unit} "
                f"lower than control. Conclusions: the approach is feasible and merits a larger trial."
            ]
        )
        self.pending_abstract = [
            (f"{a_pct}%", a_pct, "pct" in replaced),
            (f"{a_diff} {topic.unit}", a_diff, "diff" in replaced),
            (f"{a_n} participants", a_n, "n" in replaced),
        ]

        # introduction
        self.para("Heading 1")[1].add_run("Introduction")
        self.intro_idx = len(self.paras)
        sentences: list[list[Segment]] = []
        claims = list(_CLAIMS)
        rng.shuffle(claims)
        for i, ref in enumerate(cited):
            claim = claims[i % len(claims)].format(
                exposure=topic.exposure, outcome=topic.outcome
            )
            sentences.append(
                [
                    claim + " ",
                    Cite(
                        ref,
                        malformed_how.get(ref.key)
                        if ref.key in malformed_refs
                        else None,
                    ),
                    ".",
                ]
            )
        split = (len(sentences) + 1) // 2
        chunks: list[list[Segment]] = []
        for chunk in (sentences[:split], sentences[split:]):
            segs: list[Segment] = []
            for sent in chunk:
                segs += [*sent, " "]
            segs.append(rng.choice(_FILLER))
            chunks.append(segs)
        self.text_para(chunks[0])
        discussion_segs = chunks[1]
        # methods
        self.para("Heading 1")[1].add_run("Methods")
        self.text_para(
            [
                f"We ran a randomised pilot among {topic.population}. {rng.choice(_FILLER)} "
                f"The primary outcome was {topic.outcome}, measured in {topic.unit}."
            ]
        )
        ok_urls = rng.sample(pl.urls_ok, 2)
        segs = [
            "Analysis code is archived at ",
            Link(ok_urls[0].url),
            ". De-identified data are available at ",
            ok_urls[1].url,
            ".",
        ]
        idx = self.text_para(segs)
        for e in ok_urls:
            self.add_valid("valid_url", [(idx, e.url)], url=e.url, variant=e.variant)
        for _ in range(want["broken_url"]):
            variant = _URL_ORDER[self.counters.url % len(_URL_ORDER)]
            self.counters.url += 1
            cand = [
                u
                for u in pl.urls_broken
                if u.variant == variant and u.url not in self.urls
            ]
            entry = rng.choice(cand)
            as_link = rng.random() < 0.5
            idx = self.text_para(
                [
                    "A supplementary protocol is posted at ",
                    Link(entry.url) if as_link else entry.url,
                    ".",
                ]
            )
            self.add_defect(
                "broken_url",
                severity="warning",
                locations=[(idx, entry.url)],
                detail=f"URL does not resolve ({variant}).",
                url=entry.url,
                variant=variant,
                hyperlink=as_link,
            )

        # results
        self.para("Heading 1")[1].add_run("Results")
        self.text_para(
            [
                f"Of {n} participants randomised, {topic.outcome} fell by {pct:.1f}% in the intervention arm, "
                f"a difference of {diff:.1f} {topic.unit} relative to control."
            ]
        )
        # statistics
        n_valid_stats = 2
        stat_idx: list[tuple[str, Stat, str | None]] = []
        for _ in range(n_valid_stats):
            lo, hi = rng.choice([(0.0001, 0.044), (0.06, 0.9)])
            stat_idx.append(("valid", _draw_stat(rng, lo, hi), None))
        for _ in range(want["inconsistent_statistic"]):
            variant = _STAT_ORDER[self.counters.stat % len(_STAT_ORDER)]
            self.counters.stat += 1
            if variant == "flip_to_significant":
                stat_idx.append(("defect", _draw_stat(rng, 0.06, 0.19), variant))
            elif variant == "flip_to_nonsignificant":
                stat_idx.append(("defect", _draw_stat(rng, 0.021, 0.044), variant))
            else:
                stat_idx.append(("defect", _draw_stat(rng, 0.012, 0.044), variant))
        rng.shuffle(stat_idx)
        templates = [
            "The arms differed on {outcome} ({stat}).",
            "A secondary analysis of adherence gave {stat}.",
            "Change from baseline was compared across arms, {stat}.",
            "Site effects were small, {stat}.",
            "The pre-specified subgroup contrast gave {stat}.",
        ]
        tpl = list(templates)
        rng.shuffle(tpl)
        for i, (role, s, variant) in enumerate(stat_idx):
            if role == "valid":
                text = format_stat(s, _fmt_p(s.p))
            elif variant == "flip_to_significant":
                text = format_stat(s, _fmt_p(round(s.p / 4, 3)))
            elif variant == "flip_to_nonsignificant":
                text = format_stat(s, _fmt_p(round(s.p * 3, 3)))
            else:
                text = format_stat(s, _fmt_p(round(s.p / 3, 3)))
            sentence = tpl[i % len(tpl)].format(outcome=topic.outcome, stat=text)
            idx = self.text_para([sentence])
            if role == "valid":
                self.add_valid(
                    "valid_statistic", [(idx, text)], recomputed_p=round(s.p, 6)
                )
            else:
                flips = variant != "magnitude_only"
                self.add_defect(
                    "inconsistent_statistic",
                    severity="error" if flips else "warning",
                    locations=[(idx, text)],
                    detail="Reported p disagrees with the p recomputed from the reported statistic"
                    + (
                        " and the decision at .05 flips."
                        if flips
                        else " by a large factor."
                    ),
                    variant=variant,
                    recomputed_p=round(s.p, 6),
                    decision_flips=flips,
                )

        # cross-reference mentions (valid + dangling)
        targets = ["fig1", "fig2", "fig3", "tab1"]
        rng.shuffle(targets)
        broken_targets = targets[: want["dangling_crossref"]]
        for t in broken_targets:
            variant = ("deleted", "error")[self.counters.dangling % 2]
            self.counters.dangling += 1
            word = "Figure" if t.startswith("fig") else "Table"
            x = XRef(t, broken=variant)  # type: ignore[arg-type]
            idx = self.text_para(
                ["Trends over follow-up are shown in (see ", x, ")."]
                if t.startswith("fig")
                else ["Baseline characteristics are summarised in ", x, "."]
            )
            display = self.xref_display
            self.add_defect(
                "dangling_crossref",
                severity="error",
                locations=[(idx, display)],
                detail=f"{word} cross-reference points to a bookmark that does not exist.",
                variant=variant,
                target=t,
            )
        for t in targets[len(broken_targets) :]:
            x = XRef(t)
            word = "Figure" if t.startswith("fig") else "Table"
            idx = self.text_para(["Further detail appears in ", x, "."])
            self.add_valid("valid_crossref", [(idx, self.xref_display)], target=t)
        # plain mentions that exist (controls)
        idx = self.text_para(
            ["Panel layouts follow Figure 1 and the table follows Table 1."]
        )
        self.add_valid("valid_figure_mention", [(idx, "Figure 1")])
        # missing figures
        numbers = rng.sample([4, 5, 6, 7], want["missing_figure"])
        for num in numbers:
            idx = self.text_para([f"Dose-response curves are plotted in Figure {num}."])
            self.add_defect(
                "missing_figure",
                severity="error",
                locations=[(idx, f"Figure {num}")],
                detail=f"Text mentions Figure {num}; the document has only Figures 1-{n_figs}.",
                number=num,
            )
        # captions
        for i in range(1, n_figs + 1):
            self.caption(
                f"fig{i}",
                "Figure",
                i,
                rng.choice(
                    [
                        f"{topic.outcome.capitalize()} by study arm.",
                        f"Adherence over follow-up in {topic.population}.",
                        "Participant flow through the pilot.",
                        f"Distribution of {topic.outcome} at baseline.",
                    ]
                ),
            )
        self.caption(
            "tab1", "Table", 1, f"Baseline characteristics of {topic.population}."
        )
        self.text_para(
            [
                "Values are mean (standard deviation). Table body omitted in this synthetic file."
            ]
        )

        # discussion
        self.para("Heading 1")[1].add_run("Discussion")
        self.text_para(discussion_segs)
        self.text_para(
            [
                f"This synthetic {topic.name} study is invented for benchmarking and carries no findings."
            ]
        )

        # references
        self.refs_heading_idx = self.para("Heading 1")[0]
        self.paras[-1].add_run("References")
        refs = list(self.cite_numbers)
        by_key = {r.key: r for r in cited}
        ordered = (
            sorted(refs, key=lambda k: self.cite_numbers[k])
            if spec.style == "numbered"
            else sorted(refs, key=lambda k: (by_key[k].authors[0][0], by_key[k].year))
        )
        self.refs_order = [by_key[k] for k in ordered]
        bib_open: tuple[str, str] | None = {
            "zotero": (
                ' ADDIN ZOTERO_BIBL {"uncited":[],"omitted":[],"custom":[]} CSL_BIBLIOGRAPHY ',
                "",
            ),
            "mixed": (
                ' ADDIN ZOTERO_BIBL {"uncited":[],"omitted":[],"custom":[]} CSL_BIBLIOGRAPHY ',
                "",
            ),
            "mendeley": (" ADDIN Mendeley Bibliography CSL_BIBLIOGRAPHY ", ""),
            "endnote": (" ADDIN EN.REFLIST ", ""),
        }.get(spec.style)
        for i, ref in enumerate(self.refs_order):
            idx, p = self.para()
            if bib_open and i == 0:
                for r in _field_runs(bib_open[0], None)[:2]:
                    p._p.append(r)
                p._p.append(_run(extra=_fld("separate")))
            prefix = (
                f"[{self.cite_numbers[ref.key]}] " if spec.style == "numbered" else ""
            )
            p._p.append(_run(prefix + _bib_entry(ref)))
            if bib_open and i == len(self.refs_order) - 1:
                p._p.append(_run(extra=_fld("end")))
            self.bib_idx[ref.key] = idx
            self.dois.add(ref.doi)

        # truth for citations
        for ref in self.refs_order:
            locs = [*self.cite_locs[ref.key], (self.bib_idx[ref.key], ref.doi)]
            if ref.status == "fabricated":
                self.add_defect(
                    "fabricated_doi",
                    severity="warning",
                    locations=locs,
                    detail="DOI does not exist in Crossref (404) and no title match.",
                    doi=ref.doi,
                    title=ref.title,
                )
            elif ref.status == "retracted":
                self.add_defect(
                    "retracted_doi",
                    severity="error",
                    locations=locs,
                    detail="Cited work has been retracted (Crossref update-to type retraction).",
                    doi=ref.doi,
                    notice_doi=ref.notice_doi,
                )
            elif ref.key in malformed_refs:
                self.add_defect(
                    "malformed_citation_json",
                    severity="error",
                    locations=self.cite_locs[ref.key],
                    detail="Citation field JSON does not parse.",
                    variant=malformed_how[ref.key],
                    ref_key=ref.key,
                )
                self.add_valid(
                    "valid_doi", [(self.bib_idx[ref.key], ref.doi)], doi=ref.doi
                )
            else:
                self.add_valid("valid_doi", locs, doi=ref.doi)
        # abstract truth
        for snippet, num, bad in self.pending_abstract:
            if bad:
                self.add_defect(
                    "abstract_number_missing",
                    severity="error",
                    locations=[(self.abstract_idx, snippet)],
                    detail=f"Abstract states {snippet}; the number does not occur in the body.",
                    number=num,
                )
            else:
                self.add_valid(
                    "valid_abstract_number", [(self.abstract_idx, snippet)], number=num
                )

    # -- finish -----------------------------------------------------------
    def display_texts(self) -> list[str]:
        return ["".join(t.text or "" for t in p._p.iter(qn("w:t"))) for p in self.paras]

    def validate(self) -> None:
        """Raise if a planted defect is not exactly what its record says."""
        texts = self.display_texts()
        for rec in [*self.defects, *self.valid]:
            for index, snippet in rec.locations:
                if snippet not in texts[index]:
                    raise _Retry(f"snippet {snippet!r} missing from paragraph {index}")
        body = "\n".join(texts[self.intro_idx : self.refs_heading_idx])

        def present(num: str) -> bool:
            return re.search(rf"(?<![\d.]){re.escape(num)}(?!\d)", body) is not None

        for rec in self.defects:
            if rec.kind == "abstract_number_missing" and present(
                str(rec.extra["number"])
            ):
                raise _Retry("abstract defect number occurs in body")
        for rec in self.valid:
            if rec.kind == "valid_abstract_number" and not present(
                str(rec.extra["number"])
            ):
                raise _Retry("valid abstract number absent from body")
        for t in self.deleted_targets:
            if self.bookmark_name(t) in self._bookmarks():
                raise _Retry("deleted bookmark still present")

    def _bookmarks(self) -> set[str]:
        return {
            b.get(qn("w:name"))
            for p in self.paras
            for b in p._p.iter(qn("w:bookmarkStart"))
        }

    def truth(self, seed: int) -> dict[str, object]:
        texts = self.display_texts()
        return {
            "schema_version": SCHEMA_VERSION,
            "file": f"{self.spec.name}.docx",
            "seed": seed,
            "clean": not self.defects,
            "citation_style": self.spec.style,
            "topic": self.topic.name,
            "paragraph_count": len(self.paras),
            "abstract_paragraph_index": self.abstract_idx,
            "dois": sorted(self.dois),
            "urls": sorted(self.urls),
            "defects": [r.to_json(texts) for r in self.defects],
            "valid_items": [r.to_json(texts) for r in self.valid],
        }


class _Retry(Exception):
    pass


# ---------------------------------------------------------------------------
# The set
# ---------------------------------------------------------------------------


def _plan() -> list[FileSpec]:
    D = lambda *k: tuple(k)  # noqa: E731
    F, R, U, X, M, A, S, G = KINDS
    return [
        FileSpec("c01_clean", "zotero", ()),
        FileSpec("d01_mixed_a", "zotero", D(F, F, U, X, A)),
        FileSpec("d02_mixed_b", "zotero", D(R, M, S, G)),
        FileSpec("d03_endnote", "endnote", D(F, R, U, U, S)),
        FileSpec("c02_clean", "endnote", ()),
        FileSpec("d04_numbered", "numbered", D(X, X, A, A, S)),
        FileSpec("d05_zotero", "zotero", D(M, M, F, G, U)),
        FileSpec("c03_clean", "numbered", ()),
        FileSpec("d06_mendeley", "mendeley", D(R, S, S, A, X)),
        FileSpec("d07_zotero", "zotero", D(G, G, U, F, M)),
        FileSpec("d08_author_year", "author_year", D(F, R, S, A, X)),
        FileSpec("c04_clean", "mendeley", ()),
        FileSpec("d09_zotero", "zotero", D(S, S, U, M, G)),
        FileSpec("d10_mixed", "mixed", D(R, F, X, A, U, G)),
    ]


def _normalise_zip(path: Path) -> None:
    with zipfile.ZipFile(path) as zin:
        items = [(i.filename, zin.read(i.filename)) for i in zin.infolist()]
    tmp = path.with_suffix(".tmp")
    with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for name, data in items:
            zi = zipfile.ZipInfo(name, date_time=(2026, 1, 1, 0, 0, 0))
            zi.compress_type = zipfile.ZIP_DEFLATED
            zi.external_attr = 0o600 << 16
            zout.writestr(zi, data)
    tmp.replace(path)


def generate(out: Path, *, seed: int) -> Manifest:
    """Write the manuscripts, truth files and ``manifest.json`` into ``out``."""
    out.mkdir(parents=True, exist_ok=True)
    counters = Counters()
    entries: list[ManifestEntry] = []
    for spec in _plan():
        for attempt in range(_MAX_ATTEMPTS):
            trial = counters.copy()
            ms = _Manuscript(
                spec, random.Random(f"{seed}/{spec.name}/{attempt}"), trial, set()
            )
            ms.build()
            try:
                ms.validate()
            except _Retry:
                continue
            counters = trial
            break
        else:
            raise RuntimeError(
                f"could not build a consistent {spec.name} in {_MAX_ATTEMPTS} attempts"
            )
        docx_path = out / f"{spec.name}.docx"
        ms.doc.save(str(docx_path))
        _normalise_zip(docx_path)
        truth_path = out / f"{spec.name}.truth.json"
        truth_path.write_text(
            json.dumps(ms.truth(seed), indent=2, ensure_ascii=False) + "\n",
            encoding="utf-8",
        )
        entries.append(
            ManifestEntry(
                docx=docx_path.name,
                truth=truth_path.name,
                sha256=hashlib.sha256(docx_path.read_bytes()).hexdigest(),
                clean=not ms.defects,
                defects=len(ms.defects),
                kinds={k: sum(1 for d in ms.defects if d.kind == k) for k in KINDS},
            )
        )
    manifest = Manifest(
        schema_version=SCHEMA_VERSION,
        seed=seed,
        files=entries,
        totals=Totals(
            files=len(entries),
            clean=sum(1 for e in entries if e.clean),
            defects=sum(e.defects for e in entries),
            kinds={k: sum(e.kinds[k] for e in entries) for k in KINDS},
        ),
    )
    (out / "manifest.json").write_text(
        manifest.model_dump_json(indent=2) + "\n", encoding="utf-8"
    )
    return manifest


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(
        prog="python -m andamentum.compose.harness.bench.planted",
        description="Generate synthetic manuscripts with planted defects.",
    )
    ap.add_argument(
        "--out", type=Path, help="output directory for .docx and truth files"
    )
    ap.add_argument("--seed", type=int, default=0, help="seed (same seed, same bytes)")
    ap.add_argument(
        "--write-fixtures",
        type=Path,
        metavar="DIR",
        help="(re)write the recorded Crossref/URL fixtures into DIR",
    )
    args = ap.parse_args(argv)
    if args.out is None and args.write_fixtures is None:
        ap.error("give --out and/or --write-fixtures")
    if args.write_fixtures is not None:
        for path in write_fixtures(args.write_fixtures):
            print(path)
    if args.out is not None:
        m = generate(args.out, seed=args.seed)
        t = m.totals
        print(f"{t.files} files ({t.clean} clean), {t.defects} defects -> {args.out}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
