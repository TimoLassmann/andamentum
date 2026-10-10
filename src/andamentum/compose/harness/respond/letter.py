"""The response letter: one entry per reviewer point, as .md and .docx."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from pydantic import BaseModel, ConfigDict

LETTER_STEM = "response-letter"


class LetterEntry(BaseModel):
    model_config = ConfigDict(extra="forbid")

    label: str
    point: str  # the reviewer's point, quoted
    response: str
    location: str  # where the tracked change or the evidence is; "" when none


def render_markdown(entries: list[LetterEntry], *, unanswered: list[str]) -> str:
    lines = ["# Response to the reviewers", ""]
    for entry in entries:
        lines += [f"## {entry.label}", ""]
        lines += [f"> {line}" if line else ">" for line in entry.point.splitlines()]
        lines += ["", f"**Response.** {entry.response}", ""]
        if entry.location:
            lines += [f"**Location.** {entry.location}", ""]
    if unanswered:
        lines += ["## Points that could not be read reliably", ""]
        lines += [f"- {label}" for label in unanswered]
        lines.append("")
    return "\n".join(lines)


def write_docx(
    path: Path, entries: list[LetterEntry], *, unanswered: list[str]
) -> None:
    document = Document()
    document.add_heading("Response to the reviewers", level=1)
    for entry in entries:
        document.add_heading(entry.label, level=2)
        document.add_paragraph(entry.point, style="Quote")
        response = document.add_paragraph()
        response.add_run("Response. ").bold = True
        response.add_run(entry.response)
        if entry.location:
            location = document.add_paragraph()
            location.add_run("Location. ").bold = True
            location.add_run(entry.location)
    if unanswered:
        document.add_heading("Points that could not be read reliably", level=2)
        for label in unanswered:
            document.add_paragraph(label, style="List Bullet")
    document.save(str(path))


def default_letter(folder: Path) -> Path:
    """The letter's .docx when the caller names none: ``response-letter.docx`` in *folder*."""
    return folder / f"{LETTER_STEM}.docx"


def letter_paths(letter: Path) -> tuple[Path, Path]:
    """The letter's files: the .docx *letter* and the .md beside it."""
    return letter, letter.with_suffix(".md")


def existing_letters(letter: Path) -> list[Path]:
    """The letter files that already exist; they may hold the authors' edits."""
    return [path for path in letter_paths(letter) if path.exists()]


def write_letter(
    letter: Path, entries: list[LetterEntry], *, unanswered: list[str]
) -> tuple[Path, Path]:
    """Write the letter: the .docx *letter* and the .md beside it.

    Raises ``FileExistsError`` when either exists: the letter is edited by hand
    between rounds and is never overwritten.
    """
    found = existing_letters(letter)
    if found:
        raise FileExistsError(", ".join(str(path) for path in found))
    letter.parent.mkdir(parents=True, exist_ok=True)
    docx_path, md_path = letter_paths(letter)
    write_docx(docx_path, entries, unanswered=unanswered)
    md_path.write_text(
        render_markdown(entries, unanswered=unanswered), encoding="utf-8"
    )
    return docx_path, md_path
