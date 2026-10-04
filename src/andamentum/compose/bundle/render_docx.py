"""Render a bundle to a python-docx document.

The title is a Word ``Title`` paragraph and every section title is exactly
``Heading 1`` — whetstone sectionises a .docx by its heading styles, and
compose maps whetstone's sections back to bundle sections by title. Lists use
``List Bullet`` / ``List Number``, tables ``Table Grid``, image captions
``Caption``. A missing image gets a placeholder paragraph (``validate`` has
already reported the file).
"""

from __future__ import annotations

from pathlib import Path
from typing import Callable

from docx import Document as DocxDocument
from docx.document import Document as DocxDoc
from docx.shared import Inches

from .markdown import ImageBlock, ListBlock, ParagraphBlock, TableBlock, parse_blocks
from .models import Reference, Section
from .parser import find_ai_markers, inline_runs

IMAGE_WIDTH_IN = 5.5
_AI_MARKER_KEYWORD = "andamentum:contains-ai-markers"


def _stamp_ai_marker_keyword(out: DocxDoc) -> None:
    """Tag the docx's core ``keywords`` with an AI-provenance keyword.

    Visible in Word's File-Info pane and recoverable by any docx-reading tool,
    so AI assistance shows without parsing the body text.
    """
    existing = (out.core_properties.keywords or "").strip()
    if _AI_MARKER_KEYWORD in existing:
        return
    out.core_properties.keywords = (
        f"{existing}, {_AI_MARKER_KEYWORD}" if existing else _AI_MARKER_KEYWORD
    )


def _add_runs(paragraph, text: str, *, bold: bool = False) -> None:
    for chunk, styles in inline_runs(text):
        run = paragraph.add_run(chunk)
        if bold or "bold" in styles:
            run.bold = True
        if "italic" in styles:
            run.italic = True
        if "code" in styles:
            run.font.name = "Courier New"


def _emit_caption(out: DocxDoc, caption: str) -> None:
    paragraph = out.add_paragraph(style="Caption")
    _add_runs(paragraph, caption)


def _emit_body(out: DocxDoc, body: str, *, root: Path) -> None:
    for block in parse_blocks(body):
        if isinstance(block, ParagraphBlock):
            _add_runs(out.add_paragraph(), block.text)
        elif isinstance(block, ListBlock):
            style = "List Number" if block.ordered else "List Bullet"
            for item in block.items:
                _add_runs(out.add_paragraph(style=style), item)
        elif isinstance(block, TableBlock):
            table = out.add_table(rows=len(block.rows), cols=len(block.rows[0]))
            table.style = "Table Grid"
            for r, row in enumerate(block.rows):
                for c, text in enumerate(row):
                    _add_runs(table.rows[r].cells[c].paragraphs[0], text, bold=(r == 0))
        elif isinstance(block, ImageBlock):
            image = Path(block.path)
            resolved = image if image.is_absolute() else root / image
            if resolved.is_file():
                out.add_picture(str(resolved), width=Inches(IMAGE_WIDTH_IN))
            else:
                out.add_paragraph(f"[Missing figure: {block.path}]")
            if block.caption:
                _emit_caption(out, block.caption)


def render_to_docx(
    *,
    title: str,
    sections: list[Section],
    references: list[Reference],
    root: Path,
    output_path: Path,
    text_filter: Callable[[str], str] | None,
) -> None:
    """Write the bundle's content to *output_path* as a .docx file."""
    out = DocxDocument()
    out.core_properties.title = title
    out.add_heading(title, level=0)
    for section in sections:
        out.add_heading(section.title, level=1)
        body = text_filter(section.body) if text_filter else section.body
        if find_ai_markers(body):
            _stamp_ai_marker_keyword(out)
        _emit_body(out, body, root=root)
    if references:
        out.add_heading("References", level=1)
        for ref in references:
            line = f"[{ref.key}] {ref.text}" + (f" {ref.url}" if ref.url else "")
            out.add_paragraph(line)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    out.save(str(output_path))
