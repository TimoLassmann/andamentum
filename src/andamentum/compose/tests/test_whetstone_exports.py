from __future__ import annotations

import pytest
from docx import Document

import andamentum.whetstone as w
from andamentum.whetstone._confidentiality import ConfidentialityMarkerError


def test_exported():
    assert "check_confidentiality" in w.__all__
    assert "stamp_provenance" in w.__all__


def test_check_confidentiality_clean_and_marked():
    w.check_confidentiality("A plain note about the Apollo Labs pilot.")
    w.check_confidentiality("")
    with pytest.raises(ConfidentialityMarkerError):
        w.check_confidentiality(
            "This is a confidential peer review of grant applications."
        )


def test_stamp_provenance(tmp_path):
    p = tmp_path / "d.docx"
    d = Document()
    d.add_paragraph("hello")
    d.save(p)
    w.stamp_provenance(p, model="test:model")
    props = Document(p).core_properties
    assert "test:model" in (props.comments or "") + (props.keywords or "") + (
        props.subject or ""
    ) + (props.author or "") + (props.last_modified_by or "") + (props.title or "")


def test_reviewer_profile_and_provenance_line():
    assert "reviewer_profile" in w.__all__ and "provenance_line" in w.__all__
    assert w.reviewer_profile("reader").name == "reader"
    with pytest.raises(ValueError, match="unknown reviewer profile"):
        w.reviewer_profile("nobody")
    assert "test:model" in w.provenance_line(model="test:model")
