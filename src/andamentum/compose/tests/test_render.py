from __future__ import annotations

from typing import Any

from pathlib import Path

import pytest
from docx import Document

from andamentum.compose.genres import genre_for
from andamentum.compose.render import (
    ProvenanceRequired,
    citation_filter,
    marker_ids,
    plain_text,
    render_outputs,
    source_numbers,
    sources_body,
    strip_markers,
)
from andamentum.compose.schemas import Brief, Fact, Source
from andamentum.scribe import Bundle, Section

FACTS = [
    Fact(id="f1", source_id="a", statement="s", quote="q", start=0, end=1),
    Fact(id="f2", source_id="b", statement="s", quote="q", start=0, end=1),
    Fact(id="f3", source_id="a", statement="s", quote="q", start=0, end=1),
]
SOURCES = [
    Source(
        id="a",
        path="/home/jane/notes/a.md",
        label="a.md",
        title="Apollo notes",
        sha256="x",
        chars=1,
        kind="file",
    ),
    Source(
        id="b",
        path="https://example.org/b",
        label="https://example.org/b",
        title="Saltlog page",
        sha256="y",
        chars=1,
        kind="url",
    ),
]


def brief(**kw) -> Brief:
    base: dict[str, Any] = dict(
        task="t",
        genre="memo",
        title="Apollo",
        audience="board",
        purpose="p",
        key_messages=["k"],
        must_include=[],
        constraints=[],
        word_range=(1, 999),
        reader_questions=[],
    )
    base.update(kw)
    return Brief(**base)


def make_bundle(path: Path, parts: list[tuple[str, str]]) -> Bundle:
    bundle = Bundle.create(path, title="Apollo")
    for title, body in parts:
        s = bundle.add_section(title=title)
        bundle.write_section(
            s.id, body=body, expected_revision=s.revision, status="drafted"
        )
    return bundle


def test_strip_markers():
    assert (
        strip_markers("Sales rose [[f1]]. Costs fell [[f2]][[f3]].")
        == "Sales rose. Costs fell."
    )
    assert strip_markers("No markers here.") == "No markers here."
    assert marker_ids("a [[f1]] b [[f2]] c [[f1]]") == ["f1", "f2", "f1"]


def test_citation_numbers_by_first_appearance_and_collapse():
    sections = [
        Section(
            id="s01",
            title="A",
            order=10,
            revision=1,
            status="drafted",
            body="One [[f2]]. Two [[f1]][[f3]].",
        ),
    ]
    numbers = source_numbers(sections, facts=FACTS)
    assert numbers == {"b": 1, "a": 2}
    out = citation_filter(numbers, facts=FACTS)(
        "One [[f2]]. Two [[f1]][[f3]]. Gone [[f9]]."
    )
    assert out == "One [1]. Two [2]. Gone."
    body = sources_body(numbers, sources=SOURCES)
    assert body.splitlines() == [
        "1. Saltlog page (https://example.org/b)",
        "2. Apollo notes (a.md)",
    ]


def test_plain_text_has_titles_and_no_markers():
    bundle_sections = [
        Section(
            id="s01",
            title="Summary",
            order=10,
            revision=1,
            status="drafted",
            body="Fact [[f1]].",
        )
    ]
    text = plain_text(bundle_sections, title="Apollo")
    assert "Summary" in text and "[[" not in text


def test_cites_genre_docx_has_citations_and_sources(tmp_path):
    bundle = make_bundle(
        tmp_path / "b.compose",
        [
            ("Executive summary", "Sales rose [[f1]]. Costs fell [[f2]]."),
        ],
    )
    out = render_outputs(
        bundle,
        brief=brief(),
        genre=genre_for("report"),
        facts=FACTS,
        sources=SOURCES,
        model="test:m",
        formats=["docx", "md"],
    )
    assert [p.name for p in out] == ["document.docx", "document.md"]
    text = "\n".join(p.text for p in Document(str(out[0])).paragraphs)
    assert "Sales rose [1]." in text and "Costs fell [2]." in text
    assert "Sources" in text and "Apollo notes" in text
    md = out[1].read_text()
    assert "[[" not in md and "Sources" in md
    assert "andamentum-whetstone" in md  # provenance footer
    # the real bundle gains no Sources section
    assert [s.title for s in bundle.sections()] == ["Executive summary"]


def test_non_cites_genre_removes_markers_and_has_no_sources(tmp_path):
    bundle = make_bundle(tmp_path / "b.compose", [("Summary", "Plan holds [[f1]].")])
    out = render_outputs(
        bundle,
        brief=brief(),
        genre=genre_for("memo"),
        facts=FACTS,
        sources=SOURCES,
        model="test:m",
        formats=["md"],
    )
    md = out[0].read_text()
    assert "Plan holds." in md and "[[" not in md and "Sources" not in md


def test_email_txt_format(tmp_path):
    bundle = make_bundle(
        tmp_path / "b.compose",
        [("Body", "Hi Jane Doe, please send it by Friday [[f1]].")],
    )
    out = render_outputs(
        bundle,
        brief=brief(genre="email", email_subject="Apollo budget"),
        genre=genre_for("email"),
        facts=FACTS,
        sources=SOURCES,
        model="test:m",
        formats=["txt"],
    )
    assert out[0].name == "email.txt"
    text = out[0].read_text()
    assert text.startswith(
        "Subject: Apollo budget\n\nHi Jane Doe, please send it by Friday."
    )
    assert "andamentum-whetstone" in text


def test_docx_provenance_stamped(tmp_path):
    bundle = make_bundle(tmp_path / "b.compose", [("Summary", "Plan holds.")])
    out = render_outputs(
        bundle,
        brief=brief(),
        genre=genre_for("memo"),
        facts=[],
        sources=[],
        model="test:model",
        formats=["docx"],
    )
    props = Document(str(out[0])).core_properties
    joined = " ".join(
        str(getattr(props, a) or "")
        for a in (
            "comments",
            "keywords",
            "subject",
            "author",
            "last_modified_by",
            "title",
        )
    )
    assert "test:model" in joined


def test_provenance_refused_for_grant(tmp_path):
    bundle = make_bundle(tmp_path / "b.compose", [("Summary", "Plan holds.")])
    with pytest.raises(ProvenanceRequired):
        render_outputs(
            bundle,
            brief=brief(genre="grant"),
            genre=genre_for("grant"),
            facts=[],
            sources=[],
            model="m",
            provenance=False,
        )


def test_provenance_can_be_off_for_memo(tmp_path):
    bundle = make_bundle(tmp_path / "b.compose", [("Summary", "Plan holds.")])
    out = render_outputs(
        bundle,
        brief=brief(),
        genre=genre_for("memo"),
        facts=[],
        sources=[],
        model="m",
        formats=["md"],
        provenance=False,
    )
    assert "andamentum-whetstone" not in out[0].read_text()


def test_unknown_format_raises(tmp_path):
    bundle = make_bundle(tmp_path / "b.compose", [("Summary", "x.")])
    with pytest.raises(ValueError, match="unknown output format"):
        render_outputs(
            bundle,
            brief=brief(),
            genre=genre_for("memo"),
            facts=[],
            sources=[],
            model="m",
            formats=["pdf"],
        )


def _bundle_with_figure(path: Path) -> Bundle:
    from andamentum.scribe.tests._png import PNG_BYTES

    bundle = make_bundle(
        path,
        [("Executive summary", "Sales rose [[f1]].\n\n![Sales](figures/sales.png)")],
    )
    (bundle.path / "figures").mkdir(exist_ok=True)
    (bundle.path / "figures" / "sales.png").write_bytes(PNG_BYTES)
    return bundle


def test_figure_links_in_md_and_html_resolve_after_render(tmp_path):
    import re

    bundle = _bundle_with_figure(tmp_path / "b.compose")
    out = render_outputs(
        bundle,
        brief=brief(),
        genre=genre_for("report"),
        facts=FACTS,
        sources=SOURCES,
        model="test:m",
        formats=["md", "html"],
    )
    md, html = (p.read_text(encoding="utf-8") for p in out)
    md_link = re.search(r"!\[Sales\]\(([^)]+)\)", md)
    html_link = re.search(r'src="([^"]*sales\.png)"', html)
    assert md_link and html_link
    for link in (md_link.group(1), html_link.group(1)):
        assert (out[0].parent / link).resolve().is_file(), link


def test_no_absolute_path_reaches_the_document(tmp_path):
    bundle = make_bundle(
        tmp_path / "b.compose", [("Summary", "Up [[f1]]. Down [[f2]].")]
    )
    out = render_outputs(
        bundle,
        brief=brief(),
        genre=genre_for("report"),
        facts=FACTS,
        sources=SOURCES,
        model="test:m",
        formats=["md"],
    )
    md = out[0].read_text(encoding="utf-8")
    assert "/home/jane" not in md and "Apollo notes (a.md)" in md
