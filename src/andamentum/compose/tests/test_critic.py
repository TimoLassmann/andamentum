from __future__ import annotations

from typing import Literal

import json
from pathlib import Path

from docx import Document

from andamentum.compose.checks import critic
from andamentum.compose.checks.critic import map_findings, round_dir, run_critic
from andamentum.compose.genres import genre_for
from andamentum.compose.log import ActionLog
from andamentum.scribe import Bundle, Section
from andamentum.whetstone import Finding, Quote, ReviewResult


def finding(
    severity: Literal["minor", "moderate", "major"],
    *,
    sections: list[str],
    title: str = "Weak claim",
    quotes: list[Quote] | None = None,
) -> Finding:
    return Finding(
        title=title,
        severity=severity,
        confidence="high",
        rationale="why",
        fix="Do this.",
        sections_involved=sections,
        quotes=quotes or [],
    )


def sec(i: int, title: str) -> Section:
    return Section(
        id=f"s{i:02d}",
        title=title,
        order=i * 10,
        revision=1,
        status="drafted",
        body="x",
    )


SECTIONS = [sec(1, "Summary"), sec(2, "Options or Analysis")]
TITLES = {"s1": "Summary", "s2": "options or analysis"}


def test_only_major_and_moderate_become_items():
    items = map_findings(
        [
            finding("major", sections=["s1"]),
            finding("moderate", sections=["s2"]),
            finding("minor", sections=["s1"]),
        ],
        whetstone_titles=TITLES,
        sections=SECTIONS,
    )
    assert [(i.level, i.section_id) for i in items] == [
        ("section", "s01"),
        ("section", "s02"),
    ]
    assert all(i.source == "critic" for i in items)
    assert "Do this." in items[0].reason


def test_finding_without_section_is_outline_level():
    items = map_findings(
        [finding("major", sections=[])], whetstone_titles=TITLES, sections=SECTIONS
    )
    assert [(i.level, i.section_id) for i in items] == [("outline", None)]


def test_unmapped_whetstone_section_is_outline_level():
    items = map_findings(
        [finding("major", sections=["s9"])], whetstone_titles=TITLES, sections=SECTIONS
    )
    assert items[0].level == "outline"


def test_finding_in_two_sections_gives_two_items():
    items = map_findings(
        [finding("major", sections=["s1", "s2"])],
        whetstone_titles=TITLES,
        sections=SECTIONS,
    )
    assert [i.section_id for i in items] == ["s01", "s02"]


def test_quote_offsets_locate_the_bundle_section_inside_one_whetstone_section():
    text = "Apollo\nSummary\nPlan holds.\nOptions or Analysis\nTwo options exist.\n"
    at = text.index("Two options")
    f = finding(
        "major",
        sections=["s1"],
        quotes=[
            Quote(section_id="s1", char_start=at, char_end=at + 11, text="Two options")
        ],
    )
    items = map_findings(
        [f], whetstone_titles={"s1": "Apollo"}, sections=SECTIONS, text=text
    )
    assert [(i.level, i.section_id) for i in items] == [("section", "s02")]


def make_bundle(path: Path) -> Bundle:
    bundle = Bundle.create(path, title="Apollo")
    for title, body in [
        ("Summary", "Plan holds [[f1]]."),
        ("Recommendation", "Do it."),
    ]:
        s = bundle.add_section(title=title)
        bundle.write_section(
            s.id, body=body, expected_revision=s.revision, status="drafted"
        )
    return bundle


async def test_run_critic_review_mode(tmp_path, monkeypatch):
    bundle = make_bundle(tmp_path / "b.compose")
    seen: dict = {}

    async def fake_review(source, **kw):
        seen.update(kw)
        seen["titles"] = [s.title for s in source.sections]
        seen["docx_text"] = "\n".join(
            p.text for p in Document(source.docx_path).paragraphs
        )
        first = source.sections[0]
        at = source.text.index("Plan holds")
        q = Quote(
            section_id=first.id, char_start=at, char_end=at + 10, text="Plan holds"
        )
        return ReviewResult(
            findings=[
                finding("major", sections=[first.id], quotes=[q]),
                finding("minor", sections=[first.id], quotes=[q]),
            ]
        )

    monkeypatch.setattr(critic, "review_document", fake_review)
    log = ActionLog(bundle.path)
    res = await run_critic(
        bundle, genre=genre_for("memo"), model="c:m", round_no=2, log=log
    )

    assert seen["model"] == "c:m"
    assert seen["document_type"] == "external_communication"
    assert seen["confirm_own_draft"] is True
    assert seen["editor"] is False and seen["flag_novelty"] is False
    assert seen["check_novelty"] is False
    assert seen["reviewer"] is not None and seen["reviewer"].name == "reader"
    assert "[[" not in seen["docx_text"]
    assert (round_dir(bundle, round_no=2) / "draft.docx").exists()
    saved = json.loads((round_dir(bundle, round_no=2) / "review.json").read_text())
    assert len(saved["findings"]) == 2
    assert res.check.passed is False and res.major == 1
    assert [(i.level, i.section_id) for i in res.items] == [("section", "s01")]
    assert log.records()[0]["role"] == "critic"


async def test_run_critic_checklist_mode_no_major_passes(tmp_path, monkeypatch):
    bundle = make_bundle(tmp_path / "b.compose")
    seen: dict = {}

    async def fake_checklist(source, **kw):
        seen.update(kw)
        at = source.text.index("Do it")
        q = Quote(
            section_id=source.sections[0].id,
            char_start=at,
            char_end=at + 5,
            text="Do it",
        )
        return ReviewResult(findings=[finding("moderate", sections=[], quotes=[q])])

    async def not_called(*a, **k):
        raise AssertionError("review_document must not run in checklist mode")

    monkeypatch.setattr(critic, "checklist_document", fake_checklist)
    monkeypatch.setattr(critic, "review_document", not_called)
    res = await run_critic(
        bundle,
        genre=genre_for("email"),
        model="c:m",
        round_no=1,
        log=ActionLog(bundle.path),
    )
    assert seen["document_type"] == "general" and seen["reviewer"] is None
    assert res.check.passed is True and res.moderate == 1 and res.major == 0
    assert res.items[0].section_id == "s02"


async def test_critic_reviews_the_document_as_it_ships(tmp_path, monkeypatch):
    from andamentum.compose.schemas import Fact, Source

    bundle = make_bundle(tmp_path / "b.compose")
    seen: dict = {}

    async def fake_review(source, **kw):
        seen["text"] = "\n".join(p.text for p in Document(source.docx_path).paragraphs)
        at = source.text.index("Sources")
        q = Quote(
            section_id=source.sections[0].id,
            char_start=at,
            char_end=at + 7,
            text="Sources",
        )
        return ReviewResult(findings=[finding("major", sections=[], quotes=[q])])

    monkeypatch.setattr(critic, "review_document", fake_review)
    res = await run_critic(
        bundle,
        genre=genre_for("report"),
        model="c:m",
        round_no=1,
        log=ActionLog(bundle.path),
        facts=[Fact(id="f1", source_id="a", statement="s", quote="q", start=0, end=1)],
        sources=[
            Source(
                id="a",
                path="/x/a.md",
                label="a.md",
                title="Apollo notes",
                sha256="x",
                chars=1,
                kind="file",
            )
        ],
    )
    assert "Plan holds [1]." in seen["text"] and "Apollo notes (a.md)" in seen["text"]
    # a finding about the Sources list has no bundle section: outline level
    assert [(i.level, i.section_id) for i in res.items] == [("outline", None)]
    assert [s.title for s in bundle.sections()] == ["Summary", "Recommendation"]
