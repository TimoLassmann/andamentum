"""Sectionising a Word document by its own heading styles.

The contract is the same one ``test_sectionize.py`` pins for markdown —
``section.text == source[section.start:section.end]`` — except that the source
is the atom projection, so a section offset is a document address.
"""

from __future__ import annotations

from docx import Document
from lxml import etree  # type: ignore

from andamentum.docx_edit.atoms import document_atoms, projection, qn
from andamentum.whetstone.docx.reading import read_docx
from andamentum.docx_edit.textmap import TextMap
from andamentum.whetstone.pipeline.docx_source import load_docx, sectionize_docx


def _write(tmp_path, build) -> str:
    doc = Document()
    build(doc)
    path = tmp_path / "draft.docx"
    doc.save(path)
    return path


def test_headings_at_several_levels_become_sections(tmp_path) -> None:
    def build(doc):
        doc.add_heading("Introduction", 1)
        doc.add_paragraph("We study X.")
        doc.add_heading("Methods", 1)
        doc.add_paragraph("We did Y.")
        doc.add_heading("Assay", 2)
        doc.add_paragraph("Cells were counted.")

    # A budget below the whole document forces the banding to emit each
    # heading's own piece rather than one root section.
    result = load_docx(_write(tmp_path, build), target_max=30)

    assert [s.title for s in result.sections] == ["Introduction", "Methods", "Assay"]
    assert [s.id for s in result.sections] == ["s1", "s2", "s3"]
    assert result.sections[0].text.startswith("Introduction\n")


def test_a_title_styled_paragraph_opens_the_document(tmp_path) -> None:
    def build(doc):
        doc.add_heading("A Study of X", 0)  # python-docx level 0 → the Title style
        doc.add_paragraph("Author line.")
        doc.add_heading("Introduction", 1)
        doc.add_paragraph("Body.")

    result = load_docx(_write(tmp_path, build), target_max=30)

    # Title sits above Heading1, so the author line belongs to the title's
    # section instead of becoming an anonymous preamble.
    assert result.sections[0].title == "A Study of X"
    assert "Author line." in result.sections[0].text
    assert not any(s.title == "(preamble)" for s in result.sections)


def test_section_offsets_index_the_projection_exactly(tmp_path) -> None:
    def build(doc):
        doc.add_heading("Introduction", 1)
        doc.add_paragraph("We study X in a cohort of 42 mice.")
        doc.add_heading("Methods", 1)
        doc.add_paragraph("We did Y, then Z.")

    result = load_docx(_write(tmp_path, build), target_max=25)

    assert result.sections
    for section in result.sections:
        assert section.text == result.text[section.start : section.end]
    # Non-overlapping and in document order.
    ends = [s.end for s in result.sections]
    starts = [s.start for s in result.sections]
    assert starts == sorted(starts)
    assert all(end <= start for end, start in zip(ends, starts[1:]))


def test_a_document_with_no_styled_headings_yields_one_section(tmp_path) -> None:
    def build(doc):
        doc.add_paragraph("Dear editor,")
        doc.add_paragraph("Please find enclosed our manuscript.")

    result = load_docx(_write(tmp_path, build))

    assert len(result.sections) == 1
    assert result.sections[0].title == "(document)"
    assert result.sections[0].text == result.text


def test_a_heading_with_inline_formatting_reads_as_one_title(tmp_path) -> None:
    """Word splits a heading with italics into several runs; the title is whole."""

    def build(doc):
        heading = doc.add_heading("", 1)
        heading.add_run("The ")
        heading.add_run("BRCA1").italic = True
        heading.add_run(" pathway")
        doc.add_paragraph("Body.")

    result = load_docx(_write(tmp_path, build))
    assert result.sections[0].title == "The BRCA1 pathway"


def test_empty_paragraphs_between_sections_do_not_shift_offsets(tmp_path) -> None:
    def build(doc):
        doc.add_heading("Introduction", 1)
        doc.add_paragraph("We study X.")
        doc.add_paragraph("")
        doc.add_paragraph("")
        doc.add_heading("Methods", 1)
        doc.add_paragraph("We did Y.")

    result = load_docx(_write(tmp_path, build), target_max=25)
    methods = next(s for s in result.sections if s.title == "Methods")

    assert result.text[methods.start : methods.start + 7] == "Methods"
    assert methods.text == result.text[methods.start : methods.end]
    # The blank paragraphs stay in the section they follow, as two newlines.
    assert result.sections[0].text.endswith("We study X.\n\n\n")


def test_a_quote_from_a_section_is_already_a_document_address(tmp_path) -> None:
    """The payoff: reviewer coordinates ARE document coordinates.

    Converted markdown put ``aim 3,4 .`` where Word holds ``aim3,4.``,
    and 15 of 43 rewrites on a real manuscript could not be placed back. Here
    the quote is a slice of the document's own projection, so the span is found
    without a matcher having to bridge two renderings.
    """

    def build(doc):
        doc.add_heading("Methods", 1)
        paragraph = doc.add_paragraph("Mice were assigned to ")
        paragraph.add_run("aim").bold = True
        paragraph.add_run("3,4. Then we waited.")

    path = _write(tmp_path, build)
    result = load_docx(path)
    section = result.sections[0]

    start = section.text.index("aim3,4.")
    quote = section.text[start : start + len("aim3,4.")]

    textmap = TextMap(document_atoms(read_docx(path).paragraphs))
    span = textmap.find(quote)
    assert span is not None
    assert (span.start, span.end) == (
        section.start + start,
        section.start + start + len(quote),
    )
    assert textmap.slice(span) == quote


def test_sectionize_docx_works_without_a_styles_part() -> None:
    """Direct ``w:outlineLvl`` formatting is enough on its own."""
    paragraphs = []
    for text, outline in (("Results", 0), ("We found Z.", None), ("Discussion", 0)):
        p = Document().add_paragraph()._p
        if outline is not None:
            p_pr = etree.SubElement(p, qn("pPr"))
            etree.SubElement(p_pr, qn("outlineLvl")).set(qn("val"), str(outline))
            p.insert(0, p_pr)
        run = etree.SubElement(p, qn("r"))
        etree.SubElement(run, qn("t")).text = text
        paragraphs.append(p)

    result = sectionize_docx(paragraphs, styles_root=None, target_max=10)

    assert [s.title for s in result.sections] == ["Results", "Discussion"]
    assert result.text == projection(document_atoms(paragraphs))
