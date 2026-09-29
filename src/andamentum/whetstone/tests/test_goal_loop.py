"""End-to-end tests for the goal-driven loop, in the style of
``test_sequential_rounds.py``: a real ``.docx`` on disk, every LLM call
stubbed, and the round trip exercised for real.

The editor agent (``v3_editor``) is called exactly once per round for this
test's single-section document (see ``v3/editor.py:run_editor`` — it calls
every section regardless of whether findings were routed to it), so a
counter advanced on that call is a reliable "which round is this" signal for
the rest of the stub.
"""

from __future__ import annotations

import types
from contextlib import ExitStack
from pathlib import Path
from unittest.mock import patch

from docx import Document

from andamentum.whetstone.goal import GoalSpec, run_goal_loop
from andamentum.whetstone.reviewed_document import open_source
from andamentum.whetstone.v3.consolidate import _Consolidation
from andamentum.whetstone.v3.editor import EditorOutput, EditProposal
from andamentum.whetstone.v3.extract import _ClaimSpans, _Requote
from andamentum.whetstone.v3.gaps import _DemandList
from andamentum.whetstone.v3.review import _CriterionFindings, _RawFinding
from andamentum.whetstone.v3.synth import StructuredReview

ORIGINAL_SENTENCE = "The evidence layer earns its place in the analysis."
SECOND_SENTENCE = "We sampled twelve participants across three sites."


def _source_docx(path: Path) -> Path:
    document = Document()
    document.add_heading("Methods", level=1)
    document.add_paragraph(ORIGINAL_SENTENCE)
    document.add_paragraph(SECOND_SENTENCE)
    document.save(str(path))
    return path


def _stub_agents(
    *, rounds: list[dict], editor_prompts: list[str] | None = None
) -> ExitStack:
    """One entry in *rounds* per review round: ``findings`` (a list of
    ``(severity, issue)`` pairs, possibly empty) and ``edit`` (an
    ``(original, replacement)`` pair, or ``None`` for "the editor proposes
    nothing this round"). *editor_prompts*, when given, collects every prompt
    the editor is sent.
    """
    state = {"round": 0}

    def _current() -> dict:
        return rounds[min(state["round"], len(rounds) - 1)]

    def _router(defn, _model):
        name = defn.name

        class _Agent:
            def output_validator(self, fn):
                return fn

            async def run(self, _prompt, **_kwargs):
                spec = _current()
                if name == "v3_extract_claims":
                    out = _ClaimSpans(claims=[SECOND_SENTENCE])
                elif name == "v3_requote":
                    out = _Requote(quote="")
                elif name == "v3_gap_analysis":
                    out = _DemandList(demands=[])
                elif name == "v3_consolidate":
                    out = _Consolidation(groups=[])
                elif name in ("v3_synthesise", "v3_critique_revise"):
                    out = StructuredReview(
                        synopsis="A short methods note.",
                        strengths=["clear"],
                        weaknesses=["thin"],
                    )
                elif name == "v3_editor":
                    if editor_prompts is not None:
                        editor_prompts.append(_prompt)
                    edit = spec["edit"]
                    if edit is None:
                        out = EditorOutput(edits=[])
                    else:
                        original, replacement = edit
                        out = EditorOutput(
                            edits=[
                                EditProposal(
                                    title="Sharpen the idiom",
                                    rationale="clarity",
                                    original_text=original,
                                    new_text=replacement,
                                )
                            ]
                        )
                    # Editor is the last call in a round for a one-section
                    # document — advance the round counter here so the NEXT
                    # review_document_v3 call reads the next round's spec.
                    state["round"] += 1
                else:
                    out = _CriterionFindings(findings=[])
                return types.SimpleNamespace(output=out)

        return _Agent()

    def _review_agent(_criterion, _agent_model):
        class _Agent:
            def output_validator(self, fn):
                return fn

            async def run(self, _prompt, **_kwargs):
                spec = _current()
                findings = [
                    _RawFinding(issue=issue, quote=SECOND_SENTENCE, severity=severity)
                    for severity, issue in spec["findings"]
                ]
                return types.SimpleNamespace(
                    output=_CriterionFindings(findings=findings)
                )

        return _Agent()

    stack = ExitStack()
    for module in ("extract", "gaps", "consolidate", "synth", "editor"):
        stack.enter_context(
            patch(
                f"andamentum.whetstone.v3.{module}.build_pydantic_ai_agent", new=_router
            )
        )
        stack.enter_context(
            patch(f"andamentum.whetstone.v3.{module}.resolve_model", new=lambda x: None)
        )
    stack.enter_context(
        patch("andamentum.whetstone.v3.review._build_agent", new=_review_agent)
    )
    return stack


async def _open(tmp_path: Path):
    source = _source_docx(tmp_path / "draft.docx")
    return await open_source(str(source))


async def test_goal_met_on_two_consecutive_satisfied_rounds(tmp_path: Path) -> None:
    """Round 1 fails, rounds 2 and 3 pass — the loop stops after 3 rounds
    and does NOT apply round 3's own edit."""
    document = await _open(tmp_path)
    rounds = [
        {
            "findings": [("major", "too vague")],
            "edit": ("earns its place", "earns its keep"),
        },
        {"findings": [], "edit": ("earns its keep", "earns its salt")},
        {"findings": [], "edit": ("earns its salt", "earns its worth")},
    ]
    with _stub_agents(rounds=rounds):
        result = await run_goal_loop(
            document,
            GoalSpec(text="the flow is correct", check="structure", max_rounds=5),
            model="stub",
            document_type="academic",
        )

    assert result.goal_met is True
    assert result.stopped_reason == "goal_met"
    assert len(result.rounds) == 3
    # Round 2's edit landed; round 3's own edit — the satisfying one — did not.
    assert "earns its salt" in result.final_document.text
    assert "earns its worth" not in result.final_document.text
    assert result.final_review.author_questions == []


async def test_no_progress_when_the_editor_proposes_nothing(tmp_path: Path) -> None:
    document = await _open(tmp_path)
    rounds = [{"findings": [("major", "too vague")], "edit": None}]
    with _stub_agents(rounds=rounds):
        result = await run_goal_loop(
            document,
            GoalSpec(text="the flow is correct", check="structure"),
            model="stub",
            document_type="academic",
        )

    assert result.goal_met is False
    assert result.stopped_reason == "no_progress"
    assert len(result.rounds) == 1
    assert "nothing left to change" in result.reason


async def test_no_progress_on_two_back_oscillation(tmp_path: Path) -> None:
    """Round 2's edit reverses round 1's, landing back on round 0's text."""
    document = await _open(tmp_path)
    rounds = [
        {"findings": [("major", "x")], "edit": ("earns its place", "earns its keep")},
        {"findings": [("major", "x")], "edit": ("earns its keep", "earns its place")},
    ]
    with _stub_agents(rounds=rounds):
        result = await run_goal_loop(
            document,
            GoalSpec(text="the flow is correct", check="structure", max_rounds=5),
            model="stub",
            document_type="academic",
        )

    assert result.goal_met is False
    assert result.stopped_reason == "no_progress"
    assert "alternating" in result.reason
    assert len(result.rounds) == 2


async def test_cap_reached_with_the_goal_still_unmet(tmp_path: Path) -> None:
    document = await _open(tmp_path)
    rounds = [
        {"findings": [("major", "x")], "edit": ("earns its place", "earns its keep")},
        {"findings": [("major", "x")], "edit": ("earns its keep", "earns its salt")},
    ]
    with _stub_agents(rounds=rounds):
        result = await run_goal_loop(
            document,
            GoalSpec(text="the flow is correct", check="structure", max_rounds=2),
            model="stub",
            document_type="academic",
        )

    assert result.goal_met is False
    assert result.stopped_reason == "cap_reached"
    assert len(result.rounds) == 2
    assert "earns its salt" in result.final_document.text


async def test_citations_check_with_no_reference_section_is_unresolvable(
    tmp_path: Path,
) -> None:
    """No stubbing needed: the check must never reach a review call."""
    document = await _open(tmp_path)
    result = await run_goal_loop(
        document,
        GoalSpec(text="every reference is consistent", check="citations"),
        model="stub",
        document_type="academic",
    )

    assert result.goal_met is False
    assert result.stopped_reason == "unresolvable"
    assert result.rounds == []
    assert result.reason.startswith("No reference list")


async def test_a_non_met_stop_always_carries_an_author_question(tmp_path: Path) -> None:
    document = await _open(tmp_path)
    rounds = [{"findings": [("major", "too vague")], "edit": None}]
    with _stub_agents(rounds=rounds):
        result = await run_goal_loop(
            document,
            GoalSpec(text="the flow is correct", check="structure"),
            model="stub",
            document_type="academic",
        )

    assert len(result.final_review.author_questions) == 1
    question = result.final_review.author_questions[0]
    assert "the flow is correct" in question.question
    assert question.why == result.reason


async def test_judged_check_calls_the_judge_separately_every_round(
    tmp_path: Path,
) -> None:
    """The arbiter is a separate call every round, not the reviewer/editor
    pass that produced the edit — decision #3 in the design."""
    from typing import Literal

    from andamentum.llm_judge import CriterionScore, ScoreResult

    document = await _open(tmp_path)
    rounds = [
        {"findings": [], "edit": ("earns its place", "earns its keep")},
        {"findings": [], "edit": ("earns its keep", "earns its salt")},
        {"findings": [], "edit": ("earns its salt", "earns its worth")},
    ]
    verdicts: list[Literal["meets", "partial", "fails"]] = ["fails", "meets", "meets"]
    calls: list[str] = []

    async def _fake_judge_score(output, *, criteria, context=None, model):
        calls.append(model)
        overall = verdicts[len(calls) - 1]
        return ScoreResult(
            per_criterion=[
                CriterionScore(
                    criterion=criteria[0].name,
                    reasoning="stubbed",
                    meets=100 if overall == "meets" else 0,
                    partial=0,
                    fails=0 if overall == "meets" else 100,
                )
            ],
            overall=overall,
            confidence=1.0,
            doubt=0.0,
            needs_review=False,
            expected_score=1.0 if overall == "meets" else 0.0,
        )

    with (
        _stub_agents(rounds=rounds),
        patch("andamentum.llm_judge.judge_score", new=_fake_judge_score),
    ):
        result = await run_goal_loop(
            document,
            GoalSpec(text="the introduction motivates the study", check="judged"),
            model="stub",
            judge_model="judge-model",
            document_type="academic",
        )

    assert result.goal_met is True
    assert calls == ["judge-model", "judge-model", "judge-model"]


# ── --goal-check citations: the reference checker, with its problems edited ──

CITING_SENTENCE = "Costs are high [2]."


def _citing_docx(path: Path) -> Path:
    document = Document()
    document.add_heading("Methods", level=1)
    document.add_paragraph(
        f"Rare diseases are common [1]. {CITING_SENTENCE} Care is slow [3]."
    )
    document.add_heading("References", level=1)
    document.add_paragraph(
        "1. Smith J (2020) A study of rare disease costs. J Test 1:1"
    )
    document.add_paragraph("3. Lee K (2021) Delays in rare disease care. J Test 2:2")
    document.save(str(path))
    return path


async def test_reference_problems_reach_the_editor_and_the_goal_holds_once_fixed(
    tmp_path: Path,
) -> None:
    """Round 1: the checker finds [2] missing from the list, and the editor —
    shown that problem beside the review's own findings — rewrites the
    citation. Rounds 2 and 3 find nothing, so the goal is confirmed twice.

    The editor runs once per SECTION, and this document has two (Methods and
    References), so every round takes two entries of *rounds*. The rewrite
    offered to the References section does not anchor there and is dropped.
    """
    document = await open_source(str(_citing_docx(tmp_path / "draft.docx")))
    fix = {"findings": [], "edit": (CITING_SENTENCE, "Costs are high [3].")}
    quiet = {"findings": [], "edit": None}
    prompts: list[str] = []
    with _stub_agents(
        rounds=[fix, fix, quiet, quiet, quiet, quiet], editor_prompts=prompts
    ):
        result = await run_goal_loop(
            document,
            GoalSpec(
                text="every reference is consistent", check="citations", max_rounds=4
            ),
            model="stub",
            document_type="academic",
        )

    assert result.goal_met is True, result.reason
    assert len(result.rounds) == 3
    first = result.rounds[0]
    assert first.verdict.satisfied is False
    assert "Citation 2 is not in the list" in first.verdict.reason
    # The problem is a finding of the round, located at the citation itself.
    [problem] = [
        f for f in first.review.findings if f.category == "references:cited_not_listed"
    ]
    assert problem.quotes[0].text == "[2]"
    assert any("no entry numbered 2" in p for p in prompts[:2])
    assert "Costs are high [3]." in result.final_document.text
    assert all(r.verdict.satisfied for r in result.rounds[1:])


async def test_a_document_whose_citation_style_cannot_be_told_is_unresolvable(
    tmp_path: Path, monkeypatch
) -> None:
    """With no numbers, keys, notes or author-year citations to go on, the
    checker cannot decide the style — and says so before any round runs."""
    document = Document()
    document.add_heading("Methods", level=1)
    document.add_paragraph(
        "Rare diseases are common, as Smith showed in A Study of Costs."
    )
    document.add_heading("Works Cited", level=1)
    document.add_paragraph("Smith, John. A Study of Costs. Test Press.")
    path = tmp_path / "mla.docx"
    document.save(str(path))

    async def no_model(*_args, **_kwargs):
        raise AssertionError("the style is not the model's to decide here")

    monkeypatch.setattr("andamentum.references.llm.ask_style", no_model)
    result = await run_goal_loop(
        await open_source(str(path)),
        GoalSpec(text="every reference is consistent", check="citations"),
        model="stub",
        document_type="academic",
    )
    assert result.stopped_reason == "unresolvable"
    assert result.rounds == []
