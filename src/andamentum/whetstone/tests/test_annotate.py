"""Placing annotations by their own coordinates.

The three shapes here are the three the converted-markdown route measurably
could not handle, all on one real manuscript: a superscript citation run welded
to the word before it, emphasis splitting a sentence across runs, and a
sentence that occurs twice. Each one is a *matching* problem, and each one
stops existing once the reviewer's offsets are addresses in the document — so
these tests assert exact placement, not a high hit rate.
"""

from __future__ import annotations

import zipfile
from pathlib import Path

import pytest
from docx import Document
from lxml import etree  # type: ignore

from andamentum.whetstone.docx.annotate import (
    CoordinateError,
    Note,
    Rewrite,
    apply_annotations,
)
from andamentum.whetstone.docx.atoms import document_atoms, projection
from andamentum.whetstone.docx.low_level import DocxEditor
from andamentum.whetstone.docx.reading import read_docx

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def _qn(tag: str) -> str:
    return f"{{{W}}}{tag}"


# ── fixtures: documents whose text the converter route mangles ──────────


def _superscript_doc(path: Path) -> str:
    """``aim3,4.`` — the citation is its own run, with no space before it."""
    document = Document()
    document.add_heading("Methods", level=1)
    paragraph = document.add_paragraph("We met the ")
    paragraph.add_run("aim")
    paragraph.add_run("3,4").font.superscript = True
    paragraph.add_run(". Then we stopped.")
    document.save(str(path))
    return _projection_of(path)


def _emphasis_doc(path: Path) -> str:
    """One sentence, three runs, because one word is bold."""
    document = Document()
    paragraph = document.add_paragraph("The effect was ")
    paragraph.add_run("large").bold = True
    paragraph.add_run(" and clear in every cohort.")
    document.save(str(path))
    return _projection_of(path)


def _repeated_doc(path: Path) -> str:
    """The same sentence in two paragraphs. A search can only refuse."""
    document = Document()
    document.add_paragraph("The effect was large and clear.")
    document.add_paragraph("We then repeated the analysis.")
    document.add_paragraph("The effect was large and clear.")
    document.save(str(path))
    return _projection_of(path)


def _projection_of(path: Path) -> str:
    return projection(document_atoms(read_docx(path).paragraphs))


# ── readback helpers ────────────────────────────────────────────────────


def _comment_spans(path: Path) -> dict[str, str]:
    """The text each comment id brackets, read back out of the saved file."""
    with zipfile.ZipFile(path) as archive:
        root = etree.fromstring(archive.read("word/document.xml"))
    active: set[str] = set()
    covered: dict[str, str] = {}
    for element in root.iter():
        tag = element.tag
        if tag == _qn("commentRangeStart"):
            active.add(element.get(_qn("id")))
            covered.setdefault(element.get(_qn("id")), "")
        elif tag == _qn("commentRangeEnd"):
            active.discard(element.get(_qn("id")))
        elif tag == _qn("t"):
            for cid in active:
                covered[cid] = covered[cid] + (element.text or "")
    return covered


def _comment_paragraph_indexes(path: Path) -> dict[str, int]:
    """Which body paragraph each comment range STARTS in."""
    with zipfile.ZipFile(path) as archive:
        root = etree.fromstring(archive.read("word/document.xml"))
    out: dict[str, int] = {}
    for index, paragraph in enumerate(root.iter(_qn("p"))):
        for marker in paragraph.iter(_qn("commentRangeStart")):
            out[marker.get(_qn("id"))] = index
    return out


def _write(editor: DocxEditor, path: Path) -> Path:
    editor.write(str(path))
    return path


# ── the three shapes ────────────────────────────────────────────────────


def test_a_rewrite_spanning_a_superscript_citation_lands_exactly(
    tmp_path: Path,
) -> None:
    """``aim3,4.`` is three runs and no spaces. Docling renders it
    ``aim 3,4 .``, which is why the search missed; the offsets do not
    care how anything renders."""
    source = tmp_path / "src.docx"
    text = _superscript_doc(source)
    assert "aim3,4." in text

    start = text.index("aim3,4.")
    editor = DocxEditor(str(source))
    report = apply_annotations(
        editor,
        rewrites=[
            Rewrite(
                start=start,
                end=start + len("aim3,4."),
                new_text="aim 3 and 4.",
                author="t",
            )
        ],
        notes=[],
        reviewed_text=text,
    )
    assert report.applied_rewrites == 1

    out = _write(editor, tmp_path / "out.docx")
    # The projection excludes <w:del> and descends into <w:ins>, so it reads
    # the document with the change accepted.
    assert _projection_of(out) == text.replace("aim3,4.", "aim 3 and 4.")


def test_a_comment_spanning_emphasis_brackets_all_three_runs(
    tmp_path: Path,
) -> None:
    source = tmp_path / "src.docx"
    text = _emphasis_doc(source)
    quote = "was large and clear"
    start = text.index(quote)

    editor = DocxEditor(str(source))
    report = apply_annotations(
        editor,
        rewrites=[],
        notes=[
            Note(start=start, end=start + len(quote), body="crosses runs", author="t")
        ],
        reviewed_text=text,
    )
    assert report.applied_notes == 1

    out = _write(editor, tmp_path / "out.docx")
    assert _comment_spans(out) == {"1": quote}
    # …and the document still says exactly what it said. Bracketing a range
    # splits runs; splitting a run must move no character.
    assert _projection_of(out) == text


def test_the_second_of_two_identical_sentences_is_the_one_annotated(
    tmp_path: Path,
) -> None:
    """A repeated sentence identifies no passage — but an offset does.

    ``TextMap.find`` refuses an ambiguous target, and rightly: anchoring to
    the wrong copy costs the reader's trust in every other annotation. With
    coordinates there is no ambiguity to refuse.
    """
    source = tmp_path / "src.docx"
    text = _repeated_doc(source)
    sentence = "The effect was large and clear."
    first = text.index(sentence)
    second = text.index(sentence, first + 1)
    assert first != second

    editor = DocxEditor(str(source))
    report = apply_annotations(
        editor,
        rewrites=[],
        notes=[
            Note(
                start=second,
                end=second + len(sentence),
                body="the second copy",
                author="t",
            )
        ],
        reviewed_text=text,
    )
    assert report.applied_notes == 1

    out = _write(editor, tmp_path / "out.docx")
    assert _comment_spans(out) == {"1": sentence}
    # Both copies read the same, so the text cannot tell them apart — the
    # paragraph can. Paragraph 2 (0-indexed) is the second copy.
    assert _comment_paragraph_indexes(out) == {"1": 2}


def test_a_rewrite_of_the_second_copy_leaves_the_first_alone(
    tmp_path: Path,
) -> None:
    source = tmp_path / "src.docx"
    text = _repeated_doc(source)
    sentence = "The effect was large and clear."
    second = text.index(sentence, text.index(sentence) + 1)

    editor = DocxEditor(str(source))
    apply_annotations(
        editor,
        rewrites=[
            Rewrite(
                start=second,
                end=second + len(sentence),
                new_text="The effect was modest.",
                author="t",
            )
        ],
        notes=[],
        reviewed_text=text,
    )
    out = _write(editor, tmp_path / "out.docx")
    assert _projection_of(out) == (
        text[:second] + "The effect was modest." + text[second + len(sentence) :]
    )


# ── the refusals ────────────────────────────────────────────────────────


def test_a_different_document_is_refused_not_searched(tmp_path: Path) -> None:
    """The whole point: there is no fallback to look for the text."""
    source = tmp_path / "src.docx"
    text = _emphasis_doc(source)

    editor = DocxEditor(str(source))
    with pytest.raises(CoordinateError, match="does not read as the text"):
        apply_annotations(
            editor,
            rewrites=[],
            notes=[Note(start=0, end=3, body="b", author="t")],
            reviewed_text=text.replace("large", "small"),
        )


def test_an_offset_past_the_end_is_refused(tmp_path: Path) -> None:
    source = tmp_path / "src.docx"
    text = _emphasis_doc(source)

    editor = DocxEditor(str(source))
    with pytest.raises(CoordinateError, match="is not a range in a document"):
        apply_annotations(
            editor,
            rewrites=[
                Rewrite(
                    start=len(text) - 2,
                    end=len(text) + 40,
                    new_text="x",
                    author="t",
                    label="runs off the end",
                )
            ],
            notes=[],
            reviewed_text=text,
        )


def test_a_reversed_span_is_refused(tmp_path: Path) -> None:
    source = tmp_path / "src.docx"
    text = _emphasis_doc(source)
    editor = DocxEditor(str(source))
    with pytest.raises(CoordinateError):
        apply_annotations(
            editor,
            rewrites=[],
            notes=[Note(start=9, end=4, body="b", author="t")],
            reviewed_text=text,
        )


# ── overlap policy ──────────────────────────────────────────────────────


def test_two_comments_over_one_sentence_both_land(tmp_path: Path) -> None:
    """Overlapping comments are an ordinary review, not an edge case.

    Isolating one comment's span splits runs, which invalidates the run
    addresses the next comment holds — so the splits are made once, at every
    boundary, before any comment takes its runs.
    """
    source = tmp_path / "src.docx"
    text = _emphasis_doc(source)
    outer = text.index("The effect was large and clear")
    inner = text.index("was large")

    editor = DocxEditor(str(source))
    report = apply_annotations(
        editor,
        rewrites=[],
        notes=[
            Note(
                start=outer,
                end=outer + len("The effect was large and clear"),
                body="outer",
                author="t",
            ),
            Note(start=inner, end=inner + len("was large"), body="inner", author="t"),
        ],
        reviewed_text=text,
    )
    assert report.applied_notes == 2
    out = _write(editor, tmp_path / "out.docx")
    assert sorted(_comment_spans(out).values()) == sorted(
        ["The effect was large and clear", "was large"]
    )
    assert _projection_of(out) == text


def test_two_comments_on_the_identical_span_both_land(tmp_path: Path) -> None:
    source = tmp_path / "src.docx"
    text = _emphasis_doc(source)
    start = text.index("large")
    editor = DocxEditor(str(source))
    report = apply_annotations(
        editor,
        rewrites=[],
        notes=[
            Note(start=start, end=start + 5, body="one", author="t"),
            Note(start=start, end=start + 5, body="two", author="t"),
        ],
        reviewed_text=text,
    )
    assert report.applied_notes == 2
    out = _write(editor, tmp_path / "out.docx")
    assert list(_comment_spans(out).values()) == ["large", "large"]


def test_a_note_the_rewrite_covers_moves_onto_the_new_text(tmp_path: Path) -> None:
    """A finding and the fix for it belong together, not one instead of the other.

    The rewrite removes the runs its span covers, so the note cannot bracket
    the old wording — but the rewrite puts new words in their place, and the
    note goes on those. Measured on a real manuscript before this: 9 of 18
    comments dropped, every one of them the reason for a tracked change the
    author was then shown without it.
    """
    source = tmp_path / "src.docx"
    text = _emphasis_doc(source)
    start = text.index("large")

    editor = DocxEditor(str(source))
    report = apply_annotations(
        editor,
        rewrites=[
            Rewrite(start=start, end=start + 5, new_text="substantial", author="t")
        ],
        notes=[Note(start=start, end=start + 5, body="on the same word", author="t")],
        reviewed_text=text,
    )
    assert report.applied_rewrites == 1
    assert report.applied_notes == 1
    assert [n.body for n in report.moved_notes] == ["on the same word"]
    assert report.shadowed_notes == []

    out = _write(editor, tmp_path / "out.docx")
    spans = _comment_spans(out)
    assert len(spans) == 1
    # It brackets the replacement, which is what the reader needs to see.
    assert "substantial" in next(iter(spans.values()))


def test_a_note_over_a_pure_deletion_has_nothing_to_move_to(tmp_path: Path) -> None:
    """A rewrite that only removes text leaves no new wording to point at."""
    source = tmp_path / "src.docx"
    text = _emphasis_doc(source)
    start = text.index("large")

    editor = DocxEditor(str(source))
    report = apply_annotations(
        editor,
        rewrites=[Rewrite(start=start, end=start + 6, new_text="", author="t")],
        notes=[Note(start=start, end=start + 5, body="on the same word", author="t")],
        reviewed_text=text,
    )
    assert report.applied_notes == 0
    assert [n.body for n in report.shadowed_notes] == ["on the same word"]


def test_a_note_clear_of_the_rewrite_survives_it(tmp_path: Path) -> None:
    source = tmp_path / "src.docx"
    text = _emphasis_doc(source)
    editor = DocxEditor(str(source))
    report = apply_annotations(
        editor,
        rewrites=[
            Rewrite(
                start=text.index("large"),
                end=text.index("large") + 5,
                new_text="substantial",
                author="t",
            )
        ],
        notes=[
            Note(
                start=text.index("every cohort"),
                end=text.index("every cohort") + len("every cohort"),
                body="which cohorts?",
                author="t",
            )
        ],
        reviewed_text=text,
    )
    assert (report.applied_rewrites, report.applied_notes) == (1, 1)
    out = _write(editor, tmp_path / "out.docx")
    assert _comment_spans(out) == {"1": "every cohort"}


def test_two_rewrites_over_one_span_leave_one_and_report_the_other(
    tmp_path: Path,
) -> None:
    source = tmp_path / "src.docx"
    text = _emphasis_doc(source)
    start = text.index("large and clear")

    editor = DocxEditor(str(source))
    report = apply_annotations(
        editor,
        rewrites=[
            Rewrite(
                start=start,
                end=start + 5,
                new_text="substantial",
                author="t",
                label="first",
            ),
            Rewrite(
                start=start,
                end=start + 15,
                new_text="unambiguous",
                author="t",
                label="second",
            ),
        ],
        notes=[],
        reviewed_text=text,
    )
    assert report.applied_rewrites == 1
    assert [r.label for r in report.conflicted_rewrites] == ["second"]


# ── a span that crosses a paragraph ─────────────────────────────────────


def test_a_comment_may_bracket_across_a_paragraph_break(tmp_path: Path) -> None:
    source = tmp_path / "src.docx"
    text = _repeated_doc(source)
    start = text.index("analysis.")
    end = text.index("The effect was large and clear.", start) + len("The effect")

    editor = DocxEditor(str(source))
    report = apply_annotations(
        editor,
        rewrites=[],
        notes=[Note(start=start, end=end, body="spans a break", author="t")],
        reviewed_text=text,
    )
    assert report.applied_notes == 1
    out = _write(editor, tmp_path / "out.docx")
    # The paragraph break itself is not <w:t> text, so the readback joins the
    # two sides — which is exactly the span that was asked for.
    assert _comment_spans(out) == {"1": "analysis.The effect"}


def test_an_empty_batch_changes_nothing(tmp_path: Path) -> None:
    source = tmp_path / "src.docx"
    text = _emphasis_doc(source)
    editor = DocxEditor(str(source))
    report = apply_annotations(editor, rewrites=[], notes=[], reviewed_text=text)
    assert (report.total, report.applied) == (0, 0)
    out = _write(editor, tmp_path / "out.docx")
    assert _projection_of(out) == text


def test_the_review_path_never_searches_for_its_text() -> None:
    """A source scan, in the shape ``test_single_text_representation`` uses.

    An alignment module was built and deleted (65d4b9e). This is the gate that
    notices it coming back: neither the renderer nor the annotation layer may
    reach for a matcher, because by the time they run the coordinates are
    already known.
    """
    forbidden = ("TextMap(", ".find(", ".find_all(", ".closest(", "text_pattern")
    for module in ("docx/annotate.py", "renderers/docx.py"):
        path = Path(__file__).resolve().parents[1] / module
        body = "\n".join(
            line
            for line in path.read_text().splitlines()
            if not line.lstrip().startswith("#")
        )
        # The docstring is allowed to name what was removed; the code is not.
        code = body.split('"""')
        code = "".join(code[::2]) if len(code) > 2 else body
        for token in forbidden:
            if token == "TextMap(" and module == "docx/annotate.py":
                continue  # apply_edits takes one; it is never searched
            assert token not in code, f"{module} reaches for {token}"


def test_a_note_over_a_sentence_still_covers_the_sentence(tmp_path: Path) -> None:
    """A finding quotes a passage; an edit changes a few words inside it.

    That pairing is the common case — both come from a model reading the same
    text, so they land on the same troublesome sentences — and most of the
    note's span survives the rewrite. Two earlier attempts got it wrong: the
    first dropped the note, the second pinned a 79-character finding to a
    two-word insertion. The note's offsets are translated through the applied
    rewrites, so it covers its whole passage as that passage now reads.
    """
    source = tmp_path / "src.docx"
    doc = Document()
    doc.add_paragraph(
        "This means a reviewer encounters greater effort when auditing claims, "
        "which undermines the premise."
    )
    doc.save(str(source))

    editor = DocxEditor(str(source))
    text = projection(document_atoms(editor.paragraph_elements))
    sentence = "This means a reviewer encounters greater effort when auditing claims"
    note_start = text.index(sentence)
    word_start = text.index("greater effort")

    report = apply_annotations(
        editor,
        rewrites=[
            Rewrite(
                start=word_start,
                end=word_start + len("greater effort"),
                new_text="more work",
                author="t",
            )
        ],
        notes=[
            Note(
                start=note_start,
                end=note_start + len(sentence),
                body="The audit burden is asserted, never quantified.",
                author="t",
            )
        ],
        reviewed_text=text,
    )

    assert report.applied_rewrites == 1
    assert report.applied_notes == 1
    assert [n.body for n in report.moved_notes] == [
        "The audit burden is asserted, never quantified."
    ]

    out = _write(editor, tmp_path / "out.docx")
    bracketed = next(iter(_comment_spans(out).values()))
    # The whole sentence, not just the words the rewrite touched.
    assert bracketed.startswith("This means a reviewer")
    assert bracketed.rstrip().endswith("auditing claims")
    assert "more work" in bracketed


def test_a_note_clear_of_every_rewrite_is_not_reported_as_moved(tmp_path: Path) -> None:
    """Only spans a rewrite actually shifted count as moved."""
    source = tmp_path / "src.docx"
    doc = Document()
    doc.add_paragraph("First sentence stands. Second sentence gets edited later.")
    doc.save(str(source))

    editor = DocxEditor(str(source))
    text = projection(document_atoms(editor.paragraph_elements))
    later = text.index("gets edited")

    report = apply_annotations(
        editor,
        rewrites=[
            Rewrite(start=later, end=later + 11, new_text="was revised", author="t")
        ],
        notes=[Note(start=0, end=22, body="on the first", author="t")],
        reviewed_text=text,
    )
    assert report.applied_notes == 1
    assert report.moved_notes == []
