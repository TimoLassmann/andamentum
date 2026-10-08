"""The default run: review, rewrites and the reference check in one command.

The review is stubbed; the reference check runs for real on a small Word file
whose citation [2] has no entry. Crossref is refused by ``conftest.py``, so
the lookup here is either stubbed or turned off with --offline.
"""

from __future__ import annotations

import zipfile
from pathlib import Path
from typing import Any

import pytest
from docx import Document
from rich.console import Console

from andamentum.references import LookupResult
from andamentum.whetstone import cli as cli_mod
from andamentum.whetstone.references_check import REFERENCE_CHECKER
from andamentum.whetstone.schemas import Finding, Quote, ReviewResult


def _args(argv: list[str]):
    args = cli_mod.build_parser().parse_args(
        argv if argv[0] in cli_mod._COMMANDS else ["review", *argv]
    )
    cli_mod._validate_args(args)
    return args


def _draft(path: Path, *, first_line: str = "Methods") -> Path:
    document = Document()
    document.add_heading(first_line, level=1)
    document.add_paragraph(
        "Rare diseases are common [1]. Costs are high [2]. Care is slow [3]."
    )
    document.add_heading("References", level=1)
    document.add_paragraph(
        "1. Smith J (2020) A study of rare disease costs. J Test 1:1"
    )
    document.add_paragraph("3. Lee K (2021) Delays in rare disease care. J Test 2:2")
    document.save(str(path))
    return path


def _stub_review(
    monkeypatch: pytest.MonkeyPatch, *, findings: list[Finding] | None = None
) -> dict:
    """A review that behaves like the real one toward extra findings: it lists
    them after its own, which is what ``Finalize`` does."""
    seen: dict[str, Any] = {}

    async def fake_review(_source: Any, **kwargs: Any) -> ReviewResult:
        seen.update(kwargs)
        return ReviewResult(
            summary="Stubbed.",
            findings=[*(findings or []), *(kwargs.get("extra_findings") or [])],
        )

    monkeypatch.setattr("andamentum.whetstone.pipeline.review_document", fake_review)
    return seen


def _found_everything(monkeypatch: pytest.MonkeyPatch) -> list[int]:
    calls: list[int] = []

    async def lookup(entries, **_kwargs):
        calls.append(len(entries))
        return [LookupResult(entry_number=e.number, status="found") for e in entries]

    monkeypatch.setattr("andamentum.references.lookup_entries", lookup)
    return calls


def test_the_editor_the_reference_check_and_the_lookup_are_on_by_default() -> None:
    args = _args(["draft.md", "--model", "stub:none", "--out", "review.md"])
    assert (args.editor, args.references, args.offline) == (True, True, False)


async def test_reference_problems_reach_the_editor_and_the_report(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    source = _draft(tmp_path / "draft.docx")
    seen = _stub_review(monkeypatch)
    lookups = _found_everything(monkeypatch)
    out = tmp_path / "review.md"
    word = tmp_path / "review.docx"
    await cli_mod._run(
        _args(
            [str(source), "--model", "stub:none", "--out", str(out), "--out", str(word)]
        ),
        Console(stderr=True),
    )
    # In Word the checker signs its own comment, whoever signs the rest.
    with zipfile.ZipFile(word) as z:
        comments = z.read("word/comments.xml").decode()
    assert 'w:author="Reference checker (andamentum-whetstone, AI)"' in comments
    assert "Citation 2 is not in the list" in comments
    # The single reviewer's editor was handed the problem, signed by the checker.
    [problem] = seen["extra_findings"]
    assert problem.title == "Citation 2 is not in the list"
    assert problem.perspective == REFERENCE_CHECKER
    assert seen["editor"] is True
    assert lookups == [2]
    report = out.read_text()
    assert "## Reference check" in report
    assert "Crossref lookup:** 2 found as written" in report
    # The next step names the revise command, writing elsewhere.
    assert "## Next step" in report
    assert f"andamentum-whetstone revise {source}" in report
    assert "--until-clean citations" in report
    assert f"-o {tmp_path / 'review.revised.md'}" in report
    assert "--until-clean structure" not in report  # no review finding was raised


async def test_a_major_review_finding_suggests_the_structure_loop(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    source = _draft(tmp_path / "draft.docx")
    major = Finding(
        title="The claim outruns the data",
        severity="major",
        confidence="high",
        rationale="x",
        quotes=[Quote(section_id="s1", char_start=0, char_end=4, text="Rare")],
    )
    _stub_review(monkeypatch, findings=[major])
    out = tmp_path / "review.md"
    await cli_mod._run(
        _args(
            [
                str(source),
                "--model",
                "stub:none",
                "--reviewer",
                "claims",
                "--no-references",
                "--out",
                str(out),
            ]
        ),
        Console(stderr=True),
    )
    report = out.read_text()
    assert "This review raised 1 major or moderate finding." in report
    assert "--reviewer claims" in report and "--until-clean structure" in report
    assert "## Reference check" not in report


async def test_no_references_and_offline_turn_the_parts_off(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    source = _draft(tmp_path / "draft.docx")
    seen = _stub_review(monkeypatch)
    out = tmp_path / "off.md"
    await cli_mod._run(
        _args(
            [str(source), "--model", "stub:none", "--no-references", "--out", str(out)]
        ),
        Console(stderr=True),
    )
    assert seen["extra_findings"] is None
    assert "Reference check" not in out.read_text()

    # --offline: still checked, without the lookup (the conftest refuses any).
    out = tmp_path / "offline.md"
    await cli_mod._run(
        _args([str(source), "--model", "stub:none", "--offline", "--out", str(out)]),
        Console(stderr=True),
    )
    assert "Crossref lookup:** off" in out.read_text()


async def test_the_confidentiality_check_runs_before_any_reference_leaves(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    source = _draft(tmp_path / "draft.docx", first_line="Manuscript ID: ABC-123")
    _stub_review(monkeypatch)
    lookups = _found_everything(monkeypatch)
    with pytest.raises(SystemExit) as exc:
        await cli_mod._run(
            _args(
                [str(source), "--model", "stub:none", "--out", str(tmp_path / "r.md")]
            ),
            Console(stderr=True),
        )
    assert exc.value.code == 2
    assert lookups == []


def test_check_novelty_with_offline_is_a_usage_error(
    capsys: pytest.CaptureFixture[str],
) -> None:
    with pytest.raises(SystemExit) as exc:
        _args(["d.md", "--model", "m", "--out", "r.md", "--offline", "--check-novelty"])
    assert exc.value.code == 2
    assert "--check-novelty searches the scholarly sources" in capsys.readouterr().err


async def test_a_brief_reaches_the_review_as_guidelines(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """--brief is the guidelines path: its text (or the file's, with @) goes
    to the review as ``guidelines_text``, and no inline criteria are passed."""
    source = _draft(tmp_path / "draft.docx")
    brief = tmp_path / "brief.md"
    brief.write_text("Ask whether the sample is large enough.")
    for value, expected in (
        ("Is the abstract honest?", "Is the abstract honest?"),
        (f"@{brief}", "Ask whether the sample is large enough."),
    ):
        seen = _stub_review(monkeypatch)
        await cli_mod._run(
            _args(
                [
                    str(source), "--model", "stub:none", "--no-references",
                    "--brief", value, "--out", str(tmp_path / "r.md"),
                ]
            ),
            Console(stderr=True),
        )
        assert seen["guidelines_text"] == expected
        # The reviewer sets the editor criteria; the command line passes none.
        assert "criteria" not in seen and "editor_criteria" not in seen


async def test_an_unreadable_brief_file_exits_three(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    source = _draft(tmp_path / "draft.docx")
    _stub_review(monkeypatch)
    with pytest.raises(SystemExit) as exc:
        await cli_mod._run(
            _args(
                [
                    str(source), "--model", "stub:none",
                    "--brief", f"@{tmp_path / 'missing.md'}",
                    "--out", str(tmp_path / "r.md"),
                ]
            ),
            Console(stderr=True),
        )
    assert exc.value.code == 3


def test_checklist_refuses_the_review_switches(
    capsys: pytest.CaptureFixture[str],
) -> None:
    with pytest.raises(SystemExit):
        _args(
            [
                "checklist",
                "d.md",
                "--model",
                "m",
                "--out",
                "r.md",
                "--offline",
            ]
        )
    assert "unrecognized arguments: --offline" in capsys.readouterr().err
