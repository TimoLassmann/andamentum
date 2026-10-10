"""The atom stream is the single text representation. These pin its contract.

Each test here corresponds to a way one of the three previous representations
got the same paragraph wrong.
"""

from __future__ import annotations

from docx import Document
from lxml import etree  # type: ignore

from andamentum.docx_edit.atoms import (
    AtomKind,
    document_atoms,
    paragraph_atoms,
    projection,
    qn,
)


def _paragraph(*, text: str = "") -> object:
    """A bare ``<w:p>`` element, optionally with one run of text."""
    doc = Document()
    para = doc.add_paragraph()
    if text:
        para.add_run(text)
    return para._p


def _append_run(p_elem, text: str):
    run = etree.SubElement(p_elem, qn("r"))
    t = etree.SubElement(run, qn("t"))
    t.text = text
    return run


# ── the three failures this replaces ────────────────────────────────────


def test_line_break_renders_as_a_space_and_does_not_weld_words() -> None:
    """The ``<w:t>`` join produced 'readingfurther.'; a reader sees a break."""
    p = _paragraph()
    _append_run(p, "before reading")
    break_run = etree.SubElement(p, qn("r"))
    etree.SubElement(break_run, qn("br"))
    _append_run(p, "further.")

    assert projection(list(paragraph_atoms(p))) == "before reading further."


def test_hyperlink_text_is_visible() -> None:
    """Direct-child run collection lost this entirely — a hyperlink nests."""
    p = _paragraph()
    _append_run(p, "See ")
    link = etree.SubElement(p, qn("hyperlink"))
    run = etree.SubElement(link, qn("r"))
    etree.SubElement(run, qn("t")).text = "the plan"
    _append_run(p, " first.")

    assert projection(list(paragraph_atoms(p))) == "See the plan first."


def test_text_inside_an_unfamiliar_container_is_still_found() -> None:
    """Descend-by-default: a container this module has never heard of works.

    This is the property the old include-by-name walkers could not have. A
    content control, a smart tag or next year's element needs no change here.
    """
    p = _paragraph()
    _append_run(p, "Before ")
    wrapper = etree.SubElement(p, qn("sdt"))
    content = etree.SubElement(wrapper, qn("sdtContent"))
    run = etree.SubElement(content, qn("r"))
    etree.SubElement(run, qn("t")).text = "inside a content control"
    _append_run(p, " after.")

    text = projection(list(paragraph_atoms(p)))
    assert text == "Before inside a content control after."


# ── inclusion and exclusion rules ───────────────────────────────────────


def test_inserted_text_is_part_of_the_document() -> None:
    """The stream is the document with changes accepted, so a later round can
    edit what an earlier round inserted."""
    p = _paragraph()
    _append_run(p, "kept ")
    ins = etree.SubElement(p, qn("ins"))
    run = etree.SubElement(ins, qn("r"))
    etree.SubElement(run, qn("t")).text = "added"

    assert projection(list(paragraph_atoms(p))) == "kept added"


def test_deleted_text_is_not_part_of_the_document() -> None:
    """Text on its way out is not what the document says."""
    p = _paragraph()
    _append_run(p, "kept ")
    dele = etree.SubElement(p, qn("del"))
    run = etree.SubElement(dele, qn("r"))
    etree.SubElement(run, qn("delText")).text = "removed"
    _append_run(p, "tail")

    assert projection(list(paragraph_atoms(p))) == "kept tail"


def test_tab_renders_as_a_space() -> None:
    p = _paragraph()
    _append_run(p, "Column one")
    tab_run = etree.SubElement(p, qn("r"))
    etree.SubElement(tab_run, qn("tab"))
    _append_run(p, "Column two")

    assert projection(list(paragraph_atoms(p))) == "Column one Column two"


# ── provenance ──────────────────────────────────────────────────────────


def test_every_text_atom_addresses_a_run_and_an_offset() -> None:
    p = _paragraph()
    first = _append_run(p, "alpha ")
    second = _append_run(p, "beta")
    stream = list(paragraph_atoms(p))

    assert projection(stream) == "alpha beta"
    assert stream[0].run is first and stream[0].offset == 0
    assert stream[5].run is first and stream[5].offset == 5
    assert stream[6].run is second and stream[6].offset == 0
    assert all(a.kind is AtomKind.TEXT for a in stream)


def test_structural_atoms_are_not_editable() -> None:
    p = _paragraph()
    _append_run(p, "a")
    break_run = etree.SubElement(p, qn("r"))
    etree.SubElement(break_run, qn("br"))
    _append_run(p, "b")
    stream = list(paragraph_atoms(p))

    assert [a.kind for a in stream] == [
        AtomKind.TEXT,
        AtomKind.STRUCTURAL,
        AtomKind.TEXT,
    ]
    assert [a.editable for a in stream] == [True, False, True]


# ── paragraph breaks are characters ─────────────────────────────────────


def test_paragraph_break_is_a_character_in_the_stream() -> None:
    """Modelling the break as an atom is what makes a cross-paragraph edit an
    ordinary span instead of a special case."""
    doc = Document()
    doc.add_paragraph("first")
    doc.add_paragraph("second")
    paragraphs = [p._p for p in doc.paragraphs]
    stream = document_atoms(paragraphs)

    assert projection(stream) == "first\nsecond"
    break_atom = stream[5]
    assert break_atom.kind is AtomKind.PARAGRAPH_BREAK
    assert break_atom.run is None
    # Attributed to the paragraph it ends, which is where Word keeps the mark.
    assert break_atom.paragraph is paragraphs[0]
    # Deletable: removing it is how two paragraphs merge.
    assert break_atom.editable is True


def test_no_break_before_the_first_paragraph() -> None:
    doc = Document()
    doc.add_paragraph("only")
    stream = document_atoms([p._p for p in doc.paragraphs])
    assert projection(stream) == "only"


def test_projection_of_an_empty_document_is_empty() -> None:
    assert projection(document_atoms([])) == ""


# ── paragraphs nested inside paragraphs ─────────────────────────────────


def test_text_box_content_is_read_once_not_twice(tmp_path) -> None:
    """A text box holds a <w:p> inside the anchoring paragraph's run.

    The walker descends into it, so if the paragraph list also contained that
    nested paragraph the text would be read twice — once in place and once as
    a paragraph of its own. Found by comparing marker order against Docling's
    output, where the box appeared once and the projection had it twice.
    """
    from andamentum.whetstone.docx.low_level import DocxEditor

    namespace = qn("t")[1:].split("}")[0]
    doc = Document()
    doc.add_paragraph("Before the box.")
    run = doc.add_paragraph().add_run()
    run._r.append(
        etree.fromstring(
            f'<w:pict xmlns:w="{namespace}" xmlns:v="urn:schemas-microsoft-com:vml">'
            '<v:shape type="#_x0000_t202"><v:textbox><w:txbxContent>'
            "<w:p><w:r><w:t>TEXTBOX CONTENT</w:t></w:r></w:p>"
            "</w:txbxContent></v:textbox></v:shape></w:pict>"
        )
    )
    doc.add_paragraph("After the box.")
    source = tmp_path / "textbox.docx"
    doc.save(str(source))

    elements = DocxEditor(str(source)).paragraph_elements
    text = projection(document_atoms(elements))

    assert text.count("TEXTBOX CONTENT") == 1
    assert text == "Before the box.\nTEXTBOX CONTENT\nAfter the box."


def test_table_cells_are_still_their_own_paragraphs(tmp_path) -> None:
    """A <w:tbl> is a sibling of body paragraphs, so excluding nested
    paragraphs must not swallow the cells."""
    from andamentum.whetstone.docx.low_level import DocxEditor

    doc = Document()
    doc.add_paragraph("Before the table.")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Signal"
    table.cell(0, 1).text = "Coverage"
    doc.add_paragraph("After the table.")
    source = tmp_path / "table.docx"
    doc.save(str(source))

    text = projection(document_atoms(DocxEditor(str(source)).paragraph_elements))
    assert "Signal" in text and "Coverage" in text
    assert text.count("Signal") == 1
