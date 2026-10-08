"""The ``revise`` command: flag validation, plus end-to-end checks that the
command drives the revision loop and exits with the right code.

Most of this file mirrors ``test_reviewer_cli.py``'s pattern: parse with the
real parser, then call ``_validate_args`` directly and check the exit code
and message — the argument-time refusals, with no review run at all. The
end-to-end tests mirror ``test_cli_review_wiring.py`` instead: they stub
``review_document`` at the loop's own import site (not the pydantic-ai
agents underneath it — that round-by-round behaviour is
``test_revise_loop.py``'s job) and check that ``_run`` writes real output and
exits with the right code.
"""

from __future__ import annotations

from typing import Any

import pytest
from docx import Document
from rich.console import Console

from andamentum.whetstone import cli as cli_mod
from andamentum.whetstone.revise import FIXED_CONDITION_TEXT
from andamentum.whetstone.schemas import Edit, Finding, ReviewResult

BASE = ["revise", "draft.md", "--model", "stub:none", "--out", "review.md"]
UNTIL = [*BASE, "--until", "the flow is correct"]


def _args(argv: list[str]):
    return cli_mod.build_parser().parse_args(argv)


def _validate(argv: list[str]) -> None:
    cli_mod._validate_args(_args(argv))


def _refused(argv: list[str], capsys: pytest.CaptureFixture[str], *, code: int = 2) -> str:
    with pytest.raises(SystemExit) as exc:
        _validate(argv)
    assert exc.value.code == code
    return capsys.readouterr().err


# ── the stop condition ────────────────────────────────────────────────────


def test_a_condition_in_words_is_accepted() -> None:
    args = _args(UNTIL)
    cli_mod._validate_args(args)
    assert (args.until, args.until_clean) == ("the flow is correct", None)


@pytest.mark.parametrize("check", ["structure", "citations"])
def test_until_clean_needs_no_text(check: str) -> None:
    args = _args([*BASE, "--until-clean", check])
    cli_mod._validate_args(args)
    assert (args.until, args.until_clean) == (None, check)


def test_one_of_until_and_until_clean_is_required(
    capsys: pytest.CaptureFixture[str],
) -> None:
    err = _refused(BASE, capsys)
    assert "one of the arguments --until --until-clean is required" in err


def test_until_and_until_clean_are_mutually_exclusive(
    capsys: pytest.CaptureFixture[str],
) -> None:
    err = _refused([*UNTIL, "--until-clean", "structure"], capsys)
    assert "not allowed with argument" in err


def test_an_unknown_until_clean_check_is_refused_by_argparse(
    capsys: pytest.CaptureFixture[str],
) -> None:
    err = _refused([*BASE, "--until-clean", "vibes"], capsys)
    assert "invalid choice: 'vibes'" in err


def test_an_empty_condition_is_refused(capsys: pytest.CaptureFixture[str]) -> None:
    assert "--until needs a condition" in _refused([*BASE, "--until", "  "], capsys)


def test_the_round_cap_is_not_a_flag(capsys: pytest.CaptureFixture[str]) -> None:
    err = _refused([*UNTIL, "--max-rounds", "2"], capsys)
    assert "unrecognized arguments: --max-rounds 2" in err


def test_until_clean_citations_refuses_no_references(
    capsys: pytest.CaptureFixture[str],
) -> None:
    err = _refused([*BASE, "--until-clean", "citations", "--no-references"], capsys)
    assert err.strip() == (
        "error: --until-clean citations runs the reference check every round; "
        "drop --no-references"
    )


@pytest.mark.parametrize("check", ["structure", "figures"])
def test_no_references_with_another_check_is_accepted(check: str, tmp_path) -> None:
    extra = []
    if check == "figures":
        figure = tmp_path / "fig1.png"
        figure.write_bytes(b"x")
        extra = ["--figure", f"Figure 1={figure}", "--figure-model", "m"]
    _validate([*BASE, "--until-clean", check, "--no-references", *extra])


def test_a_reviewer_list_is_refused(capsys: pytest.CaptureFixture[str]) -> None:
    err = _refused([*UNTIL, "--reviewer", "claims,methods"], capsys)
    assert "revise takes one reviewer, not a list" in err


def test_one_named_reviewer_is_accepted() -> None:
    _validate([*UNTIL, "--reviewer", "reader"])


@pytest.mark.parametrize(
    "flag",
    [
        ["--until", "x"],
        ["--until-clean", "structure"],
        ["--figure", "Figure 1=f.png"],
        ["--figure-model", "m"],
    ],
)
def test_the_revise_flags_are_not_review_flags(
    flag: list[str], capsys: pytest.CaptureFixture[str]
) -> None:
    err = _refused(["review", "draft.md", "--model", "m", "--out", "r.md", *flag], capsys)
    assert "unrecognized arguments" in err


def test_revise_runs_the_editor_every_round_so_there_is_no_editor_flag(
    capsys: pytest.CaptureFixture[str],
) -> None:
    assert "unrecognized arguments: --no-editor" in _refused([*UNTIL, "--no-editor"], capsys)


# ── --figure / --figure-model ─────────────────────────────────────────────


def test_figure_without_until_clean_figures_is_refused(
    tmp_path, capsys: pytest.CaptureFixture[str]
) -> None:
    figure = tmp_path / "fig1.png"
    figure.write_bytes(b"")
    err = _refused([*UNTIL, "--figure", f"Figure 1={figure}"], capsys)
    assert "--figure only applies to --until-clean figures" in err


def test_figure_model_without_until_clean_figures_is_refused(
    capsys: pytest.CaptureFixture[str],
) -> None:
    err = _refused([*BASE, "--until-clean", "structure", "--figure-model", "m"], capsys)
    assert "--figure-model only applies" in err


def test_until_clean_figures_needs_at_least_one_figure(
    capsys: pytest.CaptureFixture[str],
) -> None:
    err = _refused(
        [*BASE, "--until-clean", "figures", "--figure-model", "ollama:qwen3.8"], capsys
    )
    assert "needs at least one --figure" in err


def test_until_clean_figures_needs_a_figure_model(
    tmp_path, capsys: pytest.CaptureFixture[str]
) -> None:
    figure = tmp_path / "fig1.png"
    figure.write_bytes(b"")
    err = _refused(
        [*BASE, "--until-clean", "figures", "--figure", f"Figure 1={figure}"], capsys
    )
    assert "needs --figure-model" in err


def test_a_figure_with_a_missing_file_exits_three(
    capsys: pytest.CaptureFixture[str],
) -> None:
    err = _refused(
        [
            *BASE,
            "--until-clean",
            "figures",
            "--figure",
            "Figure 1=does/not/exist.png",
            "--figure-model",
            "ollama:qwen3.8",
        ],
        capsys,
        code=3,
    )
    assert "no such file" in err


def test_a_figure_needs_label_and_equals(capsys: pytest.CaptureFixture[str]) -> None:
    err = _refused(
        [
            *BASE,
            "--until-clean",
            "figures",
            "--figure",
            "not-a-label-path-pair",
            "--figure-model",
            "ollama:qwen3.8",
        ],
        capsys,
    )
    assert "is not LABEL=FILE" in err


def test_a_valid_figures_run_is_accepted(tmp_path) -> None:
    figure = tmp_path / "fig1.png"
    figure.write_bytes(b"")
    args = _args(
        [
            *BASE,
            "--until-clean",
            "figures",
            "--figure",
            f"Figure 1={figure}",
            "--figure-model",
            "ollama:qwen3.8",
        ]
    )
    cli_mod._validate_args(args)
    assert args.figures == [("Figure 1", figure)]


# ── end-to-end: the command actually drives the loop ─────────────────────


def _draft(path) -> None:
    document = Document()
    document.add_heading("Methods", level=1)
    document.add_paragraph("The evidence layer earns its place in the analysis.")
    document.save(str(path))


def _review_stub(*, rounds: int):
    """Round 0 has a major finding + an edit; every later round has neither
    — a structure condition met from round 1 onward, confirmed on round 2."""
    calls: dict[str, Any] = {"n": 0, "kwargs": []}

    async def fake_review(source, **kwargs):
        idx = min(calls["n"], rounds - 1)
        calls["n"] += 1
        calls["kwargs"].append(kwargs)
        text = source.text
        if idx == 0:
            start = text.index("earns its place")
            return ReviewResult(
                summary="stub",
                findings=[
                    Finding(
                        title="vague",
                        severity="major",
                        confidence="medium",
                        rationale="x",
                        quotes=[],
                    )
                ],
                edits=[
                    Edit(
                        title="t",
                        rationale="r",
                        section_id="s1",
                        char_start=start,
                        char_end=start + len("earns its place"),
                        original_text="earns its place",
                        new_text="earns its keep",
                    )
                ],
            )
        return ReviewResult(summary="stub", findings=[], edits=[])

    return calls, fake_review


async def test_a_met_condition_writes_output_and_exits_zero(
    tmp_path, monkeypatch: pytest.MonkeyPatch, capsys: pytest.CaptureFixture[str]
) -> None:
    source = tmp_path / "draft.docx"
    _draft(source)
    out = tmp_path / "revised.docx"
    brief = tmp_path / "brief.md"
    brief.write_text("Check that every claim is supported.")

    calls, fake_review = _review_stub(rounds=3)
    monkeypatch.setattr(
        "andamentum.whetstone.revise.loop.review_document", fake_review, raising=True
    )

    code = await cli_mod._run(
        _args(
            [
                "revise",
                str(source),
                "--model",
                "stub:none",
                "--out",
                str(out),
                "--until-clean",
                "structure",
                "--brief",
                f"@{brief}",
                "--offline",
            ]
        ),
        Console(stderr=True),
    )

    assert code == 0
    assert out.exists()
    # round 0 fails; round 1 meets with nothing left to edit, so round 2
    # re-reviews the same text and confirms it.
    assert calls["n"] == 3
    # Every round reviews against the brief, and the editor runs every round.
    assert {k["guidelines_text"] for k in calls["kwargs"]} == {
        "Check that every claim is supported."
    }
    assert all(k["editor"] is True for k in calls["kwargs"])
    # The fixed wording names what the run was trying to achieve.
    assert FIXED_CONDITION_TEXT["structure"] in capsys.readouterr().err


async def test_an_unmet_condition_still_writes_output_and_exits_one(
    tmp_path, monkeypatch: pytest.MonkeyPatch
) -> None:
    source = tmp_path / "draft.docx"
    _draft(source)
    out = tmp_path / "revised.docx"

    async def always_failing(source, **_kwargs):
        return ReviewResult(
            summary="stub",
            findings=[
                Finding(
                    title="vague",
                    severity="major",
                    confidence="medium",
                    rationale="x",
                    quotes=[],
                )
            ],
            edits=[],  # no edits at all -> no_progress on round 1
        )

    monkeypatch.setattr(
        "andamentum.whetstone.revise.loop.review_document",
        always_failing,
        raising=True,
    )

    code = await cli_mod._run(
        _args(
            [
                "revise",
                str(source),
                "--model",
                "stub:none",
                "--out",
                str(out),
                "--until-clean",
                "structure",
            ]
        ),
        Console(stderr=True),
    )
    assert code == 1
    assert out.exists(), "the output is written even though the condition was not met"


async def test_until_text_is_judged_by_a_separate_call(
    tmp_path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """--until TEXT is the judged check: the text goes to llm_judge as the one
    criterion, every round, and a met verdict twice running exits 0."""
    from andamentum.llm_judge import CriterionScore, ScoreResult

    source = tmp_path / "draft.docx"
    _draft(source)
    out = tmp_path / "revised.md"
    calls, fake_review = _review_stub(rounds=3)
    monkeypatch.setattr(
        "andamentum.whetstone.revise.loop.review_document", fake_review, raising=True
    )
    judged: list[str] = []

    async def fake_judge(output, *, criteria, context=None, model):
        judged.append(criteria[0].description)
        return ScoreResult(
            per_criterion=[
                CriterionScore(
                    criterion=criteria[0].name,
                    reasoning="stubbed",
                    meets=100,
                    partial=0,
                    fails=0,
                )
            ],
            overall="meets",
            confidence=1.0,
            doubt=0.0,
            needs_review=False,
            expected_score=1.0,
        )

    monkeypatch.setattr("andamentum.llm_judge.judge_score", fake_judge)
    code = await cli_mod._run(
        _args(
            [
                "revise",
                str(source),
                "--model",
                "stub:none",
                "--out",
                str(out),
                "--until",
                "the methods are reproducible",
                "--no-references",
            ]
        ),
        Console(stderr=True),
    )
    assert code == 0
    assert judged == ["the methods are reproducible"] * 2


@pytest.mark.parametrize("check", ["structure", "citations"])
async def test_until_clean_names_the_check_the_loop_runs(
    tmp_path, monkeypatch: pytest.MonkeyPatch, check: str
) -> None:
    """--until-clean KIND reaches revise_document as that check, with its
    fixed wording; the loop itself is stubbed out."""
    seen: dict[str, Any] = {}

    async def fake_revise(_document: Any, *, spec: Any, **_kwargs: Any) -> Any:
        seen["spec"] = spec
        raise RuntimeError("stop here")

    monkeypatch.setattr("andamentum.whetstone.revise.revise_document", fake_revise)
    source = tmp_path / "draft.docx"
    _draft(source)
    argv = ["revise", str(source), "--model", "stub:none", "--out", str(tmp_path / "r.md")]
    with pytest.raises(SystemExit) as exc:
        await cli_mod._run(
            _args([*argv, "--until-clean", check, "--offline"]), Console(stderr=True)
        )
    assert exc.value.code == 4
    assert seen["spec"].check == check
    assert seen["spec"].text == FIXED_CONDITION_TEXT[check]
