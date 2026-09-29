"""There is one answer to "what does this document say", and one matcher.

A source scan, in the spirit of ``test_verdict_single_source.py`` and
``test_doubt_scale_single_source.py`` elsewhere in the repo. It exists because
the duplication it forbids is how the whole family of anchoring bugs was born:
three text projections that disagreed on the same paragraph, and two matchers
over them, so every fix had to be made more than once and a span resolved
against one could not be trusted against another.

    <w:t> join            'See the protocol at the plan before readingfurther.'
    visible_text()        'See the protocol at the plan before reading further.'
    direct-child runs     'See the protocol at  before reading further.'

These tests fail the moment a fourth appears.
"""

from __future__ import annotations

import ast
from pathlib import Path

DOCX_PACKAGE = Path(__file__).resolve().parent.parent / "docx"
ATOMS = DOCX_PACKAGE / "atoms.py"

# Reading many <w:t> nodes at once is document-text extraction, and that is
# atoms.py's job. Reading ONE run's <w:t> to splice it is mutation, which is
# surgery.py's job — hence the file allowance rather than a blanket ban.
_MULTI_TEXT_READS = (
    ".//w:t",
    'iter(qn("t"))',
    "iter(TEXT_TAG)",
)
_EXTRACTION_ALLOWED = {"atoms.py"}

# Only one module may implement locating a target in the document.
_MATCHER_ALLOWED = {"textmap.py"}


def _sources() -> list[Path]:
    return sorted(p for p in DOCX_PACKAGE.glob("*.py") if p.name != "__init__.py")


def test_only_atoms_extracts_text_from_many_nodes() -> None:
    offenders = []
    for path in _sources():
        if path.name in _EXTRACTION_ALLOWED:
            continue
        body = path.read_text()
        for pattern in _MULTI_TEXT_READS:
            if pattern in body:
                offenders.append(f"{path.name}: {pattern}")
    assert not offenders, (
        "document text must come from atoms.py, which is the single "
        "representation. Found direct multi-node text extraction in:\n  "
        + "\n  ".join(offenders)
    )


def test_only_textmap_implements_normalised_matching() -> None:
    """One matcher. Edits used to use a raw substring test and comments a
    normalised index, and the two disagreed on 16 of 27 real targets."""
    offenders = []
    for path in _sources():
        if path.name in _MATCHER_ALLOWED:
            continue
        if "normalize_for_match" in path.read_text():
            offenders.append(path.name)
    assert not offenders, (
        "normalisation belongs to textmap.py (over core.text_match). "
        f"Also found in: {offenders}"
    )


def test_no_module_rebuilds_a_paragraph_from_a_string() -> None:
    """Reconstructing a paragraph from a text projection destroys everything
    the projection dropped — line breaks, tabs, and earlier tracked changes.
    Edits cut the runs their span covers instead."""
    offenders = []
    for path in _sources():
        body = path.read_text()
        if "_clear_runs" in body and path.name != "low_level.py":
            offenders.append(f"{path.name}: clears a paragraph's runs")
        if "ndiff" in body:
            offenders.append(f"{path.name}: diffs a whole paragraph")
    assert not offenders, "\n  ".join(offenders)


def test_paragraph_data_holds_no_text_of_its_own() -> None:
    """The two fields whose disagreement started this: ``original`` and
    ``modified``. Text is a derived view now, never stored state."""
    tree = ast.parse((DOCX_PACKAGE / "low_level.py").read_text())
    paragraph_data = next(
        node
        for node in ast.walk(tree)
        if isinstance(node, ast.ClassDef) and node.name == "ParagraphData"
    )
    assigned = {
        target.attr
        for node in ast.walk(paragraph_data)
        if isinstance(node, ast.Assign)
        for target in node.targets
        if isinstance(target, ast.Attribute)
    }
    assert not ({"original", "modified", "_visible"} & assigned), (
        f"ParagraphData stores text again: {sorted(assigned)}"
    )


def test_the_atom_projection_is_what_visible_text_returns() -> None:
    """The one representation has one definition, checked rather than trusted."""
    from docx import Document
    from lxml import etree  # type: ignore

    from andamentum.whetstone.docx.atoms import paragraph_atoms, projection, qn
    from andamentum.whetstone.docx.low_level import visible_text

    doc = Document()
    para = doc.add_paragraph()
    para.add_run("Accuracy was")
    etree.SubElement(etree.SubElement(para._p, qn("r")), qn("tab"))
    para.add_run("ninety-four percent")
    etree.SubElement(etree.SubElement(para._p, qn("r")), qn("br"))
    link = etree.SubElement(para._p, qn("hyperlink"))
    run = etree.SubElement(link, qn("r"))
    etree.SubElement(run, qn("t")).text = "across every cohort."

    assert projection(list(paragraph_atoms(para._p))) == visible_text(para._p)
