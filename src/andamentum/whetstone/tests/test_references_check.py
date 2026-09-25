"""The references subcommand, end to end on a real .docx, offline."""

from __future__ import annotations

import zipfile
from pathlib import Path

import pytest
from docx import Document

from andamentum.whetstone import cli as cli_mod


def _draft(path: Path) -> None:
    document = Document()
    document.add_heading("Introduction", level=1)
    document.add_paragraph("Rare diseases are common [1]. Costs are high [3].")
    document.add_heading("References", level=1)
    document.add_paragraph(
        "1. Yang G, Cintina I (2022) The national economic burden of rare "
        "disease in the United States in 2019. Orphanet J Rare Dis 17:163"
    )
    document.add_paragraph(
        "2. Zurynski Y, Deverell M (2017) Australian children living with rare "
        "diseases. Orphanet J Rare Dis 12:68"
    )
    document.save(str(path))


def test_a_missing_citation_becomes_a_comment_on_the_citation(tmp_path: Path) -> None:
    source = tmp_path / "draft.docx"
    _draft(source)
    out_docx, out_md = tmp_path / "checked.docx", tmp_path / "checked.md"
    with pytest.raises(SystemExit) as exit_info:
        cli_mod.main(
            [
                "references",
                str(source),
                "--no-lookup",
                "--out",
                str(out_docx),
                "--out",
                str(out_md),
            ]
        )
    assert exit_info.value.code == 0

    markdown = out_md.read_text()
    assert "Citation 3 is not in the list" in markdown
    assert "Reference 2 is never cited" in markdown
    assert "Crossref lookup:** off" in markdown

    with zipfile.ZipFile(out_docx) as z:
        comments = z.read("word/comments.xml").decode()
        body = z.read("word/document.xml").decode()
    assert "Must fix: Citation 3 is not in the list" in comments
    assert "Reference checker (andamentum-whetstone, AI)" in comments
    # The comment brackets the citation itself.
    start = body.index("commentRangeStart")
    end = body.index("commentRangeEnd", start)
    assert "[3]" in body[start:end]


def test_an_output_with_an_unknown_extension_is_refused(tmp_path: Path) -> None:
    source = tmp_path / "draft.docx"
    _draft(source)
    with pytest.raises(SystemExit) as exit_info:
        cli_mod.main(["references", str(source), "--out", str(tmp_path / "x.txt")])
    assert exit_info.value.code == 1
