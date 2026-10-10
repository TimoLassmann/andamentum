"""One index, used by edits and comments alike."""

from __future__ import annotations

from docx import Document
from lxml import etree  # type: ignore

from andamentum.docx_edit.atoms import document_atoms, qn
from andamentum.docx_edit.textmap import AtomSpan, TextMap


def _map(*paragraph_texts: str) -> TextMap:
    doc = Document()
    for text in paragraph_texts:
        doc.add_paragraph(text)
    return TextMap(document_atoms([p._p for p in doc.paragraphs]))


def test_resolves_plain_text_to_atom_indices() -> None:
    tm = _map("The evidence layer earns its place.")
    span = tm.find("earns its place")
    assert span is not None
    assert tm.slice(span) == "earns its place"


def test_span_is_half_open() -> None:
    tm = _map("abcdef")
    span = tm.find("cd")
    assert span is not None
    assert span == AtomSpan(2, 4)
    assert len(span) == 2


def test_matches_across_markdown_markers_and_case() -> None:
    """The quote comes from harvested markdown; Word holds plain text."""
    tm = _map("The evidence layer earns its place.")
    span = tm.find("**Earns its place**")
    assert span is not None
    assert tm.slice(span) == "earns its place"


def test_matches_across_a_paragraph_boundary() -> None:
    tm = _map("A heading", "and the body that follows it.")
    span = tm.find("A heading and the body")
    assert span is not None
    assert tm.slice(span) == "A heading\nand the body"


def test_matches_across_a_line_break() -> None:
    """The break is one space in the projection, so a quote reads through it."""
    doc = Document()
    p = doc.add_paragraph()
    p.add_run("before reading")
    p.add_run().add_break()
    p.add_run("further.")
    tm = TextMap(document_atoms([p._p]))

    span = tm.find("reading further")
    assert span is not None
    assert tm.slice(span) == "reading further"


def test_matches_text_inside_a_hyperlink() -> None:
    doc = Document()
    para = doc.add_paragraph()
    para.add_run("See ")
    link = etree.SubElement(para._p, qn("hyperlink"))
    run = etree.SubElement(link, qn("r"))
    etree.SubElement(run, qn("t")).text = "the registered plan"
    tm = TextMap(document_atoms([para._p]))

    span = tm.find("the registered plan")
    assert span is not None
    assert tm.slice(span) == "the registered plan"


# ── ambiguity ───────────────────────────────────────────────────────────


def test_a_repeated_target_is_refused_not_guessed() -> None:
    tm = _map("The result holds. The result holds.")
    assert tm.find("The result holds") is None
    assert len(tm.find_all("The result holds")) == 2


def test_absent_target_returns_none() -> None:
    assert _map("Some prose.").find("not in the document") is None


def test_find_all_is_bounded() -> None:
    tm = _map(" ".join(["repeat"] * 20))
    assert len(tm.find_all("repeat", limit=3)) == 3


# ── converter decoration ────────────────────────────────────────────────


def test_a_table_row_quote_matches_cell_paragraphs() -> None:
    """Markdown pipes never existed in Word; the cells are separate paragraphs."""
    tm = _map("Validation", "94%")
    span = tm.find("| Validation | 94% |")
    assert span is not None


def test_a_list_marker_quote_matches_the_list_text() -> None:
    tm = _map("First point in the list")
    assert tm.find("- First point in the list") is not None


# ── diagnostics ─────────────────────────────────────────────────────────


def test_closest_reports_a_near_miss() -> None:
    tm = _map("The evidence layer earns its place.")
    score, snippet = tm.closest("The evidence layer earns its keep")
    assert 0.5 < score < 1.0
    assert "evidence layer" in snippet


def test_closest_on_an_empty_map_is_silent() -> None:
    assert TextMap([]).closest("anything") == (0.0, "")


# ── ported from the DocIndex suite this replaces ────────────────────────


def test_a_match_may_span_two_runs_of_one_paragraph() -> None:
    from andamentum.docx_edit.atoms import paragraph_atoms

    doc = Document()
    para = doc.add_paragraph()
    for chunk in ("the methods were ", "significantly ", "robust and clear"):
        para.add_run(chunk)
    tm = TextMap(list(paragraph_atoms(para._p)))

    span = tm.find("significantly robust")
    assert span is not None
    assert tm.slice(span) == "significantly robust"
    # The two ends really are in different runs.
    assert tm.atoms[span.start].run is not tm.atoms[span.end - 1].run


def test_a_paragraph_boundary_does_not_fuse_the_words_either_side() -> None:
    """Without a separator, "introduction" + "studying" would match as one."""
    tm = _map("introduction", "studying")
    assert tm.find("introductionstudying") is None
    assert tm.find("introduction studying") is not None


def test_closest_scores_the_local_window_not_the_whole_document() -> None:
    """A long target almost entirely present must score high even in a big
    document — a global ratio scored about 0.01 for exactly this case."""
    body = (
        "the widget catalogue is released under a permissive licence an "
        "archived release is deposited in a public repository the evaluation "
        "corpus comprising benchmark snapshots for all forty scenarios"
    )
    tm = _map("lorem ipsum " * 200 + body + " dolor sit " * 200)
    target = (
        "under a permissive licence; an archived release is deposited in "
        "[a public repository]. The evaluation corpus, comprising benchmark "
        "snapshots for all forty scenarios"
    )
    score, snippet = tm.closest(target)
    assert score > 0.8, f"local score should be high, got {score}"
    assert "public repository" in snippet
