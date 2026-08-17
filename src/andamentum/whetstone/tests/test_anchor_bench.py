"""Regression gates over the deterministic anchoring benchmark.

The benchmark itself lives in :mod:`.anchor_bench` and prints a per-construct
table; this file turns the two properties that must not regress into test
failures. No LLM, no network — the corpus is generated and the ground truth
is known by construction.
"""

from __future__ import annotations

from pathlib import Path

import pytest

from andamentum.whetstone.docx.low_level import DocxEditor, visible_text
from andamentum.whetstone.tests.anchor_bench import run_benchmark

# Headroom below the measured 96.3% so ordinary drift doesn't fail the
# suite, but a real regression does. Raise it when the rate rises.
_MIN_CORRECT_RATE = 90.0


@pytest.fixture(scope="module")
def results(tmp_path_factory: pytest.TempPathFactory) -> dict:
    import asyncio

    directory = tmp_path_factory.mktemp("anchor_bench")
    return asyncio.run(run_benchmark(directory))


def test_no_quote_is_ever_anchored_to_the_wrong_place(results: dict) -> None:
    """The invariant that outranks every other number here.

    A comment on the wrong sentence is worse than no comment: the reader
    cannot tell it has moved, so it quietly discredits the annotations that
    ARE right. Anchoring used to return the first of several matches, which
    put a quote from the second occurrence of a repeated sentence onto the
    first.
    """
    offenders = {
        name: outcome.wrong_detail for name, outcome in results.items() if outcome.wrong
    }
    assert not offenders, f"wrong anchors: {offenders}"


def test_correct_anchor_rate_holds(results: dict) -> None:
    total_correct = sum(o.correct for o in results.values())
    total = sum(o.total for o in results.values())
    rate = 100.0 * total_correct / total
    assert rate >= _MIN_CORRECT_RATE, f"correct rate {rate:.1f}% below floor"


@pytest.mark.parametrize(
    "construct",
    ["plain", "tabs", "tables", "lists", "runs", "headings", "typography"],
)
def test_construct_anchors_completely(results: dict, construct: str) -> None:
    """Each of these was measured at 100% and should stay there.

    ``repeats`` is deliberately excluded: its quotes are genuinely
    ambiguous and refusing them is the correct answer, not a shortfall.
    """
    outcome = results[construct]
    assert outcome.correct == outcome.total, (
        f"{construct}: {outcome.correct}/{outcome.total} correct; "
        f"missed={outcome.missed_detail}"
    )


def test_ambiguous_repeats_are_refused_not_guessed(results: dict) -> None:
    """The repeated sentence must produce a refusal, never a wrong anchor."""
    outcome = results["repeats"]
    assert outcome.wrong == 0
    assert outcome.missed >= 1, "expected the ambiguous quote to be refused"


# ── the document-corruption regression, independent of the benchmark ────


def test_run_normalisation_preserves_tabs_and_breaks(tmp_path: Path) -> None:
    """Opening and writing a .docx must not delete the author's tabs.

    ``_normalize_runs`` splits runs so each holds one token. It used to
    rebuild them from ``<w:t>`` text and drop the original, discarding every
    sibling element — so a document went in with a tab and came out without
    one. Anything else living in a multi-token run (drawings, footnote
    references, field plumbing) went the same way. This is data loss in the
    user's own manuscript, so it is pinned separately from the anchoring
    numbers it also caused.
    """
    import zipfile

    import docx

    source = tmp_path / "in.docx"
    document = docx.Document()
    document.add_paragraph("Left column\tRight column with several more words")
    paragraph = document.add_paragraph("Line one of two")
    paragraph.add_run().add_break()
    paragraph.add_run("Line two with several more words following it")
    document.save(str(source))

    def counts(path: Path) -> tuple[int, int]:
        xml = zipfile.ZipFile(path).read("word/document.xml").decode()
        return xml.count("<w:tab/>"), xml.count("<w:br/>")

    out = tmp_path / "out.docx"
    DocxEditor(str(source)).write(str(out))

    assert counts(out) == counts(source) == (1, 1)


def test_visible_text_reads_tabs_as_the_space_they_are(tmp_path: Path) -> None:
    """A tab is whitespace to a reader, so it must be whitespace to a match.

    Joining only ``<w:t>`` yields "Accuracy was94%", which no quote taken
    from the converted markdown ("Accuracy was 94%") can ever match — and
    no normalisation can fix, because collapsing whitespace cannot invent a
    space that is not in the string.
    """
    import docx

    source = tmp_path / "tabbed.docx"
    document = docx.Document()
    document.add_paragraph("Accuracy was\tninety-four percent overall")
    document.save(str(source))

    paragraph = DocxEditor(str(source)).paragraphs[0]
    assert "was ninety-four" in paragraph.text
    assert visible_text(paragraph.p_elem) == paragraph.text
