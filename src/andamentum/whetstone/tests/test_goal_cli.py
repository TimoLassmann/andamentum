"""``--goal`` / ``--goal-check`` / ``--max-rounds`` flag validation, plus two
end-to-end checks that the CLI dispatches into the goal loop correctly.

Most of this file mirrors ``test_reviewer_cli.py``'s pattern: parse with the
real parser, then call ``_validate_args`` directly and check the exit code
and message — the argument-time refusals, with no review run at all. The
last two tests mirror ``test_cli_review_wiring.py`` instead: they stub
``review_document_v3`` at the loop's own import site (not the pydantic-ai
agents underneath it — that round-by-round behaviour is
``test_goal_loop.py``'s job) and check that ``_run`` writes real output and
exits with the right code.
"""

from __future__ import annotations

import pytest
from docx import Document
from rich.console import Console

from andamentum.whetstone import cli as cli_mod
from andamentum.whetstone.schemas import Edit, Finding, ReviewResult

BASE = ["draft.md", "--model", "stub:none", "--out", "review.md"]


def _args(argv: list[str]):
    return cli_mod._build_parser().parse_args(argv)


def _validate(argv: list[str]) -> None:
    cli_mod._validate_args(_args(argv))


def test_goal_defaults_to_none() -> None:
    assert _args(BASE).goal is None


def test_goal_check_defaults_to_judged() -> None:
    assert _args(BASE).goal_check == "judged"


def test_max_rounds_defaults_to_four() -> None:
    assert _args(BASE).max_rounds == 4


def test_a_plain_review_with_no_goal_is_accepted() -> None:
    _validate(BASE)


def test_a_goal_with_the_default_reviewer_is_accepted() -> None:
    _validate([*BASE, "--goal", "the flow is correct"])


def test_goal_check_without_goal_is_refused(
    capsys: pytest.CaptureFixture[str],
) -> None:
    with pytest.raises(SystemExit) as exc:
        _validate([*BASE, "--goal-check", "structure"])
    assert exc.value.code == 1
    assert "--goal-check has no effect without --goal" in capsys.readouterr().err


def test_max_rounds_without_goal_is_refused(
    capsys: pytest.CaptureFixture[str],
) -> None:
    with pytest.raises(SystemExit) as exc:
        _validate([*BASE, "--max-rounds", "2"])
    assert exc.value.code == 1
    assert "--max-rounds has no effect without --goal" in capsys.readouterr().err


def test_max_rounds_at_its_own_default_is_tolerated_without_goal() -> None:
    _validate([*BASE, "--max-rounds", "4"])


def test_goal_under_checklist_is_refused(
    capsys: pytest.CaptureFixture[str],
) -> None:
    with pytest.raises(SystemExit) as exc:
        _validate([*BASE, "--mode", "checklist", "--goal", "the flow is correct"])
    assert exc.value.code == 1
    assert "--goal only applies to review mode" in capsys.readouterr().err


def test_goal_under_panel_is_refused(capsys: pytest.CaptureFixture[str]) -> None:
    with pytest.raises(SystemExit) as exc:
        _validate(
            [
                *BASE,
                "--mode",
                "panel",
                "--i-am-the-author",
                "--goal",
                "the flow is correct",
            ]
        )
    assert exc.value.code == 1
    assert "--goal only applies to review mode" in capsys.readouterr().err


def test_goal_with_a_reviewer_list_is_refused(
    capsys: pytest.CaptureFixture[str],
) -> None:
    with pytest.raises(SystemExit) as exc:
        _validate(
            [*BASE, "--reviewer", "claims,methods", "--goal", "the flow is correct"]
        )
    assert exc.value.code == 1
    assert "--goal takes one reviewer, not a list" in capsys.readouterr().err


def test_max_rounds_below_one_is_refused(capsys: pytest.CaptureFixture[str]) -> None:
    with pytest.raises(SystemExit) as exc:
        _validate([*BASE, "--goal", "the flow is correct", "--max-rounds", "0"])
    assert exc.value.code == 1
    assert "--max-rounds must be at least 1" in capsys.readouterr().err


def test_an_unknown_goal_check_is_refused_by_argparse() -> None:
    with pytest.raises(SystemExit) as exc:
        _args([*BASE, "--goal-check", "vibes"])
    assert exc.value.code == 2


# ── --goal-figure / --goal-figure-model ──────────────────────────────────


def test_goal_figure_without_goal_check_figures_is_refused(
    tmp_path, capsys: pytest.CaptureFixture[str]
) -> None:
    figure = tmp_path / "fig1.png"
    figure.write_bytes(b"")
    with pytest.raises(SystemExit) as exc:
        _validate(
            [
                *BASE,
                "--goal",
                "the flow is correct",
                "--goal-figure",
                f"Figure 1={figure}",
            ]
        )
    assert exc.value.code == 1
    assert (
        "--goal-figure only applies to --goal-check figures" in capsys.readouterr().err
    )


def test_goal_figure_model_without_goal_check_figures_is_refused(
    capsys: pytest.CaptureFixture[str],
) -> None:
    with pytest.raises(SystemExit) as exc:
        _validate(
            [
                *BASE,
                "--goal",
                "the flow is correct",
                "--goal-figure-model",
                "ollama:qwen3.8",
            ]
        )
    assert exc.value.code == 1
    assert "--goal-figure-model only applies" in capsys.readouterr().err


def test_goal_check_figures_needs_at_least_one_goal_figure(
    capsys: pytest.CaptureFixture[str],
) -> None:
    with pytest.raises(SystemExit) as exc:
        _validate(
            [
                *BASE,
                "--goal",
                "figures match",
                "--goal-check",
                "figures",
                "--goal-figure-model",
                "ollama:qwen3.8",
            ]
        )
    assert exc.value.code == 1
    assert "needs at least one --goal-figure" in capsys.readouterr().err


def test_goal_check_figures_needs_a_figure_model(
    tmp_path, capsys: pytest.CaptureFixture[str]
) -> None:
    figure = tmp_path / "fig1.png"
    figure.write_bytes(b"")
    with pytest.raises(SystemExit) as exc:
        _validate(
            [
                *BASE,
                "--goal",
                "figures match",
                "--goal-check",
                "figures",
                "--goal-figure",
                f"Figure 1={figure}",
            ]
        )
    assert exc.value.code == 1
    assert "needs --goal-figure-model" in capsys.readouterr().err


def test_goal_figure_with_a_missing_file_is_refused(
    capsys: pytest.CaptureFixture[str],
) -> None:
    with pytest.raises(SystemExit) as exc:
        _validate(
            [
                *BASE,
                "--goal",
                "figures match",
                "--goal-check",
                "figures",
                "--goal-figure",
                "Figure 1=does/not/exist.png",
                "--goal-figure-model",
                "ollama:qwen3.8",
            ]
        )
    assert exc.value.code == 2
    assert "no such file" in capsys.readouterr().err


def test_goal_figure_needing_label_and_equals_is_refused(
    capsys: pytest.CaptureFixture[str],
) -> None:
    with pytest.raises(SystemExit) as exc:
        _validate(
            [
                *BASE,
                "--goal",
                "figures match",
                "--goal-check",
                "figures",
                "--goal-figure",
                "not-a-label-path-pair",
                "--goal-figure-model",
                "ollama:qwen3.8",
            ]
        )
    assert exc.value.code == 2
    assert "is not LABEL=PATH" in capsys.readouterr().err


def test_a_valid_goal_check_figures_run_is_accepted(tmp_path) -> None:
    figure = tmp_path / "fig1.png"
    figure.write_bytes(b"")
    args = _args(
        [
            *BASE,
            "--goal",
            "figures match",
            "--goal-check",
            "figures",
            "--goal-figure",
            f"Figure 1={figure}",
            "--goal-figure-model",
            "ollama:qwen3.8",
        ]
    )
    cli_mod._validate_args(args)
    assert args.goal_figures == [("Figure 1", figure)]


# ── end-to-end: the CLI actually drives the loop ─────────────────────────


def _goal_draft(path) -> None:
    document = Document()
    document.add_heading("Methods", level=1)
    document.add_paragraph("The evidence layer earns its place in the analysis.")
    document.save(str(path))


def _goal_review_stub(*, rounds: int):
    """Round 0 has a major finding + an edit; every later round has neither
    — a structure goal met from round 1 onward, confirmed on round 2."""
    calls = {"n": 0}

    async def fake_review(source, **_kwargs):
        idx = min(calls["n"], rounds - 1)
        calls["n"] += 1
        text = source.text
        if idx == 0:
            start = text.index("earns its place")
            return ReviewResult(
                summary="stub",
                findings=[
                    Finding(
                        title="vague",
                        severity="major",
                        confidence="medium",
                        rationale="x",
                        quotes=[],
                    )
                ],
                edits=[
                    Edit(
                        title="t",
                        rationale="r",
                        section_id="s1",
                        char_start=start,
                        char_end=start + len("earns its place"),
                        original_text="earns its place",
                        new_text="earns its keep",
                    )
                ],
            )
        return ReviewResult(summary="stub", findings=[], edits=[])

    return calls, fake_review


async def test_a_goal_met_run_writes_output_and_exits_zero(
    tmp_path, monkeypatch: pytest.MonkeyPatch
) -> None:
    source = tmp_path / "draft.docx"
    _goal_draft(source)
    out = tmp_path / "reviewed.docx"

    _calls, fake_review = _goal_review_stub(rounds=3)
    monkeypatch.setattr(
        "andamentum.whetstone.goal.loop.review_document_v3", fake_review, raising=True
    )

    await cli_mod._run(
        _args(
            [
                str(source),
                "--model",
                "stub:none",
                "--out",
                str(out),
                "--goal",
                "the flow is correct",
                "--goal-check",
                "structure",
            ]
        ),
        Console(stderr=True),
    )

    assert out.exists()
    # round 0 fails; round 1 meets with nothing left to edit, so round 2
    # re-reviews the same text and confirms it.
    assert _calls["n"] == 3


async def test_a_goal_not_met_run_still_writes_output_and_exits_four(
    tmp_path, monkeypatch: pytest.MonkeyPatch
) -> None:
    source = tmp_path / "draft.docx"
    _goal_draft(source)
    out = tmp_path / "reviewed.docx"

    async def always_failing(source, **_kwargs):
        return ReviewResult(
            summary="stub",
            findings=[
                Finding(
                    title="vague",
                    severity="major",
                    confidence="medium",
                    rationale="x",
                    quotes=[],
                )
            ],
            edits=[],  # no edits at all -> no_progress on round 1
        )

    monkeypatch.setattr(
        "andamentum.whetstone.goal.loop.review_document_v3",
        always_failing,
        raising=True,
    )

    with pytest.raises(SystemExit) as exc:
        await cli_mod._run(
            _args(
                [
                    str(source),
                    "--model",
                    "stub:none",
                    "--out",
                    str(out),
                    "--goal",
                    "the flow is correct",
                    "--goal-check",
                    "structure",
                ]
            ),
            Console(stderr=True),
        )
    assert exc.value.code == 4
    assert out.exists(), "the output is written even though the goal was not met"
