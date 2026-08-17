"""Editing by cutting only the covered runs.

The invariant under test throughout: an edit changes the text it was asked to
change, and nothing else in the document moves.
"""

from __future__ import annotations

from docx import Document
from lxml import etree  # type: ignore

from andamentum.whetstone.docx.atoms import document_atoms, projection, qn
from andamentum.whetstone.docx.surgery import Edit, apply_edits
from andamentum.whetstone.docx.textmap import TextMap


def _doc(*paragraph_texts: str):
    doc = Document()
    for text in paragraph_texts:
        doc.add_paragraph(text)
    return doc, [p._p for p in doc.paragraphs]


def _accepted(paragraphs) -> str:
    """What the document says once the tracked changes are accepted."""
    return projection(document_atoms(_live(paragraphs)))


def _live(paragraphs):
    """Re-read the paragraph list from the body — splits add new ones."""
    body = paragraphs[0].getparent()
    return body.findall(qn("p"))


def _count(paragraphs, tag: str) -> int:
    body = paragraphs[0].getparent()
    return len(body.findall(f".//{qn(tag)}"))


def _deleted_text(paragraphs) -> str:
    body = paragraphs[0].getparent()
    return "".join(t.text or "" for t in body.iter(qn("delText")))


def _edit(textmap, target: str, new_text: str, author: str = "Reviewer") -> Edit:
    span = textmap.find(target)
    assert span is not None, f"target did not resolve: {target!r}"
    return Edit(span=span, new_text=new_text, author=author)


# ── narrow marks ────────────────────────────────────────────────────────


def test_only_the_changed_words_are_marked() -> None:
    """The span-scoped diff keeps the mark on the word, not the whole quote."""
    _, paragraphs = _doc("The evidence layer earns its place today.")
    tm = TextMap(document_atoms(paragraphs))
    apply_edits(tm, [_edit(tm, "earns its place", "earns its keep")])

    assert _accepted(paragraphs) == "The evidence layer earns its keep today."
    assert _deleted_text(paragraphs) == "place"


def test_a_replacement_spanning_several_runs() -> None:
    _, paragraphs = _doc("")
    p = paragraphs[0]
    for chunk in ("The data ", "shows ", "a clear ", "trend."):
        run = etree.SubElement(p, qn("r"))
        etree.SubElement(run, qn("t")).text = chunk
    tm = TextMap(document_atoms(paragraphs))
    apply_edits(tm, [_edit(tm, "shows a clear trend", "show a marked trend")])

    assert _accepted(paragraphs) == "The data show a marked trend."


def test_pure_insertion_adds_without_deleting() -> None:
    _, paragraphs = _doc("The result holds.")
    tm = TextMap(document_atoms(paragraphs))
    apply_edits(tm, [_edit(tm, "The result", "The main result")])

    assert _accepted(paragraphs) == "The main result holds."
    assert _deleted_text(paragraphs) == ""


def test_pure_deletion_removes_without_inserting() -> None:
    _, paragraphs = _doc("The very clear result holds.")
    tm = TextMap(document_atoms(paragraphs))
    apply_edits(tm, [_edit(tm, "The very clear result", "The result")])

    assert _accepted(paragraphs) == "The result holds."
    assert "very" in _deleted_text(paragraphs)


# ── what must survive ───────────────────────────────────────────────────


def test_a_line_break_outside_the_change_survives() -> None:
    """The old rebuild took this paragraph's break count from 1 to 0."""
    doc = Document()
    para = doc.add_paragraph()
    para.add_run("The evidence layer earns its place")
    para.add_run().add_break()
    para.add_run("and the deliberation layer does not.")
    paragraphs = [para._p]

    tm = TextMap(document_atoms(paragraphs))
    apply_edits(tm, [_edit(tm, "earns its place", "earns its keep")])

    assert _count(paragraphs, "br") == 1
    assert _accepted(paragraphs) == (
        "The evidence layer earns its keep and the deliberation layer does not."
    )


def test_a_tab_outside_the_change_survives() -> None:
    doc = Document()
    para = doc.add_paragraph()
    para.add_run("Column one")
    para.add_run().add_tab()
    para.add_run("Column two follows.")
    paragraphs = [para._p]

    tm = TextMap(document_atoms(paragraphs))
    apply_edits(tm, [_edit(tm, "Column two follows.", "Column two comes next.")])

    assert _count(paragraphs, "tab") == 1
    assert _accepted(paragraphs) == "Column one Column two comes next."


def test_formatting_is_preserved_on_both_sides_of_a_cut() -> None:
    doc = Document()
    para = doc.add_paragraph()
    run = para.add_run("emphasised text here")
    run.bold = True
    paragraphs = [para._p]

    tm = TextMap(document_atoms(paragraphs))
    apply_edits(tm, [_edit(tm, "text", "prose")])

    assert _accepted(paragraphs) == "emphasised prose here"
    # Every surviving run of that paragraph still carries the bold property.
    runs = paragraphs[0].findall(f".//{qn('r')}")
    assert runs and all(r.find(f"{qn('rPr')}/{qn('b')}") is not None for r in runs)


def test_an_earlier_authors_tracked_change_survives() -> None:
    """The old rebuild silently accepted a co-author's revisions and erased
    their name whenever it touched the same paragraph."""
    doc = Document()
    para = doc.add_paragraph()
    para.add_run("Our results are ")
    existing = etree.SubElement(para._p, qn("ins"))
    existing.set(qn("id"), "1")
    existing.set(qn("author"), "Dr Human")
    existing.set(qn("date"), "2026-01-01T00:00:00")
    run = etree.SubElement(existing, qn("r"))
    etree.SubElement(run, qn("t")).text = "modestly "
    para.add_run("encouraging overall.")
    paragraphs = [para._p]

    tm = TextMap(document_atoms(paragraphs))
    apply_edits(tm, [_edit(tm, "encouraging overall", "promising overall", "AI")])

    authors = {
        e.get(qn("author"))
        for e in paragraphs[0].iter()
        if e.tag in (qn("ins"), qn("del"))
    }
    assert "Dr Human" in authors
    assert "AI" in authors
    assert _accepted(paragraphs) == "Our results are modestly promising overall."


def test_a_hyperlink_outside_the_change_survives() -> None:
    doc = Document()
    para = doc.add_paragraph()
    para.add_run("See ")
    link = etree.SubElement(para._p, qn("hyperlink"))
    run = etree.SubElement(link, qn("r"))
    etree.SubElement(run, qn("t")).text = "the plan"
    para.add_run(" before starting.")
    paragraphs = [para._p]

    tm = TextMap(document_atoms(paragraphs))
    apply_edits(tm, [_edit(tm, "before starting", "before you start")])

    assert _count(paragraphs, "hyperlink") == 1
    assert _accepted(paragraphs) == "See the plan before you start."


# ── paragraph boundaries are ordinary characters ────────────────────────


def test_an_edit_across_a_paragraph_break_merges_the_paragraphs() -> None:
    _, paragraphs = _doc("The first thought ends here", "and the second begins.")
    tm = TextMap(document_atoms(paragraphs))
    apply_edits(tm, [_edit(tm, "ends here\nand the second begins", "ends here and continues")])

    # The first paragraph's mark is marked deleted: Word merges on accept.
    ppr = paragraphs[0].find(qn("pPr"))
    assert ppr is not None
    assert ppr.find(f"{qn('rPr')}/{qn('del')}") is not None


def test_a_newline_in_the_replacement_creates_a_paragraph() -> None:
    _, paragraphs = _doc("One thought. Another thought entirely.")
    tm = TextMap(document_atoms(paragraphs))
    report = apply_edits(
        tm, [_edit(tm, "One thought.", "One thought.\nA new paragraph starts.")]
    )

    assert report.paragraphs_created == 1
    live = _live(paragraphs)
    assert len(live) == 2
    # The first paragraph's mark is marked inserted: Word's "break added here".
    ppr = live[0].find(qn("pPr"))
    assert ppr is not None
    assert ppr.find(f"{qn('rPr')}/{qn('ins')}") is not None
    assert "A new paragraph starts." in _accepted(paragraphs)


# ── conflicts ───────────────────────────────────────────────────────────


def test_overlapping_edits_are_refused_not_composed() -> None:
    _, paragraphs = _doc("We found that the results show a significant effect here.")
    tm = TextMap(document_atoms(paragraphs))
    report = apply_edits(
        tm,
        [
            _edit(tm, "the results show a significant effect", "the results suggest an effect"),
            _edit(tm, "a significant effect here", "a large effect here"),
        ],
    )

    assert report.applied == 1
    assert len(report.conflicted) == 1
    text = _accepted(paragraphs)
    assert text == "We found that the results suggest an effect here."


def test_disjoint_edits_in_one_paragraph_both_apply() -> None:
    _, paragraphs = _doc("We found that the effect was clear in every cohort.")
    tm = TextMap(document_atoms(paragraphs))
    report = apply_edits(
        tm,
        [
            _edit(tm, "We found", "We observed"),
            _edit(tm, "every cohort", "each cohort"),
        ],
    )

    assert report.applied == 2
    assert _accepted(paragraphs) == (
        "We observed that the effect was clear in each cohort."
    )


# ── repeated rounds ─────────────────────────────────────────────────────


def test_a_second_round_edits_what_the_first_round_inserted() -> None:
    _, paragraphs = _doc("The evidence layer earns its place.")

    first = TextMap(document_atoms(paragraphs))
    apply_edits(first, [_edit(first, "earns its place", "earns its keep", "Reviewer A")])

    second = TextMap(document_atoms(_live(paragraphs)))
    assert second.text == "The evidence layer earns its keep."
    apply_edits(
        second,
        [_edit(second, "earns its keep", "earns its salt", "Reviewer B")],
        first_id=5000,
    )

    assert _accepted(paragraphs) == "The evidence layer earns its salt."
    authors = {
        e.get(qn("author"))
        for e in paragraphs[0].iter()
        if e.tag in (qn("ins"), qn("del"))
    }
    assert authors == {"Reviewer A", "Reviewer B"}
