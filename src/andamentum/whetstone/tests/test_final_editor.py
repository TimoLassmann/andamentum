"""One editor pass after a list of reviewers.

A list of reviewers does not edit round by round. Each reviewer reads the same
unchanged text and raises findings; the line editor then runs ONCE over all of
them. Per-round editing gave five reviewers a style quota over one abstract,
each rewriting the last one's insertions, and the page that came back was
tracked changes with no paper underneath.

Two spellings ask for the pass and mean the same thing: ``--reviewer
claims,methods --editor`` and ``--reviewer claims,methods,editor``. A single
reviewer keeps its in-graph editor, which already sees all of its findings.
"""

from __future__ import annotations

from pathlib import Path
from typing import Any

import pytest
from docx import Document
from rich.console import Console

from andamentum.whetstone import cli as cli_mod
from andamentum.whetstone.schemas import Edit, Finding, Quote, ReviewResult
from andamentum.whetstone.v3 import editor as editor_mod
from andamentum.whetstone.v3.editor import Declined, _to_editor_finding, edit_findings
from andamentum.whetstone.v3.model import Section
from andamentum.whetstone.v3.reviewers import EDITOR_PROFILE, PROFILES, author_for

BASE = ["draft.md", "--model", "stub:none", "--out", "review.md"]


def _args(argv: list[str]):
    args = cli_mod._build_parser().parse_args(argv)
    cli_mod._validate_args(args)
    return args


# ── the command line ────────────────────────────────────────────────────────


def test_a_list_with_the_editor_flag_ends_in_one_editor_pass() -> None:
    args = _args([*BASE, "--reviewer", "claims,methods", "--editor"])
    assert args.reviewers == ["claims", "methods"]
    assert args.final_editor is True


def test_the_editor_named_in_a_list_is_the_final_pass_not_a_round() -> None:
    """As a round it would run claim extraction, synthesis and self-critique
    to produce nothing; its one job is the pass at the end."""
    args = _args([*BASE, "--reviewer", "claims,methods,editor"])
    assert args.reviewers == ["claims", "methods"]
    assert args.final_editor is True


def test_the_two_spellings_agree() -> None:
    flagged = _args([*BASE, "--reviewer", "claims,methods", "--editor"])
    listed = _args([*BASE, "--reviewer", "claims,methods,editor"])
    assert (flagged.reviewers, flagged.final_editor) == (
        listed.reviewers,
        listed.final_editor,
    )


def test_the_editor_is_the_final_pass_wherever_it_is_named() -> None:
    """Its position in the list is not where it runs."""
    args = _args([*BASE, "--reviewer", "editor,claims"])
    assert args.reviewers == ["claims"]
    assert args.final_editor is True


def test_a_list_without_the_editor_has_no_final_pass() -> None:
    args = _args([*BASE, "--reviewer", "claims,methods"])
    assert args.reviewers == ["claims", "methods"]
    assert args.final_editor is False


def test_a_single_reviewer_keeps_its_own_editor() -> None:
    """One set of findings; the in-graph pass already sees all of it."""
    args = _args([*BASE, "--reviewer", "claims", "--editor"])
    assert args.reviewers == ["claims"]
    assert args.final_editor is False
    alone = _args([*BASE, "--reviewer", "editor", "--editor"])
    assert alone.reviewers == ["editor"]
    assert alone.final_editor is False


def test_the_line_editor_signs_the_final_render() -> None:
    """The rounds sign their own comments on the way through; the pass at the
    end is the one still unwritten, and it is the editor's."""
    args = _args([*BASE, "--reviewer", "claims,methods", "--editor"])
    assert cli_mod._resolve_author(args) == author_for(profile=EDITOR_PROFILE)
    plain = _args([*BASE, "--reviewer", "claims,methods"])
    assert cli_mod._resolve_author(plain) == author_for(profile=PROFILES["methods"])


def test_checklist_still_reads_the_list_as_typed(
    capsys: pytest.CaptureFixture[str],
) -> None:
    """Checklist refuses a list and refuses the line editor; splitting the
    editor off first would have let ``claims,editor`` through as ``claims``."""
    with pytest.raises(SystemExit):
        _args([*BASE, "--mode", "checklist", "--reviewer", "claims,editor"])
    assert "line editor runs no checks" in capsys.readouterr().err
    with pytest.raises(SystemExit):
        _args([*BASE, "--mode", "checklist", "--reviewer", "claims,methods"])
    assert "one reviewer" in capsys.readouterr().err


# ── the conversion the pass runs on ─────────────────────────────────────────


def _finding(**kw: Any) -> Finding:
    base: dict[str, Any] = dict(
        title="Evidence support",
        severity="moderate",
        confidence="medium",
        rationale="The verb outruns the data.",
        quotes=[Quote(section_id="s1", char_start=10, char_end=22, text="proves that")],
        perspective="Claims reviewer",
    )
    base.update(kw)
    return Finding(**base)


def test_a_rendered_finding_maps_back_to_what_the_editor_reads() -> None:
    internal = _to_editor_finding(_finding())
    assert internal is not None
    assert internal.criterion == "Evidence support"
    assert internal.severity == "moderate"
    assert internal.span is not None
    assert (internal.span.section_id, internal.span.start, internal.span.end) == (
        "s1",
        10,
        22,
    )


def test_the_reviewer_is_written_into_the_issue() -> None:
    """Several reviewers' findings share one list; 'two said this' and 'one
    said this twice' must not look alike."""
    internal = _to_editor_finding(_finding())
    assert internal is not None
    assert internal.issue == "[Claims reviewer] The verb outruns the data."
    untagged = _to_editor_finding(_finding(perspective=None))
    assert untagged is not None
    assert untagged.issue == "The verb outruns the data."


def test_a_finding_with_no_quote_is_not_sent() -> None:
    assert _to_editor_finding(_finding(quotes=[])) is None


async def test_one_call_over_all_findings_carries_the_disagreements(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    seen: dict[str, Any] = {}

    async def fake_run_editor(
        sections: Any, **kwargs: Any
    ) -> tuple[list[Edit], list[Declined]]:
        seen["sections"] = sections
        seen.update(kwargs)
        return [], []

    monkeypatch.setattr(editor_mod, "run_editor", fake_run_editor)
    sections = [Section(id="s1", title="Intro", text="x" * 40, start=0, end=40)]
    await edit_findings(
        sections,
        [_finding(), _finding(perspective="Methods reviewer"), _finding(quotes=[])],
        model="stub:none",
        criteria=["clarity"],
        brief="Do not evaluate the argument.",
        disagreements=["whether the sample is large enough"],
    )
    assert seen["sections"] is sections
    assert seen["criteria"] == ["clarity"]
    assert seen["agent_model"] == "stub:none"
    assert len(seen["findings"]) == 2
    assert seen["brief"].startswith("Do not evaluate the argument.")
    assert "DISAGREE" in seen["brief"]
    assert "whether the sample is large enough" in seen["brief"]
    assert "can_fix=false" in seen["brief"]


async def test_no_disagreements_adds_nothing_to_the_brief(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    seen: dict[str, Any] = {}

    async def fake_run_editor(
        sections: Any, **kwargs: Any
    ) -> tuple[list[Edit], list[Declined]]:
        seen.update(kwargs)
        return [], []

    monkeypatch.setattr(editor_mod, "run_editor", fake_run_editor)
    sections = [Section(id="s1", title="Intro", text="x" * 40, start=0, end=40)]
    await edit_findings(
        sections, [], model="stub:none", criteria=["clarity"], brief="Brief."
    )
    assert seen["brief"] == "Brief."


# ── the run ─────────────────────────────────────────────────────────────────


def _draft(path: Path) -> None:
    document = Document()
    document.add_heading("Methods", level=1)
    document.add_paragraph("We met the goal. Then we stopped.")
    document.save(str(path))


async def test_the_rounds_do_not_edit_and_the_editor_runs_once_over_all_of_them(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    source = tmp_path / "draft.docx"
    _draft(source)
    rounds: list[dict[str, Any]] = []
    editor_calls: list[dict[str, Any]] = []

    async def fake_review(_source: Any, **kwargs: Any) -> ReviewResult:
        rounds.append(kwargs)
        name = kwargs["reviewer"].identity
        return ReviewResult(
            summary=f"{name} stub.",
            findings=[
                Finding(
                    title=name,
                    severity="moderate",
                    confidence="medium",
                    rationale=f"{name} objects.",
                    quotes=[
                        Quote(
                            section_id="s1",
                            char_start=8,
                            char_end=16,
                            text="met the goal",
                        )
                    ],
                )
            ],
        )

    async def fake_reconcile(reviews: Any, *, model: str) -> Any:
        from types import SimpleNamespace as N

        return N(
            verdict="Two views.",
            agreed=[],
            disagreements=[
                N(about="whether the goal was met", positions=["yes", "no"])
            ],
            raised_by_one=[],
        )

    async def fake_edit_findings(sections: Any, findings: Any, **kwargs: Any) -> Any:
        editor_calls.append({"sections": sections, "findings": findings, **kwargs})
        edit = Edit(
            title="Tighten",
            severity="minor",
            confidence="high",
            rationale="shorter",
            section_id="s1",
            char_start=8,
            char_end=16,
            original_text="met the goal",
            new_text="succeeded",
        )
        return [edit], [
            Declined(section_id="s1", criterion="c", issue="i", why="author's call")
        ]

    monkeypatch.setattr("andamentum.whetstone.v3.review_document_v3", fake_review)
    monkeypatch.setattr(
        "andamentum.whetstone.v3.reconcile.reconcile_reviewers", fake_reconcile
    )
    monkeypatch.setattr(
        "andamentum.whetstone.v3.editor.edit_findings", fake_edit_findings
    )

    rendered: dict[str, Any] = {}

    def fake_render_markdown(result: Any, out: Any, **kwargs: Any) -> None:
        rendered["result"] = result
        Path(out).write_text("stub")

    monkeypatch.setattr(
        "andamentum.whetstone.renderers.render_markdown", fake_render_markdown
    )

    args = _args(
        [
            str(source),
            "--model",
            "stub:none",
            "--out",
            str(tmp_path / "review.md"),
            "--reviewer",
            "claims,methods",
            "--editor",
        ]
    )
    await cli_mod._run(args, Console(stderr=True))

    # Two rounds, neither editing.
    assert [r["reviewer"].name for r in rounds] == ["claims", "methods"]
    assert all(r["editor"] is False for r in rounds)

    # Only round one flags the draft's novelty claims. The claims do not
    # change because a second reviewer read the draft, so a later round
    # spending the flag call would buy a list nobody keeps.
    assert [r["flag_novelty"] for r in rounds] == [True, False]

    # One editor call, with both reviewers' findings, tagged, and the
    # disagreement the reconciliation found.
    assert len(editor_calls) == 1
    call = editor_calls[0]
    assert [f.perspective for f in call["findings"]] == [
        "Claims reviewer",
        "Methods reviewer",
    ]
    assert call["disagreements"] == ["whether the goal was met"]
    assert call["criteria"] == list(EDITOR_PROFILE.editor_criteria or ())
    assert call["brief"] == EDITOR_PROFILE.brief
    assert call["model"] == "stub:none"

    # The rendered result carries the editor's edits and what it declined.
    result = rendered["result"]
    assert [e.new_text for e in result.edits] == ["succeeded"]
    assert [d.why for d in result.declined_findings] == ["author's call"]
    assert result.metrics.edits_count == 1


async def test_editor_criteria_typed_on_the_command_line_win(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    source = tmp_path / "draft.docx"
    _draft(source)
    editor_calls: list[dict[str, Any]] = []

    async def fake_review(_source: Any, **kwargs: Any) -> ReviewResult:
        return ReviewResult(summary="stub")

    async def fake_reconcile(reviews: Any, *, model: str) -> Any:
        from types import SimpleNamespace as N

        return N(verdict="", agreed=[], disagreements=[], raised_by_one=[])

    async def fake_edit_findings(sections: Any, findings: Any, **kwargs: Any) -> Any:
        editor_calls.append(kwargs)
        return [], []

    monkeypatch.setattr("andamentum.whetstone.v3.review_document_v3", fake_review)
    monkeypatch.setattr(
        "andamentum.whetstone.v3.reconcile.reconcile_reviewers", fake_reconcile
    )
    monkeypatch.setattr(
        "andamentum.whetstone.v3.editor.edit_findings", fake_edit_findings
    )
    monkeypatch.setattr(
        "andamentum.whetstone.renderers.render_markdown",
        lambda r, out, **kw: Path(out).write_text("stub"),
    )

    args = _args(
        [
            str(source),
            "--model",
            "stub:none",
            "--out",
            str(tmp_path / "review.md"),
            "--reviewer",
            "claims,methods,editor",
            "--editor-criteria",
            "voice, tone",
        ]
    )
    await cli_mod._run(args, Console(stderr=True))
    assert editor_calls[0]["criteria"] == ["voice", "tone"]
