"""Opening a source so its text and its .docx are the same thing.

Two input paths, one property: ``document.text`` is the projection of
``document.docx_path``, character for character. Everything downstream —
sectionising, quoting, annotating — inherits its coordinates from that, so it
is the property worth pinning on both paths rather than each consequence of it
separately.
"""

from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document

from andamentum.whetstone.docx.atoms import document_atoms, projection
from andamentum.whetstone.docx.reading import read_docx
from andamentum.whetstone.reviewed_document import ReviewedDocument, open_source


def _projection_of(path: Path) -> str:
    return projection(document_atoms(read_docx(path).paragraphs))


def _manuscript(path: Path) -> None:
    # Sibling H1s, not a Title with H1s under it: the size banding only splits
    # a section that exceeds its budget, so a short document nested under one
    # Title is legitimately one section — which would make this fixture prove
    # nothing about heading detection.
    document = Document()
    document.add_heading("Introduction", level=1)
    document.add_paragraph("We set out to measure the thing.")
    document.add_heading("Methods", level=1)
    paragraph = document.add_paragraph("We met the ")
    paragraph.add_run("aim")
    paragraph.add_run("3,4").font.superscript = True
    paragraph.add_run(". Then we stopped.")
    document.add_heading("Results", level=1)
    document.add_paragraph("Accuracy reached 94%.")
    document.save(str(path))


# ── the .docx path: provenance from reading ─────────────────────────────


async def test_a_docx_is_reviewed_as_its_own_projection(tmp_path: Path) -> None:
    source = tmp_path / "draft.docx"
    _manuscript(source)

    document = await open_source(str(source))
    try:
        assert document.docx_path == source
        assert document.docx_is_temp is False
        assert document.text == _projection_of(source)
        # The three shapes the converter mangles are simply absent: no
        # inserted spaces, no ``**`` around the emphasis, no ``|---|`` rows.
        assert "aim3,4." in document.text
        assert "**" not in document.text
    finally:
        document.cleanup()
    assert source.exists(), "an input the user owns must never be deleted"


async def test_a_docx_is_sectionised_by_its_heading_styles(tmp_path: Path) -> None:
    source = tmp_path / "draft.docx"
    _manuscript(source)

    document = await open_source(str(source))
    try:
        assert [s.title for s in document.sections] == [
            "Introduction",
            "Methods",
            "Results",
        ]
        for section in document.sections:
            assert section.text == document.text[section.start : section.end]
    finally:
        document.cleanup()


async def test_a_quote_from_a_docx_section_is_already_a_document_address(
    tmp_path: Path,
) -> None:
    """The payoff, stated as an equation rather than a hit rate."""
    source = tmp_path / "draft.docx"
    _manuscript(source)

    document = await open_source(str(source))
    try:
        methods = next(s for s in document.sections if s.title == "Methods")
        offset = methods.text.index("aim3,4.")
        start = methods.start + offset
        assert document.text[start : start + len("aim3,4.")] == "aim3,4."
    finally:
        document.cleanup()


# ── the markdown path: provenance from writing ──────────────────────────


_MARKDOWN = """# Introduction

We set out to measure the thing.

# Methods

Participants were recruited locally.

# Results

Accuracy reached 94%.
"""


async def test_markdown_is_reviewed_as_the_baseline_it_generated(
    tmp_path: Path,
) -> None:
    """Written first, reviewed second — so the offsets are valid here too.

    The pipeline used to harvest twice: once for the review and once, later,
    for the baseline. Nothing tied the two renderings together, and for an
    ambiguous HTML page harvest races two extractors, so they need not even
    have been the same text.
    """
    source = tmp_path / "draft.md"
    source.write_text(_MARKDOWN)

    document = await open_source(str(source))
    try:
        assert document.docx_is_temp is True
        assert document.docx_path.exists()
        assert document.text == _projection_of(document.docx_path)
        # Headings are carried in the paragraph style, so the reviewed text
        # holds the title and not the ``#`` that produced it.
        assert "# Introduction" not in document.text
        assert document.text.startswith("Introduction\n")
    finally:
        document.cleanup()
    assert not document.docx_path.exists()


async def test_the_generated_baseline_keeps_the_markdown_structure(
    tmp_path: Path,
) -> None:
    source = tmp_path / "draft.md"
    source.write_text(_MARKDOWN)

    document = await open_source(str(source))
    try:
        assert [s.title for s in document.sections] == [
            "Introduction",
            "Methods",
            "Results",
        ]
        for section in document.sections:
            assert section.text == document.text[section.start : section.end]
    finally:
        document.cleanup()


async def test_raw_markdown_content_is_opened_without_a_file(tmp_path: Path) -> None:
    document = await open_source(_MARKDOWN)
    try:
        assert "Participants were recruited locally." in document.text
        assert document.docx_is_temp is True
    finally:
        document.cleanup()


async def test_a_hash_that_is_not_a_heading_stays_body_text() -> None:
    document = await open_source("Use the #hashtag convention.\n\nAnd carry on.")
    try:
        assert document.text.startswith("Use the #hashtag convention.")
    finally:
        document.cleanup()


async def test_a_missing_path_raises_rather_than_being_reviewed() -> None:
    with pytest.raises(FileNotFoundError, match="not found"):
        await open_source("/tmp/whetstone-no-such-draft-xyz.md")


# ── the confidentiality gate sees what the body projection does not ─────


async def test_a_marker_in_a_running_header_still_reaches_the_tripwire(
    tmp_path: Path,
) -> None:
    """A header is not in the body, and the body is all the reviewer reads.

    The gate has to see it anyway: "CONFIDENTIAL — FOR PEER REVIEW ONLY" is a
    running header far more often than it is a paragraph of the draft, and a
    review that reads the body directly would otherwise walk straight past it.
    """
    source = tmp_path / "draft.docx"
    document = Document()
    document.sections[0].header.paragraphs[
        0
    ].text = "CONFIDENTIAL — Manuscript ID: BIO-2026-0042"
    document.add_heading("Methods", level=1)
    document.add_paragraph("We did the thing.")
    document.save(str(source))

    opened = await open_source(str(source))
    try:
        assert "Manuscript ID:" not in opened.text
        assert "Manuscript ID:" in opened.confidentiality_text

        from andamentum.whetstone._confidentiality import (
            ConfidentialityMarkerError,
            check_confidentiality,
        )

        check_confidentiality(opened.text)  # the body alone says nothing
        with pytest.raises(ConfidentialityMarkerError):
            check_confidentiality(opened.confidentiality_text)
    finally:
        opened.cleanup()


async def test_a_plain_document_scans_as_just_its_text(tmp_path: Path) -> None:
    source = tmp_path / "draft.docx"
    _manuscript(source)
    document = await open_source(str(source))
    try:
        assert document.confidentiality_text == document.text
    finally:
        document.cleanup()


# ── the warning a thin review would otherwise be ────────────────────────


async def test_a_long_document_with_no_heading_styles_says_so(
    tmp_path: Path,
) -> None:
    """Bold-and-bigger is not a heading, and the result is one 40-page unit.

    Legitimate input, but not what the author expects to have happened, and
    it shows up only as a suspiciously thin review. Say it out loud instead.
    """
    source = tmp_path / "flat.docx"
    document = Document()
    for _ in range(60):
        document.add_paragraph(
            "Sentence after sentence, with no structure at all. " * 8
        )
    document.save(str(source))

    opened = await open_source(str(source))
    try:
        assert len(opened.sections) == 1
        assert (
            opened.warnings
            and "no paragraphs styled as headings" in (opened.warnings[0])
        )
    finally:
        opened.cleanup()


async def test_a_short_document_with_one_section_is_not_flagged(
    tmp_path: Path,
) -> None:
    source = tmp_path / "note.docx"
    document = Document()
    document.add_paragraph("A short note with nothing else in it.")
    document.save(str(source))

    opened = await open_source(str(source))
    try:
        assert opened.warnings == []
    finally:
        opened.cleanup()


# ── cleanup ownership ───────────────────────────────────────────────────


def test_cleanup_is_a_no_op_for_a_document_the_caller_supplied(
    tmp_path: Path,
) -> None:
    source = tmp_path / "draft.docx"
    _manuscript(source)
    ReviewedDocument(
        text="x", sections=[], docx_path=source, docx_is_temp=False
    ).cleanup()
    assert source.exists()


# ── end to end: open, review, annotate ──────────────────────────────────


async def test_every_annotation_lands_on_a_document_of_hard_shapes(
    tmp_path: Path,
) -> None:
    """The regression that pins the 15-of-43.

    One document carrying all three shapes the converted-markdown route lost —
    a superscript citation with no space before it, a sentence broken across
    runs by emphasis, and a table (which Word holds as one paragraph per cell,
    and Docling as ``|---|---|`` rows). Every rewrite and every comment must
    land; ``failed`` is zero by construction, not by tuning.
    """
    from andamentum.whetstone.renderers.docx import render_docx
    from andamentum.whetstone.schemas import Edit, Finding, Quote, ReviewResult

    source = tmp_path / "hard.docx"
    document = Document()
    document.add_heading("Methods", level=1)
    paragraph = document.add_paragraph("We met the ")
    paragraph.add_run("aim")
    paragraph.add_run("3,4").font.superscript = True
    paragraph.add_run(". Then we stopped.")
    document.add_heading("Results", level=1)
    emphasis = document.add_paragraph("The effect was ")
    emphasis.add_run("large").bold = True
    emphasis.add_run(" and clear in every cohort.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Metric"
    table.cell(0, 1).text = "Value"
    table.cell(1, 0).text = "Validation"
    table.cell(1, 1).text = "94%"
    document.save(str(source))

    opened = await open_source(str(source))
    text = opened.text
    assert "aim3,4." in text

    def span(needle: str) -> tuple[int, int]:
        start = text.index(needle)
        return start, start + len(needle)

    cite = span("aim3,4.")
    emph = span("was large and clear")
    cell = span("94%")

    result = ReviewResult(
        summary="A review.",
        edits=[
            Edit(
                title="Space the citation",
                rationale="Readability.",
                section_id="s1",
                char_start=cite[0],
                char_end=cite[1],
                original_text=text[cite[0] : cite[1]],
                new_text="aim 3 and 4.",
            )
        ],
        findings=[
            Finding(
                title="Unquantified",
                severity="moderate",
                confidence="high",
                rationale="How large?",
                quotes=[
                    Quote(
                        section_id="s2",
                        char_start=emph[0],
                        char_end=emph[1],
                        text=text[emph[0] : emph[1]],
                    )
                ],
            ),
            Finding(
                title="No interval",
                severity="minor",
                confidence="medium",
                rationale="A point estimate with no interval.",
                quotes=[
                    Quote(
                        section_id="s2",
                        char_start=cell[0],
                        char_end=cell[1],
                        text="94%",
                    )
                ],
            ),
        ],
    )

    out = tmp_path / "reviewed.docx"
    try:
        report = render_docx(
            result, document=opened, output_path=out, visible_watermark=False
        )
    finally:
        opened.cleanup()

    assert report.applied_rewrites == 1
    assert report.applied_notes == 2
    assert report.shadowed_notes == []
    assert report.conflicted_rewrites == []
    assert out.exists()

    # The rewrite reads back with the change accepted, in place, and nothing
    # else in the document moved.
    after = _projection_of(out)
    assert "aim 3 and 4." in after
    assert after.endswith(text.replace("aim3,4.", "aim 3 and 4."))


async def test_rendering_against_a_document_that_was_not_reviewed_is_refused(
    tmp_path: Path,
) -> None:
    """No silent fallback to searching — that is the whole change."""
    from andamentum.whetstone.docx.annotate import CoordinateError
    from andamentum.whetstone.renderers.docx import render_docx
    from andamentum.whetstone.schemas import Finding, Quote, ReviewResult

    reviewed = tmp_path / "a.docx"
    _manuscript(reviewed)
    other = tmp_path / "b.docx"
    document = Document()
    document.add_paragraph("A completely different document.")
    document.save(str(other))

    opened = await open_source(str(reviewed))
    try:
        mismatched = ReviewedDocument(
            text=opened.text, sections=[], docx_path=other, docx_is_temp=False
        )
        result = ReviewResult(
            findings=[
                Finding(
                    title="x",
                    severity="minor",
                    confidence="low",
                    rationale="r",
                    quotes=[
                        Quote(section_id="s1", char_start=0, char_end=5, text="Intro")
                    ],
                )
            ]
        )
        with pytest.raises(CoordinateError, match="does not read as the text"):
            render_docx(
                result,
                document=mismatched,
                output_path=tmp_path / "out.docx",
                visible_watermark=False,
            )
    finally:
        opened.cleanup()
