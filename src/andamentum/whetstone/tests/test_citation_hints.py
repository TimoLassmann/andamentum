"""What the Word file says about citations that its text does not."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from lxml import etree  # type: ignore

from andamentum.whetstone.docx.citation_hints import citation_hints
from andamentum.whetstone.reviewed_document import open_source


async def _hints(path: Path):
    document = await open_source(str(path))
    return document.text, citation_hints(
        document.docx_path, expected_text=document.text
    )


async def test_superscript_runs_become_hints_at_their_offsets(tmp_path: Path) -> None:
    doc = Document()
    p = doc.add_paragraph()
    p.add_run("Diagnosis is slow")
    p.add_run("12").font.superscript = True
    p.add_run(". Area of 3 m")
    p.add_run("2").font.superscript = True
    path = tmp_path / "raised.docx"
    doc.save(str(path))

    text, hints = await _hints(path)
    assert [(h.kind, text[h.start : h.end]) for h in hints] == [
        ("superscript", "12"),
        ("superscript", "2"),
    ]


async def test_a_zotero_field_result_is_one_citation_hint(tmp_path: Path) -> None:
    """begin / instruction / separate / result / end, as Zotero writes it."""
    doc = Document()
    p = doc.add_paragraph()
    p.add_run("Rare diseases are common ")
    runs = [p.add_run() for _ in range(5)]
    runs[0]._r.append(_fld("begin"))
    instr = etree.SubElement(runs[1]._r, qn("w:instrText"))
    instr.text = ' ADDIN ZOTERO_ITEM CSL_CITATION {"citationItems":[]} '
    runs[2]._r.append(_fld("separate"))
    runs[3].text = "(Smith et al., 2020)"
    runs[4]._r.append(_fld("end"))
    p.add_run(".")
    path = tmp_path / "zotero.docx"
    doc.save(str(path))

    text, hints = await _hints(path)
    fields = [h for h in hints if h.kind == "field_citation"]
    assert [text[h.start : h.end] for h in fields] == ["(Smith et al., 2020)"]
    assert "ZOTERO_ITEM" in fields[0].payload


async def test_a_field_that_is_not_a_citation_is_ignored(tmp_path: Path) -> None:
    doc = Document()
    p = doc.add_paragraph()
    runs = [p.add_run() for _ in range(5)]
    runs[0]._r.append(_fld("begin"))
    instr = etree.SubElement(runs[1]._r, qn("w:instrText"))
    instr.text = " REF _Ref123 \\h "
    runs[2]._r.append(_fld("separate"))
    runs[3].text = "Figure 2"
    runs[4]._r.append(_fld("end"))
    path = tmp_path / "ref.docx"
    doc.save(str(path))

    _, hints = await _hints(path)
    assert hints == []


async def test_a_different_text_gets_no_hints(tmp_path: Path) -> None:
    doc = Document()
    doc.add_paragraph().add_run("x1").font.superscript = True
    path = tmp_path / "x.docx"
    doc.save(str(path))
    assert citation_hints(path, expected_text="something else") == []


def _fld(kind: str):
    element = etree.Element(qn("w:fldChar"))
    element.set(qn("w:fldCharType"), kind)
    return element


async def test_a_footnote_mark_is_a_note_hint_at_its_position(tmp_path: Path) -> None:
    """The mark has no character of its own; its hint sits where it falls in
    the text, and carries the note's text from footnotes.xml."""
    import re
    import zipfile
    from io import BytesIO

    doc = Document()
    doc.add_paragraph("Introduction")
    p = doc.add_paragraph()
    p.add_run("Rare diseases are common.")
    p.add_run("MARK")
    p.add_run(" More text.")
    path = tmp_path / "notes.docx"
    doc.save(str(path))

    w = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    out = BytesIO()
    with zipfile.ZipFile(path) as src, zipfile.ZipFile(out, "w") as dst:
        for item in src.infolist():
            body = src.read(item.filename)
            if item.filename == "word/document.xml":
                body = re.sub(
                    rb"<w:r><w:t>MARK</w:t></w:r>",
                    b'<w:r><w:footnoteReference w:id="1"/></w:r>',
                    body,
                )
            elif item.filename == "word/_rels/document.xml.rels":
                body = body.replace(
                    b"</Relationships>",
                    b'<Relationship Id="rIdFn" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/footnotes" Target="footnotes.xml"/></Relationships>',
                )
            dst.writestr(item, body)
        dst.writestr(
            "word/footnotes.xml",
            f'<w:footnotes xmlns:w="{w}">'
            '<w:footnote w:type="separator" w:id="-1"><w:p><w:r><w:separator/></w:r></w:p></w:footnote>'
            '<w:footnote w:id="1"><w:p><w:r><w:t>Michael Pollan, The Omnivore’s Dilemma (2006), 3.</w:t></w:r></w:p></w:footnote>'
            "</w:footnotes>",
        )
    path.write_bytes(out.getvalue())

    text, hints = await _hints(path)
    notes = [h for h in hints if h.kind == "note"]
    assert len(notes) == 1
    assert text[: notes[0].start].endswith("common.")
    assert notes[0].payload == "Michael Pollan, The Omnivore’s Dilemma (2006), 3."
