"""The two invariants, checked across every document shape we can construct.

These are the tests that are supposed to catch the *next* problem rather than
the last one. Every bug in this area so far has been an instance of one of two
failures, so both are asserted directly:

I1  An edit modifies only the text it was asked to modify. Nothing else in the
    document moves — not a line break, not a hyperlink, not a footnote
    reference, not somebody else's tracked change.

I2  There is one text representation. What the matcher searched is what the
    reader ends up with.

The corpus deliberately contains constructs this module knows nothing special
about. That is the point: the walker descends by default, so an element nobody
anticipated has to work without a rule of its own.
"""

from __future__ import annotations

import pytest
from docx import Document
from lxml import etree  # type: ignore

from andamentum.whetstone.docx.atoms import document_atoms, projection, qn
from andamentum.whetstone.docx.surgery import Edit, apply_edits
from andamentum.whetstone.docx.textmap import TextMap

SENTENCE = "The evidence layer earns its place in the analysis."
TAIL = "A closing sentence that mentions the deliberation layer."

# Elements whose count must never change unless an edit's span covered them.
WATCHED = ("br", "tab", "hyperlink", "drawing", "footnoteReference", "bookmarkStart")


def _run(parent, text: str):
    run = etree.SubElement(parent, qn("r"))
    element = etree.SubElement(run, qn("t"))
    element.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    element.text = text
    return run


# ── document shapes ─────────────────────────────────────────────────────


def plain():
    doc = Document()
    doc.add_paragraph(SENTENCE)
    doc.add_paragraph(TAIL)
    return doc


def with_line_break():
    doc = Document()
    para = doc.add_paragraph()
    _run(para._p, SENTENCE)
    etree.SubElement(etree.SubElement(para._p, qn("r")), qn("br"))
    _run(para._p, TAIL)
    return doc


def with_tab():
    doc = Document()
    para = doc.add_paragraph()
    _run(para._p, SENTENCE)
    etree.SubElement(etree.SubElement(para._p, qn("r")), qn("tab"))
    _run(para._p, TAIL)
    return doc


def with_hyperlink():
    doc = Document()
    para = doc.add_paragraph()
    _run(para._p, SENTENCE)
    link = etree.SubElement(para._p, qn("hyperlink"))
    _run(link, " See the plan. ")
    _run(para._p, TAIL)
    return doc


def with_content_control():
    doc = Document()
    para = doc.add_paragraph()
    _run(para._p, SENTENCE)
    content = etree.SubElement(etree.SubElement(para._p, qn("sdt")), qn("sdtContent"))
    _run(content, " Controlled text. ")
    _run(para._p, TAIL)
    return doc


def with_bookmark():
    doc = Document()
    para = doc.add_paragraph()
    _run(para._p, SENTENCE)
    mark = etree.SubElement(para._p, qn("bookmarkStart"))
    mark.set(qn("id"), "1")
    mark.set(qn("name"), "anchor")
    _run(para._p, TAIL)
    return doc


def with_footnote_reference():
    doc = Document()
    para = doc.add_paragraph()
    _run(para._p, SENTENCE)
    run = etree.SubElement(para._p, qn("r"))
    etree.SubElement(run, qn("footnoteReference")).set(qn("id"), "2")
    _run(para._p, TAIL)
    return doc


def with_formatting():
    doc = Document()
    para = doc.add_paragraph()
    _run(para._p, "The evidence layer ")
    bold = para.add_run("earns its place")
    bold.bold = True
    italic = para.add_run(" in the analysis.")
    italic.italic = True
    doc.add_paragraph(TAIL)
    return doc


def with_someone_elses_change():
    doc = Document()
    para = doc.add_paragraph()
    _run(para._p, SENTENCE)
    ins = etree.SubElement(para._p, qn("ins"))
    ins.set(qn("id"), "900")
    ins.set(qn("author"), "Dr Human")
    ins.set(qn("date"), "2026-01-01T00:00:00")
    _run(ins, " Added by a co-author. ")
    dele = etree.SubElement(para._p, qn("del"))
    dele.set(qn("id"), "901")
    dele.set(qn("author"), "Dr Human")
    dele.set(qn("date"), "2026-01-01T00:00:00")
    run = etree.SubElement(dele, qn("r"))
    etree.SubElement(run, qn("delText")).text = "struck by a co-author"
    _run(para._p, TAIL)
    return doc


def split_across_runs():
    doc = Document()
    para = doc.add_paragraph()
    for word in SENTENCE.split(" "):
        _run(para._p, word + " ")
    doc.add_paragraph(TAIL)
    return doc


SHAPES = [
    plain,
    with_line_break,
    with_tab,
    with_hyperlink,
    with_content_control,
    with_bookmark,
    with_footnote_reference,
    with_formatting,
    with_someone_elses_change,
    split_across_runs,
]

REPLACEMENTS = [
    ("earns its place", "earns its keep"),
    ("evidence layer", "evidence tier"),
    ("The evidence layer earns its place", "This layer earns its keep"),
    ("place in the analysis", "place in the write-up"),
]


def _paragraphs(doc):
    return doc.element.body.findall(qn("p"))


def _counts(doc) -> dict[str, int]:
    body = doc.element.body
    return {tag: len(body.findall(f".//{qn(tag)}")) for tag in WATCHED}


@pytest.mark.parametrize("shape", SHAPES, ids=lambda f: f.__name__)
@pytest.mark.parametrize("target,replacement", REPLACEMENTS, ids=lambda v: v[:18])
def test_an_edit_changes_only_what_it_was_asked_to_change(
    shape, target, replacement
) -> None:
    doc = shape()
    before_counts = _counts(doc)
    before_map = TextMap(document_atoms(_paragraphs(doc)))
    span = before_map.find(target)
    assert span is not None, f"{shape.__name__}: {target!r} did not resolve"

    expected = (
        before_map.text[: span.start] + replacement + before_map.text[span.end :]
    )

    apply_edits(before_map, [Edit(span=span, new_text=replacement, author="AI")])

    after = projection(document_atoms(_paragraphs(doc)))
    # I2: what the matcher searched is what the reader gets.
    assert after == expected
    # I1: nothing else in the document moved.
    assert _counts(doc) == before_counts


@pytest.mark.parametrize("shape", SHAPES, ids=lambda f: f.__name__)
def test_an_earlier_authors_change_is_never_silently_accepted(shape) -> None:
    """Editing a paragraph must not flatten a tracked change already in it."""
    doc = shape()
    paragraphs = _paragraphs(doc)
    before = {
        (e.tag, e.get(qn("author")))
        for e in doc.element.body.iter()
        if e.tag in (qn("ins"), qn("del"))
    }
    textmap = TextMap(document_atoms(paragraphs))
    span = textmap.find("earns its place")
    assert span is not None
    apply_edits(textmap, [Edit(span=span, new_text="earns its keep", author="AI")])

    after = {
        (e.tag, e.get(qn("author")))
        for e in doc.element.body.iter()
        if e.tag in (qn("ins"), qn("del"))
    }
    assert before <= after


@pytest.mark.parametrize("shape", SHAPES, ids=lambda f: f.__name__)
def test_edits_are_repeatable_over_their_own_output(shape) -> None:
    """A second round reads the first round's result and edits it correctly."""
    doc = shape()
    first = TextMap(document_atoms(_paragraphs(doc)))
    span = first.find("earns its place")
    assert span is not None
    apply_edits(first, [Edit(span=span, new_text="earns its keep", author="A")])

    second = TextMap(document_atoms(_paragraphs(doc)))
    assert "earns its keep" in second.text
    span2 = second.find("earns its keep")
    assert span2 is not None
    apply_edits(
        second, [Edit(span=span2, new_text="earns its salt", author="B")], first_id=9000
    )

    final = projection(document_atoms(_paragraphs(doc)))
    assert "earns its salt" in final
    assert "earns its keep" not in final
    authors = {
        e.get(qn("author"))
        for e in doc.element.body.iter()
        if e.tag in (qn("ins"), qn("del"))
    }
    assert {"A", "B"} <= authors


@pytest.mark.parametrize("shape", SHAPES, ids=lambda f: f.__name__)
def test_opening_and_saving_changes_nothing(shape, tmp_path) -> None:
    """A round trip with no patches must leave the document exactly as it was.

    It did not: the loader split every run into one run per token before any
    edit was requested, rewriting the author's run structure just by opening
    the file. Nothing rewrites anything now unless a span asks it to.
    """
    from andamentum.whetstone.docx.low_level import DocxEditor

    source = tmp_path / "in.docx"
    shape().save(str(source))
    before_runs = len(Document(str(source)).element.body.findall(f".//{qn('r')}"))
    before_counts = _counts(Document(str(source)))
    before_text = projection(
        document_atoms(Document(str(source)).element.body.findall(qn("p")))
    )

    out = tmp_path / "out.docx"
    DocxEditor(str(source), author="nobody").write(str(out))

    written = Document(str(out))
    assert len(written.element.body.findall(f".//{qn('r')}")) == before_runs
    assert _counts(written) == before_counts
    assert projection(document_atoms(written.element.body.findall(qn("p")))) == (
        before_text
    )
