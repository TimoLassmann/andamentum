"""Deterministic benchmark for .docx comment anchoring. No LLM, no network.

Scope note: this measures locating a supplied text
pattern in a Word file. The review path no longer searches for anything (see
:mod:`andamentum.whetstone.docx.annotate`), so its placement rate is 100% by
construction and a benchmark of it would only ever be measuring a bug.

A finding's quote comes from harvested markdown; a Word comment has to be
pinned back into the original .docx. Those are two different texts — the
converter reshapes tables, drops characters Word stores as structure, and
invents text Word generates from styles — so the pin is a search, and the
search can miss or, worse, hit the wrong place.

This module measures that, exactly, with no model in the loop.

How the ground truth is built
-----------------------------
Every block of every synthetic document carries a unique marker
(``BLK07``). The marker is the answer key, never the question:

* it identifies which markdown the block produced, so a quote can be
  sampled from the right place;
* it identifies which Word paragraph the quote *should* anchor to, so a
  landed anchor can be checked;
* quotes are sampled from the block's prose with the marker REMOVED, so a
  distinctive token never helps the matcher do its job.

The corpus is generated rather than collected, so the ground truth is
known by construction rather than annotated, and the whole run is
reproducible on any machine with no fixtures to keep in sync.

The metric that matters
-----------------------
Three outcomes per quote — ``correct``, ``wrong``, ``missed``. Anchor rate
alone is a bad target because guessing maximises it. ``wrong`` is the
number to drive to zero and keep there: a comment on the wrong sentence is
worse than no comment, because the reader cannot tell it is misplaced.

Run it::

    uv run python -m andamentum.whetstone.tests.anchor_bench
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Callable

# ── the answer key ──────────────────────────────────────────────────────

_MARKER = re.compile(r"\bBLK(\d{2})\b")


def marker(block_id: int) -> str:
    return f"BLK{block_id:02d}"


def marker_in(text: str) -> int | None:
    """The block id a piece of text belongs to, or None if unmarked."""
    match = _MARKER.search(text)
    return int(match.group(1)) if match else None


def markers_in(text: str) -> frozenset[int]:
    """Every block id mentioned — a converted table row names several."""
    return frozenset(int(m) for m in _MARKER.findall(text))


def strip_marker(text: str) -> str:
    """Remove the answer key, leaving the prose a quote is sampled from."""
    return " ".join(_MARKER.sub("", text).split())


# ── corpus ──────────────────────────────────────────────────────────────


@dataclass
class Case:
    """One synthetic document exercising one construct."""

    name: str
    build: Callable[[object], None]
    note: str = ""


@dataclass
class QuoteProbe:
    """One sampled quote and the block(s) it may legitimately anchor to.

    Usually one block. A table row is the exception: the converter renders
    a whole row as one markdown line, so a quote taken from it spans
    several Word cells and landing in any cell OF THAT ROW is correct.
    Insisting on one particular cell would score a correct anchor as wrong.
    """

    case: str
    expected: frozenset[int]
    quote: str
    kind: str  # full | prefix | interior — how much of the block was taken


def _para(document, block_id: int, text: str, **kw) -> None:
    document.add_paragraph(f"{marker(block_id)} {text}", **kw)


def _build_plain(document) -> None:
    _para(document, 1, "The assay detected degradation in stored plasma samples.")
    _para(document, 2, "Recruitment ran for six weeks across two clinical sites.")


def _build_tabbed(document) -> None:
    _para(document, 3, "Accuracy was\tninety-four percent across the cohort.")
    _para(document, 4, "Sensitivity\tand specificity were computed per site.")


def _build_table(document) -> None:
    _para(document, 5, "Results are summarised below for each cohort.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = f"{marker(6)} Cohort"
    table.cell(0, 1).text = f"{marker(7)} Accuracy"
    table.cell(1, 0).text = f"{marker(8)} Validation cohort"
    table.cell(1, 1).text = f"{marker(9)} ninety-four percent"


def _build_lists(document) -> None:
    _para(document, 10, "Samples were prepared as follows.", style="List Number")
    _para(
        document, 11, "Absorbance was measured at 280 nanometres.", style="List Number"
    )
    _para(document, 12, "Plates were washed three times.", style="List Bullet")


def _build_runs(document) -> None:
    paragraph = document.add_paragraph(f"{marker(13)} See ")
    paragraph.add_run("Table 1").bold = True
    paragraph.add_run(" for the full cohort breakdown of enrolled participants.")
    paragraph2 = document.add_paragraph(f"{marker(14)} The ")
    paragraph2.add_run("absorbance ratio").italic = True
    paragraph2.add_run(" exceeded the threshold in every degraded sample.")


def _build_headings(document) -> None:
    document.add_heading(f"{marker(15)} Methods and materials", level=2)
    _para(document, 16, "Plasma was collected under a standard venepuncture protocol.")
    document.add_heading(f"{marker(17)} Statistical analysis", level=3)
    _para(document, 18, "Confidence intervals were obtained by bootstrap resampling.")


def _build_typography(document) -> None:
    _para(document, 19, "The assay’s output was “degraded” or “intact”.")
    _para(document, 20, "Well‑characterised controls were run on every plate.")
    _para(document, 21, "Storage at −20 °C preserved the signal.")


def _build_line_breaks(document) -> None:
    paragraph = document.add_paragraph(f"{marker(22)} Degradation was observed")
    paragraph.add_run().add_break()
    paragraph.add_run("in every sample stored beyond the recommended window.")


def _build_repeats(document) -> None:
    """Genuinely ambiguous text — the same sentence twice, far apart.

    Correct behaviour is either the right block (disambiguated by position)
    or a refusal. What must NEVER happen is a confident anchor on the wrong
    one, and only a case like this can catch that.
    """
    _para(document, 23, "All differences were statistically significant.")
    _para(document, 24, "An unrelated sentence separates the two repeats here.")
    _para(document, 25, "Another unrelated sentence adds further separation.")
    _para(document, 26, "All differences were statistically significant.")


CASES: list[Case] = [
    Case("plain", _build_plain, "ordinary prose — the control"),
    Case("tabs", _build_tabbed, "Word stores tabs as structure, not text"),
    Case("tables", _build_table, "cells are separate paragraphs; markdown is one row"),
    Case("lists", _build_lists, "the '1. ' exists in no Word run"),
    Case("runs", _build_runs, "bold/italic split a sentence across runs"),
    Case("headings", _build_headings, "style-driven heading levels"),
    Case("typography", _build_typography, "smart quotes, non-breaking hyphen/space"),
    Case("line_breaks", _build_line_breaks, "in-paragraph break, like tabs"),
    Case("repeats", _build_repeats, "identical text twice — ambiguity, not absence"),
]


def build_corpus(directory: Path) -> list[tuple[Case, Path]]:
    """Write one .docx per case. Returns (case, path) pairs."""
    import docx

    directory.mkdir(parents=True, exist_ok=True)
    built: list[tuple[Case, Path]] = []
    for case in CASES:
        document = docx.Document()
        case.build(document)
        path = directory / f"{case.name}.docx"
        document.save(str(path))
        built.append((case, path))
    return built


# ── quote sampling (deterministic) ──────────────────────────────────────

_MIN_QUOTE_CHARS = 20


def sample_quotes(case_name: str, markdown: str) -> list[QuoteProbe]:
    """Take the quotes a reviewer would take, from the converted markdown.

    One probe per marked line, in three shapes, because they fail
    differently: a ``full`` line carries whatever decoration the converter
    added (list numbers, table pipes), a ``prefix`` stops before it, and an
    ``interior`` span starts past it. A construct that only breaks the full
    line is a decoration problem; one that breaks all three is a text
    problem.

    Deterministic: same markdown always yields the same probes, so two runs
    of the benchmark are comparable line for line.
    """
    probes: list[QuoteProbe] = []
    for line in markdown.splitlines():
        expected = markers_in(line)
        if not expected:
            continue
        full = line.strip()
        prose = strip_marker(full)
        if len(prose) < _MIN_QUOTE_CHARS:
            continue

        probes.append(QuoteProbe(case_name, expected, prose, "full"))

        words = prose.split()
        if len(words) >= 6:
            prefix = " ".join(words[: max(3, len(words) // 2)])
            if len(prefix) >= _MIN_QUOTE_CHARS:
                probes.append(QuoteProbe(case_name, expected, prefix, "prefix"))
            interior = " ".join(words[2:])
            if len(interior) >= _MIN_QUOTE_CHARS:
                probes.append(QuoteProbe(case_name, expected, interior, "interior"))
    return probes


# ── running ─────────────────────────────────────────────────────────────


@dataclass
class Outcome:
    correct: int = 0
    wrong: int = 0
    missed: int = 0
    wrong_detail: list[str] = field(default_factory=list)
    missed_detail: list[str] = field(default_factory=list)

    @property
    def total(self) -> int:
        return self.correct + self.wrong + self.missed

    @property
    def anchor_rate(self) -> float:
        return 100.0 * (self.correct + self.wrong) / self.total if self.total else 0.0

    @property
    def correct_rate(self) -> float:
        return 100.0 * self.correct / self.total if self.total else 0.0


def _paragraph_texts(docx_path: Path) -> list[str]:
    from andamentum.docx_edit.atoms import paragraph_atoms, projection
    from andamentum.whetstone.docx.low_level import DocxEditor

    editor = DocxEditor(str(docx_path))
    return [projection(list(paragraph_atoms(p))) for p in editor.paragraph_elements]


def _marker_free_atoms(paragraph_texts: list[str]) -> list:
    """One atom per character of the corpus, with the block markers removed.

    The markers are the harness's answer key, not part of the corpus: left in,
    one sitting between two table cells would break a cross-cell match and
    score the matcher down for the scaffolding. ``paragraph`` carries the
    paragraph number so a match can be scored against the key.
    """
    from andamentum.docx_edit.atoms import Atom, AtomKind

    stream: list[Atom] = []
    for number, text in enumerate(paragraph_texts):
        if number > 0:
            stream.append(
                Atom(
                    char="\n",
                    kind=AtomKind.PARAGRAPH_BREAK,
                    run=None,
                    paragraph=number - 1,
                    offset=-1,
                )
            )
        for offset, char in enumerate(strip_marker(text)):
            stream.append(
                Atom(
                    char=char,
                    kind=AtomKind.TEXT,
                    run=None,
                    paragraph=number,
                    offset=offset,
                )
            )
    return stream


def run_case(case: Case, docx_path: Path, markdown: str) -> Outcome:
    """Anchor every sampled quote and score it against the answer key.

    The markers are removed from the indexed text as well as from the
    quotes. They are metadata about the corpus, not part of it: left in,
    a marker sitting between two table cells would break a cross-cell
    match and score the matcher down for the harness's own scaffolding.
    Block identity is captured BEFORE stripping, so the answer key
    survives while staying invisible to the thing being measured.
    """
    from andamentum.docx_edit.textmap import TextMap

    raw = _paragraph_texts(docx_path)
    block_of = [marker_in(t) for t in raw]
    # The index is built over synthetic atoms carrying the paragraph number, so
    # the harness measures the same TextMap the product uses rather than a
    # parallel implementation of it.
    index = TextMap(_marker_free_atoms(raw))
    outcome = Outcome()

    for probe in sample_quotes(case.name, markdown):
        span = index.find(probe.quote)
        if span is None:
            outcome.missed += 1
            outcome.missed_detail.append(f"[{probe.kind}] {probe.quote[:70]}")
            continue
        landed = block_of[index.atoms[span.start].paragraph]
        if landed in probe.expected:
            outcome.correct += 1
        else:
            outcome.wrong += 1
            outcome.wrong_detail.append(
                f"[{probe.kind}] {probe.quote[:50]!r} → block {landed} "
                f"(expected {sorted(probe.expected)})"
            )
    return outcome


async def run_benchmark(directory: Path) -> dict[str, Outcome]:
    """Build the corpus, convert each document, anchor every quote."""
    from andamentum.harvest import extract

    results: dict[str, Outcome] = {}
    for case, path in build_corpus(directory):
        markdown = await extract(path)
        results[case.name] = run_case(case, path, markdown)
    return results


def format_report(results: dict[str, Outcome]) -> str:
    lines = [
        "",
        f"{'construct':<14} {'quotes':>6} {'correct':>8} {'wrong':>6} {'missed':>7}   note",
        "-" * 88,
    ]
    notes = {c.name: c.note for c in CASES}
    totals = Outcome()
    for name, outcome in results.items():
        totals.correct += outcome.correct
        totals.wrong += outcome.wrong
        totals.missed += outcome.missed
        flag = "  <-- WRONG ANCHORS" if outcome.wrong else ""
        lines.append(
            f"{name:<14} {outcome.total:>6} "
            f"{outcome.correct:>7}{'':1} {outcome.wrong:>6} {outcome.missed:>7}   "
            f"{notes.get(name, '')}{flag}"
        )
    lines.append("-" * 88)
    lines.append(
        f"{'TOTAL':<14} {totals.total:>6} {totals.correct:>7}  "
        f"{totals.wrong:>6} {totals.missed:>7}   "
        f"correct {totals.correct_rate:.1f}% / anchored {totals.anchor_rate:.1f}%"
    )

    for name, outcome in results.items():
        if outcome.wrong:
            lines.append(f"\nWRONG ANCHORS in {name} (must be zero):")
            lines.extend(f"  {d}" for d in outcome.wrong_detail)
    for name, outcome in results.items():
        if outcome.missed_detail:
            lines.append(f"\nmissed in {name}:")
            lines.extend(f"  {d}" for d in outcome.missed_detail[:6])
    return "\n".join(lines)


def main() -> None:
    import asyncio
    import tempfile

    with tempfile.TemporaryDirectory() as tmp:
        results = asyncio.run(run_benchmark(Path(tmp)))
    print(format_report(results))


if __name__ == "__main__":
    main()
