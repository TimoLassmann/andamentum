"""The baseline .docx written for a non-Word source keeps what a reference
check needs: superscripts, footnotes, and the characters docling escaped."""

from __future__ import annotations

import zipfile

from docx import Document

from andamentum.whetstone.docx.citation_hints import citation_hints
from andamentum.whetstone.reviewed_document import _write_baseline_docx, open_source

MARKDOWN = """## Introduction

Diagnosis is slow<sup>1</sup>. Li &amp; Durbin agree.[^1-1] More text.[^9]

[^1-1]: Michael Pollan, The Omnivore’s Dilemma (New York: Penguin &amp; Co, 2006), 99.
"""


def test_sup_becomes_a_superscript_run_and_entities_are_decoded() -> None:
    path = _write_baseline_docx(MARKDOWN)
    paragraph = Document(str(path)).paragraphs[1]
    runs = [(r.text, bool(r.font.superscript)) for r in paragraph.runs]
    assert ("1", True) in runs
    assert "Li & Durbin" in paragraph.text and "&amp;" not in paragraph.text


def test_a_markdown_footnote_becomes_a_word_footnote() -> None:
    path = _write_baseline_docx(MARKDOWN)
    with zipfile.ZipFile(path) as archive:
        body = archive.read("word/document.xml").decode()
        notes = archive.read("word/footnotes.xml").decode()
    assert body.count("<w:footnoteReference") == 1
    assert "Penguin &amp; Co" in notes  # escaped once, as XML — not twice
    # A mark with no definition stays as written; nothing is invented.
    assert "[^9]" in "".join(p.text for p in Document(str(path)).paragraphs)
    # The definition line is not body text.
    assert "Pollan" not in "".join(p.text for p in Document(str(path)).paragraphs)


async def test_the_reference_checker_sees_them_as_hints() -> None:
    document = await open_source(MARKDOWN)
    hints = citation_hints(document.docx_path, expected_text=document.text)
    kinds = sorted(h.kind for h in hints)
    assert kinds == ["note", "superscript"]
    note = next(h for h in hints if h.kind == "note")
    assert note.payload.startswith("Michael Pollan") and "&amp;" not in note.payload
