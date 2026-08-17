"""Heading detection reads the document, not a rendering of it.

Every test here is a case where looking for ``#`` in converted markdown gives
the wrong answer, or where a naive reading of the style would.
"""

from __future__ import annotations

from docx import Document
from lxml import etree  # type: ignore

from andamentum.whetstone.docx.atoms import document_atoms, projection, qn
from andamentum.whetstone.docx.headings import (
    StyleLevels,
    document_headings,
    heading_level,
    paragraph_spans,
)
from andamentum.whetstone.docx.reading import body_paragraphs, read_docx

_STYLES_XML = """<w:styles xmlns:w="{ns}">
  <w:style w:type="paragraph" w:styleId="Heading1">
    <w:name w:val="heading 1"/><w:pPr><w:outlineLvl w:val="0"/></w:pPr>
  </w:style>
  <w:style w:type="paragraph" w:styleId="Heading2">
    <w:name w:val="heading 2"/><w:pPr><w:outlineLvl w:val="1"/></w:pPr>
  </w:style>
  <w:style w:type="paragraph" w:styleId="Title"><w:name w:val="Title"/></w:style>
  <w:style w:type="paragraph" w:styleId="TOCHeading">
    <w:name w:val="TOC Heading"/><w:pPr><w:outlineLvl w:val="9"/></w:pPr>
  </w:style>
  <w:style w:type="paragraph" w:styleId="berschrift1"><w:name w:val="heading 1"/></w:style>
  <w:style w:type="paragraph" w:styleId="LabSectionHead">
    <w:name w:val="Lab Section Head"/><w:basedOn w:val="Heading2"/>
  </w:style>
  <w:style w:type="paragraph" w:styleId="BodyText"><w:name w:val="Body Text"/></w:style>
  <w:style w:type="paragraph" w:styleId="Loop"><w:name w:val="Loop"/>
    <w:basedOn w:val="Loop"/>
  </w:style>
</w:styles>""".format(ns="http://schemas.openxmlformats.org/wordprocessingml/2006/main")


def _styles():
    return StyleLevels(etree.fromstring(_STYLES_XML.encode()))


def _paragraph(text: str = "", *, style: str | None = None, outline: int | None = None):
    """A bare ``<w:p>``, optionally styled and/or carrying a direct outline level."""
    p = Document().add_paragraph()._p
    if style is not None or outline is not None:
        p_pr = etree.SubElement(p, qn("pPr"))
        # Word's schema order: pStyle first, then outlineLvl.
        if style is not None:
            etree.SubElement(p_pr, qn("pStyle")).set(qn("val"), style)
        if outline is not None:
            etree.SubElement(p_pr, qn("outlineLvl")).set(qn("val"), str(outline))
        p.insert(0, p_pr)
    if text:
        run = etree.SubElement(p, qn("r"))
        etree.SubElement(run, qn("t")).text = text
    return p


# ── style resolution ────────────────────────────────────────────────────


def test_built_in_heading_styles_resolve_to_their_level() -> None:
    styles = _styles()
    assert heading_level(_paragraph("Methods", style="Heading1"), styles=styles) == 1
    assert heading_level(_paragraph("Assay", style="Heading2"), styles=styles) == 2


def test_body_text_is_not_a_heading() -> None:
    styles = _styles()
    assert (
        heading_level(_paragraph("We did Y.", style="BodyText"), styles=styles) is None
    )
    assert heading_level(_paragraph("Unstyled prose."), styles=styles) is None


def test_title_opens_the_document_above_heading_one() -> None:
    """Level 0, so everything before the first Heading1 nests inside it."""
    assert heading_level(_paragraph("The Study", style="Title"), styles=_styles()) == 0


def test_direct_outline_level_outranks_the_style() -> None:
    """Direct formatting wins in Word; it has to win here too."""
    p = _paragraph("Results", style="BodyText", outline=1)
    assert heading_level(p, styles=_styles()) == 2


def test_outline_level_nine_means_body_text_not_level_ten() -> None:
    """Word reserves 9 for 'body text' — ``TOCHeading`` says exactly this."""
    styles = _styles()
    assert (
        heading_level(_paragraph("Contents", style="TOCHeading"), styles=styles) is None
    )
    assert heading_level(_paragraph("Contents", outline=9), styles=styles) is None


def test_localized_style_id_resolves_by_its_canonical_name() -> None:
    """A German file holds styleId 'berschrift1' with name 'heading 1'."""
    p = _paragraph("Einleitung", style="berschrift1")
    assert heading_level(p, styles=_styles()) == 1


def test_custom_style_resolves_through_based_on() -> None:
    p = _paragraph("Protocol", style="LabSectionHead")
    assert heading_level(p, styles=_styles()) == 2


def test_dangling_style_reference_falls_back_to_the_style_id() -> None:
    """The document uses a style it never defines; the id is all that is left."""
    p = _paragraph("Discussion", style="Heading3")
    assert heading_level(p, styles=_styles()) == 3


def test_based_on_cycle_does_not_hang() -> None:
    assert heading_level(_paragraph("Odd", style="Loop"), styles=_styles()) is None


def test_no_styles_part_still_reads_direct_formatting() -> None:
    styles = StyleLevels(None)
    assert heading_level(_paragraph("Results", outline=0), styles=styles) == 1
    assert heading_level(_paragraph("Results", style="Heading2"), styles=styles) == 2


# ── paragraph spans over the atom stream ────────────────────────────────


def test_paragraph_spans_index_the_projection_exactly() -> None:
    paragraphs = [_paragraph("One"), _paragraph(""), _paragraph("Three")]
    stream = document_atoms(paragraphs)
    text = projection(stream)

    spans = paragraph_spans(stream, paragraphs)
    assert [text[s.start : s.end] for s in spans] == ["One", "", "Three"]


def test_an_empty_paragraph_holds_its_place() -> None:
    """It contributes no atoms, so identity-grouping would silently drop it."""
    paragraphs = [_paragraph("One"), _paragraph(""), _paragraph("Three")]
    spans = paragraph_spans(document_atoms(paragraphs), paragraphs)
    assert (spans[1].start, spans[1].end) == (4, 4)
    assert spans[2].start == 5


def test_mismatched_stream_and_paragraph_list_is_refused() -> None:
    """Silently wrong offsets are worse than a crash — every anchor moves."""
    paragraphs = [_paragraph("One"), _paragraph("Two")]
    stream = document_atoms(paragraphs)
    try:
        paragraph_spans(stream, paragraphs[:1])
    except ValueError:
        pass
    else:  # pragma: no cover - the assertion below reports the failure
        raise AssertionError("expected a ValueError for a desynced paragraph list")


# ── whole-document detection ────────────────────────────────────────────


def _doc_headings(paragraphs, *, styles_root=None):
    stream = document_atoms(paragraphs)
    return document_headings(paragraphs, stream=stream, styles_root=styles_root)


def test_headings_carry_projection_offsets_and_stripped_titles() -> None:
    paragraphs = [
        _paragraph("Introduction", style="Heading1"),
        _paragraph("We study X."),
        _paragraph("Methods", style="Heading1"),
    ]
    stream = document_atoms(paragraphs)
    text = projection(stream)
    found = _doc_headings(
        paragraphs, styles_root=etree.fromstring(_STYLES_XML.encode())
    )

    assert [(h.level, h.title) for h in found] == [(1, "Introduction"), (1, "Methods")]
    assert [text[h.start : h.end] for h in found] == ["Introduction", "Methods"]


def test_a_heading_split_across_runs_reads_as_one_title() -> None:
    """Bold or italic inside a heading splits it into several runs in Word."""
    p = _paragraph(style="Heading1")
    for chunk in ("The ", "BRCA1", " pathway"):
        run = etree.SubElement(p, qn("r"))
        etree.SubElement(run, qn("t")).text = chunk

    found = _doc_headings([p], styles_root=etree.fromstring(_STYLES_XML.encode()))
    assert [h.title for h in found] == ["The BRCA1 pathway"]


def test_an_empty_styled_paragraph_is_not_a_heading() -> None:
    paragraphs = [_paragraph("", style="Heading1"), _paragraph("Body.")]
    assert _doc_headings(paragraphs, styles_root=None) == []


def test_a_heading_styled_table_cell_is_not_a_section_boundary(tmp_path) -> None:
    """Word holds each cell as its own ``<w:p>``; a styled one is still a cell."""
    doc = Document()
    doc.add_heading("Results", 1)
    table = doc.add_table(rows=1, cols=1)
    cell_paragraph = table.rows[0].cells[0].paragraphs[0]
    cell_p = cell_paragraph._p
    p_pr = etree.SubElement(cell_p, qn("pPr"))
    etree.SubElement(p_pr, qn("pStyle")).set(qn("val"), "Heading1")
    cell_p.insert(0, p_pr)
    cell_paragraph.add_run("Cohort")
    path = tmp_path / "table.docx"
    doc.save(path)

    parts = read_docx(path)
    found = _doc_headings(parts.paragraphs, styles_root=parts.styles_root)
    assert [h.title for h in found] == ["Results"]


# ── reading a real package ──────────────────────────────────────────────


def test_read_docx_returns_the_same_paragraphs_the_editor_would(tmp_path) -> None:
    doc = Document()
    doc.add_heading("Title", 0)
    doc.add_paragraph("Body.")
    path = tmp_path / "round.docx"
    doc.save(path)

    from andamentum.whetstone.docx.low_level import DocxEditor

    editor = DocxEditor(str(path))
    expected = projection(document_atoms(editor.paragraph_elements))
    assert projection(document_atoms(read_docx(path).paragraphs)) == expected


def test_body_paragraphs_skips_paragraphs_nested_in_a_paragraph() -> None:
    """A text box's content is read in place by the atom walker, once."""
    doc = Document()
    outer = doc.add_paragraph("Anchor")._p
    run = etree.SubElement(outer, qn("r"))
    etree.SubElement(etree.SubElement(run, qn("txbxContent")), qn("p"))

    assert body_paragraphs(outer.getparent()) == [outer]
