"""Rendering to docx, md, html, txt."""

import pytest
from docx import Document as DocxDocument

from andamentum.scribe import Bundle, ScribeError

from ._png import PNG_BYTES as PNG


BODY = (
    "A paragraph with **bold**, *italic* and `code`.\n\n"
    "- first\n- second\n\n"
    "1. one\n2. two\n\n"
    "| Name | Score |\n|---|---|\n| A | 1 |\n\n"
    "![A caption](figures/plot.png){#fig:plot}"
)


@pytest.fixture
def bundle(tmp_path):
    b = Bundle.create(tmp_path / "d.compose", title="Apollo Report")
    (tmp_path / "d.compose" / "figures" / "plot.png").write_bytes(PNG)
    b.add_section(title="Summary", body=BODY)
    b.add_section(title="Details", body="Plain. See [@k1].")
    b.add_reference(key="k1", text="Doe. A paper.", url="https://example.org/p")
    return b


def test_docx_styles(bundle, tmp_path):
    out = bundle.render(tmp_path / "o.docx", format="docx")
    doc = DocxDocument(str(out))
    styles = [(p.style.name if p.style else "", p.text) for p in doc.paragraphs]
    assert styles[0] == ("Title", "Apollo Report")
    headings = [t for s, t in styles if s == "Heading 1"]
    assert headings == ["Summary", "Details", "References"]
    assert [t for s, t in styles if s == "List Bullet"] == ["first", "second"]
    assert [t for s, t in styles if s == "List Number"] == ["one", "two"]
    assert ("Caption", "A caption") in styles
    assert len(doc.inline_shapes) == 1
    assert doc.tables[0].style is not None and doc.tables[0].style.name == "Table Grid"
    assert doc.tables[0].rows[0].cells[0].paragraphs[0].runs[0].bold
    assert doc.tables[0].rows[1].cells[1].text == "1"
    para = next(p for p in doc.paragraphs if p.text.startswith("A paragraph"))
    assert [r.text for r in para.runs if r.bold] == ["bold"]
    assert [r.text for r in para.runs if r.italic] == ["italic"]
    assert "[k1] Doe. A paper. https://example.org/p" in [t for _, t in styles]


def test_docx_text_filter_applies_to_bodies_only(bundle, tmp_path):
    out = bundle.render(
        tmp_path / "o.docx",
        format="docx",
        text_filter=lambda t: t.replace("Plain", "Filtered"),
    )
    text = "\n".join(p.text for p in DocxDocument(str(out)).paragraphs)
    assert "Filtered. See" in text and "Plain" not in text


def test_docx_missing_image_placeholder(bundle, tmp_path):
    bundle.add_section(title="Gap", body="![c](figures/gone.png)")
    out = bundle.render(tmp_path / "o.docx", format="docx")
    text = "\n".join(p.text for p in DocxDocument(str(out)).paragraphs)
    assert "[Missing figure: figures/gone.png]" in text


def test_md_render_with_relative_image_and_references(bundle):
    out = bundle.render(bundle.path / "out" / "o.md", format="md")
    text = out.read_text()
    assert text.startswith("# Apollo Report\n\n## Summary\n\n")
    assert "![A caption](../figures/plot.png){#fig:plot}" in text
    assert "- first\n- second" in text
    assert "## References\n\n- [k1] Doe. A paper. <https://example.org/p>" in text
    assert text.endswith("\n")


def test_txt_render(bundle, tmp_path):
    text = bundle.render(tmp_path / "o.txt", format="txt").read_text()
    assert text.startswith(
        "Apollo Report\n\nSummary\n\nA paragraph with bold, italic and code."
    )
    assert "- first\n- second" in text and "1. one\n2. two" in text
    assert "Name | Score" in text and "[Figure: A caption]" in text
    assert "References" not in text and "**" not in text


def test_html_render(bundle, tmp_path):
    out = bundle.render(tmp_path / "o.html", format="html")
    html = out.read_text()
    assert "<title>Apollo Report</title>" in html
    assert ">Summary</h2>" in html and ">Details</h2>" in html
    assert "<strong>bold</strong>" in html and "<li>first</li>" in html
    assert "<table" in html and "<img" in html and "plot.png" in html
    assert "[k1] Doe. A paper." in html


def test_render_is_deterministic_for_text_formats(bundle, tmp_path):
    a = bundle.render(tmp_path / "a.md", format="md").read_bytes()
    b = bundle.render(tmp_path / "b.md", format="md").read_bytes()
    assert a == b


def test_unknown_format(bundle, tmp_path):
    with pytest.raises(ScribeError, match="unknown format"):
        bundle.render(tmp_path / "o.rtf", format="rtf")  # type: ignore[arg-type]


def test_docx_ai_marker_keyword(bundle, tmp_path):
    clean = DocxDocument(str(bundle.render(tmp_path / "a.docx", format="docx")))
    assert "andamentum:contains-ai-markers" not in (
        clean.core_properties.keywords or ""
    )
    bundle.add_section(title="Drafted", body="Some text. [ai-drafted]")
    marked = DocxDocument(str(bundle.render(tmp_path / "b.docx", format="docx")))
    assert "andamentum:contains-ai-markers" in (marked.core_properties.keywords or "")


def test_md_render_image_root_override(bundle, tmp_path):
    out_dir = tmp_path / "elsewhere"
    out = bundle.render(out_dir / "o.md", format="md", image_root=bundle.path)
    assert "![A caption](" in out.read_text(encoding="utf-8")
    assert (out_dir / "../d.compose/figures/plot.png").resolve().is_file()
