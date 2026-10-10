"""Every corpus fixture opens with python-docx and has its key feature."""

from __future__ import annotations

import zipfile
from pathlib import Path

import pytest
from docx import Document as DocxDocument
from docx.document import Document as WordDocument

from andamentum.tests.docx_corpus.builder import (
    ALL_FIXTURES,
    FORMATTED_RUNS_TEXT,
    HYPERLINK_TARGET,
    LIST_ITEMS,
    TABLE_CELLS,
    TRACKED_ACCEPTED_TEXT,
    accepted_text,
    document_xml,
    existing_comments,
    existing_tracked,
    fields_and_controls,
    formatted_runs,
    headed_sections,
    hyperlink_and_footnote,
    plain_two_paragraphs,
    table_and_list,
)


def _open(path: Path) -> WordDocument:
    return DocxDocument(str(path))


def _style_name(paragraph: object) -> str:
    """Style name of a paragraph, or "" when the paragraph has no style."""
    style = paragraph.style  # type: ignore[attr-defined]
    return style.name if style is not None else ""


def test_plain_two_paragraphs(tmp_path: Path) -> None:
    doc = _open(plain_two_paragraphs(tmp_path))
    assert len(doc.paragraphs) == 2
    assert all(len(p.runs) == 1 for p in doc.paragraphs)
    assert doc.paragraphs[0].text.startswith("Jane Doe recorded")


def test_formatted_runs(tmp_path: Path) -> None:
    doc = _open(formatted_runs(tmp_path))
    assert len(doc.paragraphs) == 1
    paragraph = doc.paragraphs[0]
    assert paragraph.text == FORMATTED_RUNS_TEXT
    runs = paragraph.runs
    assert any(r.bold and r.text == "Sali" for r in runs)
    assert any(r.text == "ne" and not r.bold for r in runs)  # mid-word boundary
    assert any(r.italic for r in runs)
    assert any(r.font.superscript for r in runs)
    assert any(r.font.color.rgb is not None for r in runs)


def test_headed_sections(tmp_path: Path) -> None:
    doc = _open(headed_sections(tmp_path))
    assert len(doc.paragraphs) == 9
    styles = [_style_name(p) for p in doc.paragraphs]
    assert styles[0] == "Title"
    assert "Heading 1" in styles
    assert "Heading 2" in styles
    bibliography = [p for p in doc.paragraphs if p.text == "Bibliography"]
    assert len(bibliography) == 1
    assert _style_name(bibliography[0]) == "Heading 1"


def test_hyperlink_and_footnote(tmp_path: Path) -> None:
    path = hyperlink_and_footnote(tmp_path)
    doc = _open(path)
    assert len(doc.paragraphs) == 2
    xml = document_xml(path)
    assert "<w:hyperlink " in xml
    assert "<w:bookmarkStart " in xml and 'w:name="PilotStart"' in xml
    assert 'w:footnoteReference w:id="1"' in xml
    with zipfile.ZipFile(path) as archive:
        assert "word/footnotes.xml" in archive.namelist()
    link_ids = [
        rel.rId for rel in doc.part.rels.values() if rel.reltype.endswith("/hyperlink")
    ]
    assert len(link_ids) == 1
    assert doc.part.rels[link_ids[0]].target_ref == HYPERLINK_TARGET


def test_existing_tracked(tmp_path: Path) -> None:
    path = existing_tracked(tmp_path)
    doc = _open(path)
    assert len(doc.paragraphs) == 1
    xml = document_xml(path)
    assert "<w:ins " in xml and "<w:del " in xml
    assert xml.count('w:author="Jane Doe"') == 2
    assert accepted_text(path) == TRACKED_ACCEPTED_TEXT


def test_existing_comments(tmp_path: Path) -> None:
    path = existing_comments(tmp_path)
    doc = _open(path)
    assert len(doc.paragraphs) == 1
    with zipfile.ZipFile(path) as archive:
        names = set(archive.namelist())
        assert {"word/comments.xml", "word/commentsExtended.xml"} <= names
        extended = archive.read("word/commentsExtended.xml").decode("utf-8")
    assert 'w15:paraIdParent="20000001"' in extended
    xml = document_xml(path)
    assert xml.count("<w:commentReference") == 3
    assert xml.count("<w:commentRangeStart") == 3


def test_table_and_list(tmp_path: Path) -> None:
    doc = _open(table_and_list(tmp_path))
    assert len(doc.tables) == 1
    table = doc.tables[0]
    assert len(table.rows) == 2 and len(table.columns) == 2
    assert [c.text for c in table.rows[0].cells] == list(TABLE_CELLS[0])
    assert [c.text for c in table.rows[1].cells] == list(TABLE_CELLS[1])
    assert len(doc.paragraphs) == 1 + len(LIST_ITEMS)
    list_items = doc.paragraphs[1:]
    assert all(p._p.pPr is not None and p._p.pPr.numPr is not None for p in list_items)
    assert [p.text for p in list_items] == list(LIST_ITEMS)


def test_fields_and_controls(tmp_path: Path) -> None:
    path = fields_and_controls(tmp_path)
    doc = _open(path)
    assert len(doc.paragraphs) == 2
    xml = document_xml(path)
    assert 'w:name="_Ref200"' in xml
    assert "REF _Ref200" in xml
    assert "PAGE" in xml
    assert 'w:tag w:val="am:claim-1"' in xml
    assert "w:fldSimple" in xml


def test_accepted_text_drops_deletions(tmp_path: Path) -> None:
    path = existing_tracked(tmp_path)
    assert "on paper" not in accepted_text(path)
    assert "every meal" in accepted_text(path)


def test_accepted_text_joins_paragraphs_with_newlines(tmp_path: Path) -> None:
    path = plain_two_paragraphs(tmp_path)
    lines = accepted_text(path).split("\n")
    assert lines == [
        "Jane Doe recorded the first SaltLog reading at Apollo.",
        "The second paragraph follows on from the first.",
    ]


def test_document_xml_is_the_body_part(tmp_path: Path) -> None:
    path = plain_two_paragraphs(tmp_path)
    assert document_xml(path).startswith("<?xml")
    assert "<w:document" in document_xml(path)


@pytest.mark.parametrize("builder", ALL_FIXTURES, ids=lambda f: f.__name__)
def test_fixture_writes_named_file_in_given_directory(builder, tmp_path: Path) -> None:
    path = builder(tmp_path)
    assert path == tmp_path / f"{builder.__name__}.docx"
    assert path.is_file()
    assert _open(path).paragraphs is not None
