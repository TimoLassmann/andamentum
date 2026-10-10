"""Characterisation of compose's DocxAdapter writer (Engine C).

These tests pin the behaviour the writer has today on the synthetic corpus in
``andamentum.tests.docx_corpus``. They are not a specification: a failing test
means the behaviour changed and the change needs a decision.
"""

from __future__ import annotations

import zipfile
from pathlib import Path

import docx
import pytest
from lxml import etree  # type: ignore

from andamentum.compose.harness.docx import writer as writer_module
from andamentum.compose.harness.docx.ooxml import w, w15
from andamentum.compose.harness.docx.writer import DocxAdapter, DocxApplyError
from andamentum.compose.harness.kernel import (
    Bind,
    CitationItem,
    CommentOp,
    DocumentView,
    Insert,
    InsertCitation,
    InsertParagraph,
    Replace,
    Reply,
    Transaction,
    UpdateBibliography,
)
from andamentum.compose.harness.spike.make import (
    build_docx,
    complex_field,
    paragraph,
    run,
)
from andamentum.tests.docx_corpus.builder import (
    existing_comments,
    fields_and_controls,
    formatted_runs,
    headed_sections,
    hyperlink_and_footnote,
    plain_two_paragraphs,
    table_and_list,
)

AUTHOR = "andamentum compose"

# Synthetic works, invented for these tests.
SMITH = {
    "id": "smith2021",
    "type": "article-journal",
    "title": "Weekly salt logging in invented clinics",
    "container-title": "Journal of Invented Results",
    "author": [
        {"family": "Smith", "given": "Jane Q."},
        {"family": "Jones", "given": "Kim"},
    ],
    "issued": {"date-parts": [[2021]]},
    "DOI": "10.5555/saltlog.smith",
}
DOE = {
    "id": "doe2020",
    "title": "Sodium and salt logging in invented cohorts",
    "author": [{"family": "Doe", "given": "Jane"}],
    "issued": {"date-parts": [[2020]]},
    "DOI": "10.9999/spike.001",
}


@pytest.fixture
def adapter() -> DocxAdapter:
    return DocxAdapter()


def _txn(*ops) -> Transaction:
    return Transaction(ops=list(ops), author=AUTHOR, reason="test", source="test")


def _text(view: DocumentView, paragraph_id: str) -> str:
    return next(p.text for p in view.paragraphs if p.id == paragraph_id)


def _span(view: DocumentView, paragraph_id: str, needle: str) -> tuple[int, int]:
    start = _text(view, paragraph_id).index(needle)
    return start, start + len(needle)


def _parts(path: Path) -> dict[str, bytes]:
    with zipfile.ZipFile(path) as archive:
        return {name: archive.read(name) for name in archive.namelist()}


def _root(path: Path):
    return etree.fromstring(_parts(path)["word/document.xml"])


def _body_paragraphs(path: Path) -> list:
    return list(_root(path).find(w("body")).iter(w("p")))


def _tracked_by(root, tag: str, author: str) -> list:
    return [e for e in root.iter(w(tag)) if e.get(w("author")) == author]


def _in_field_result(root) -> list[str]:
    """Revision, comment and control tags between a ``separate`` and its ``end``."""
    inside: list[str] = []
    depth = 0
    for element in root.iter():
        if element.tag == w("fldChar"):
            kind = element.get(w("fldCharType"))
            if kind == "separate":
                depth += 1
            elif kind == "end":
                depth = max(0, depth - 1)
        elif depth and element.tag in {
            w("ins"),
            w("del"),
            w("commentRangeStart"),
            w("commentRangeEnd"),
            w("commentReference"),
            w("sdt"),
        }:
            inside.append(element.tag.rpartition("}")[2])
    return inside


# ------------------------------------------------------------ replace and insert


def test_replace_with_matching_expect_is_a_tracked_change(
    adapter: DocxAdapter, tmp_path: Path
):
    source = plain_two_paragraphs(tmp_path)
    view = adapter.open(source)
    start, end = _span(view, view.paragraphs[0].id, "Jane Doe")
    result = adapter.apply(
        view,
        _txn(
            Replace(
                paragraph_id="p0",
                start=start,
                end=end,
                text="Joan Roe",
                expect="Jane Doe",
            )
        ),
        output=tmp_path / "out.docx",
    )
    assert result.applied == 1 and result.refused == []
    assert result.untouched.ok and result.untouched.changed_outside == []
    out = adapter.open(result.output)
    assert _text(out, "p0") == (
        "Joan Roe recorded the first SaltLog reading at Apollo."
    )
    assert _text(out, "p1") == _text(view, "p1")
    root = _root(result.output)
    deleted = "".join(
        t.text or ""
        for d in _tracked_by(root, "del", AUTHOR)
        for t in d.iter(w("delText"))
    )
    inserted = "".join(
        t.text or "" for i in _tracked_by(root, "ins", AUTHOR) for t in i.iter(w("t"))
    )
    assert "Jane" in deleted and "Doe" in deleted
    assert "Joan" in inserted and "Roe" in inserted
    # Nothing outside the run is marked: the untouched word survives unchanged.
    assert "recorded the first" not in deleted and "recorded the first" not in inserted


def test_insert_is_a_tracked_insertion_and_keeps_the_text_around_it(
    adapter: DocxAdapter, tmp_path: Path
):
    view = adapter.open(plain_two_paragraphs(tmp_path))
    at = _text(view, "p1").index("first") + len("first")
    result = adapter.apply(
        view,
        _txn(Insert(paragraph_id="p1", at=at, text=" again")),
        output=tmp_path / "out.docx",
    )
    assert result.applied == 1 and result.refused == []
    out = adapter.open(result.output)
    assert _text(out, "p1") == "The second paragraph follows on from the first again."
    inserted = _tracked_by(_root(result.output), "ins", AUTHOR)
    assert len(inserted) == 1
    assert "".join(t.text or "" for t in inserted[0].iter(w("t"))) == " again"


# ------------------------------------------------------------------- refusals


def test_expect_mismatch_is_refused_and_leaves_the_paragraph_byte_identical(
    adapter: DocxAdapter, tmp_path: Path
):
    source = plain_two_paragraphs(tmp_path)
    view = adapter.open(source)
    result = adapter.apply(
        view,
        _txn(
            Replace(
                paragraph_id="p0", start=0, end=8, text="Joan Roe", expect="Jane Roe"
            )
        ),
        output=tmp_path / "out.docx",
    )
    assert result.applied == 0
    assert result.refused == [
        "replace p0[0:8]: expected 'Jane Roe' but the range reads 'Jane Doe'"
    ]
    before = _body_paragraphs(source)
    after = _body_paragraphs(result.output)
    assert [etree.tostring(p) for p in before] == [etree.tostring(p) for p in after]
    assert _root(result.output).find(f".//{w('ins')}") is None


def test_newline_in_an_edit_is_refused_and_points_at_insert_paragraph(
    adapter: DocxAdapter, tmp_path: Path
):
    view = adapter.open(plain_two_paragraphs(tmp_path))
    result = adapter.apply(
        view,
        _txn(
            Replace(
                paragraph_id="p0", start=0, end=8, text="Joan\nRoe", expect="Jane Doe"
            ),
            Insert(paragraph_id="p1", at=0, text="Next: "),
        ),
        output=tmp_path / "out.docx",
    )
    assert result.applied == 1
    assert result.refused == [
        "replace p0[0:8]: text contains a newline; use InsertParagraph"
    ]
    out = adapter.open(result.output)
    assert _text(out, "p0") == _text(view, "p0")
    assert _text(out, "p1") == "Next: The second paragraph follows on from the first."


def test_edit_into_a_field_result_is_refused_in_both_forms(
    adapter: DocxAdapter, tmp_path: Path
):
    source = fields_and_controls(tmp_path)
    view = adapter.open(source)
    # "See Table 2 on page 1 ...": the REF result "Table 2" sits at 4..11.
    assert _text(view, "p1")[4:11] == "Table 2"
    result = adapter.apply(
        view,
        _txn(
            Replace(
                paragraph_id="p1", start=4, end=11, text="Table 3", expect="Table 2"
            ),
            Insert(paragraph_id="p1", at=6, text="X"),
        ),
        output=tmp_path / "out.docx",
    )
    assert result.applied == 0
    assert result.refused == [
        "replace p1[4:11]: range edits a field result, which is regenerated on refresh",
        "insert p1@6: insertion inside a field result, which is regenerated on refresh",
    ]
    assert _in_field_result(_root(result.output)) == []
    before = _body_paragraphs(source)[1]
    after = _body_paragraphs(result.output)[1]
    assert etree.tostring(before) == etree.tostring(after)


# ------------------------------------------------------------- new paragraphs


def test_insert_paragraph_is_a_tracked_paragraph_with_its_own_mark(
    adapter: DocxAdapter, tmp_path: Path
):
    view = adapter.open(plain_two_paragraphs(tmp_path))
    result = adapter.apply(
        view,
        _txn(
            InsertParagraph(after_paragraph_id="p0", text="Inserted note."),
            InsertParagraph(
                after_paragraph_id="p0", text="Added heading.", style="heading 1"
            ),
        ),
        output=tmp_path / "out.docx",
    )
    assert result.applied == 2 and result.refused == []
    out = adapter.open(result.output)
    texts = [p.text for p in out.paragraphs]
    assert texts == [
        "Jane Doe recorded the first SaltLog reading at Apollo.",
        "Inserted note.",
        "Added heading.",
        "The second paragraph follows on from the first.",
    ]
    assert out.paragraphs[1].style == "Normal"
    assert out.paragraphs[2].style == "heading 1"
    assert out.paragraphs[2].heading_level == 1
    root = _root(result.output)
    for text in ("Inserted note.", "Added heading."):
        (new,) = [p for p in root.iter(w("p")) if "".join(p.itertext()) == text]
        mark = new.find(f"{w('pPr')}/{w('rPr')}/{w('ins')}")
        assert mark is not None and mark.get(w("author")) == AUTHOR
        assert new.find(w("ins")) is not None


def test_insert_paragraph_with_an_unknown_style_is_refused(
    adapter: DocxAdapter, tmp_path: Path
):
    source = plain_two_paragraphs(tmp_path)
    view = adapter.open(source)
    result = adapter.apply(
        view,
        _txn(
            InsertParagraph(after_paragraph_id="p0", text="X.", style="No Such Style")
        ),
        output=tmp_path / "out.docx",
    )
    assert result.applied == 0
    assert result.refused == [
        "insert_paragraph after p0: style 'No Such Style' is not in the document"
    ]
    assert [p.text for p in adapter.open(result.output).paragraphs] == [
        p.text for p in view.paragraphs
    ]


# ------------------------------------------------------------ comments, replies


def test_comment_on_a_plain_document_creates_the_comments_part(
    adapter: DocxAdapter, tmp_path: Path
):
    view = adapter.open(plain_two_paragraphs(tmp_path))
    start, end = _span(view, "p0", "Jane Doe")
    result = adapter.apply(
        view,
        _txn(CommentOp(paragraph_id="p0", start=start, end=end, text="Which Jane?")),
        output=tmp_path / "out.docx",
    )
    assert result.applied == 1 and result.refused == []
    out = adapter.open(result.output)
    (comment,) = out.comments
    assert comment.author == AUTHOR and comment.text == "Which Jane?"
    # A document with no comments part starts its comment ids at 1.
    assert comment.parent_id is None and comment.id == "1"
    assert out.paragraphs[0].text[comment.start : comment.end] == "Jane Doe"
    assert _text(out, "p0") == _text(view, "p0")
    names = set(_parts(result.output))
    assert "word/comments.xml" in names
    assert b"comments+xml" in _parts(result.output)["[Content_Types].xml"]


def test_reply_joins_the_thread_of_its_root_and_extends_the_thread_part(
    adapter: DocxAdapter, tmp_path: Path
):
    source = existing_comments(tmp_path)
    view = adapter.open(source)
    result = adapter.apply(
        view,
        _txn(
            Reply(parent_comment_id="2", text="Agreed, I will reword it."),
            Reply(parent_comment_id="1", text="The unit is grams."),
        ),
        output=tmp_path / "out.docx",
    )
    assert result.applied == 2 and result.refused == []
    out = adapter.open(result.output)
    new = [c for c in out.comments if c.author == AUTHOR]
    assert [c.id for c in new] == ["3", "4"]
    assert [c.parent_id for c in new] == ["0", "1"]
    assert [c.text for c in new] == [
        "Agreed, I will reword it.",
        "The unit is grams.",
    ]
    by_id = {c.id: c for c in out.comments}
    assert (new[0].paragraph_id, new[0].start, new[0].end) == (
        by_id["0"].paragraph_id,
        by_id["0"].start,
        by_id["0"].end,
    )
    assert (new[1].start, new[1].end) == (by_id["1"].start, by_id["1"].end)
    marks = [e.get(w("id")) for e in _root(result.output).iter(w("commentReference"))]
    # A reply goes after the last member of its thread: reply 4 (to comment 1)
    # sits after comment 1's marker, reply 3 (to comment 2) after comment 2's.
    assert marks == ["0", "1", "4", "2", "3"]
    extended = _parts(result.output)["word/commentsExtended.xml"]
    entries = etree.fromstring(extended).findall(w15("commentEx"))
    parents = [e.get(w15("paraIdParent")) for e in entries]
    assert parents == [None, None, "20000001", "20000001", "20000002"]


def test_reply_to_an_unknown_comment_is_refused(adapter: DocxAdapter, tmp_path: Path):
    view = adapter.open(existing_comments(tmp_path))
    result = adapter.apply(
        view,
        _txn(Reply(parent_comment_id="99", text="Hello?")),
        output=tmp_path / "out.docx",
    )
    assert result.applied == 0
    assert result.refused == ["reply to comment 99: no such comment"]


def test_comment_range_inside_formatted_runs_keeps_the_text_and_the_formatting(
    adapter: DocxAdapter, tmp_path: Path
):
    view = adapter.open(formatted_runs(tmp_path))
    start, end = _span(view, "p0", "Saline")  # "Sali" is bold, "ne" is plain
    result = adapter.apply(
        view,
        _txn(CommentOp(paragraph_id="p0", start=start, end=end, text="Spelling?")),
        output=tmp_path / "out.docx",
    )
    assert result.applied == 1
    out = adapter.open(result.output)
    (comment,) = out.comments
    assert out.paragraphs[0].text[comment.start : comment.end] == "Saline"
    assert out.paragraphs[0].text == view.paragraphs[0].text
    bold = [
        r
        for r in _root(result.output).iter(w("r"))
        if r.find(w("t")) is not None
        and r.find(w("t")).text == "Sali"
        and r.find(f"{w('rPr')}/{w('b')}") is not None
    ]
    assert len(bold) == 1


# ----------------------------------------------------------------------- binds


def test_bind_wraps_the_range_in_a_hidden_tagged_control(
    adapter: DocxAdapter, tmp_path: Path
):
    view = adapter.open(plain_two_paragraphs(tmp_path))
    result = adapter.apply(
        view,
        _txn(
            Bind(paragraph_id="p0", start=0, end=8, anchor_id="person-1", kind="span")
        ),
        output=tmp_path / "out.docx",
    )
    assert result.applied == 1 and result.refused == []
    out = adapter.open(result.output)
    (anchor,) = out.anchors
    assert anchor.id == "person-1" and anchor.kind == "span"
    assert out.paragraphs[0].text[anchor.start : anchor.end] == "Jane Doe"
    assert _text(out, "p0") == _text(view, "p0")
    sdt = next(
        s
        for s in _root(result.output).iter(w("sdt"))
        if s.find(f"{w('sdtPr')}/{w('tag')}").get(w("val")) == "am:person-1"
    )
    props = sdt.find(w("sdtPr"))
    assert props.find(w("alias")).get(w("val")) == "span"
    assert props.find(w15("appearance")).get(w15("val")) == "hidden"
    assert props.find(w("lock")) is None


def test_bind_refusals_cut_fields_collide_with_anchors_and_edits(
    adapter: DocxAdapter, tmp_path: Path
):
    view = adapter.open(fields_and_controls(tmp_path))
    result = adapter.apply(
        view,
        _txn(
            # Inside the REF result "Table 2": cuts the field.
            Bind(paragraph_id="p1", start=6, end=9, anchor_id="n1", kind="span"),
            # The anchor "claim-1" already exists in the document.
            Bind(paragraph_id="p1", start=26, end=30, anchor_id="claim-1", kind="span"),
            # Overlaps the Replace in this transaction.
            Bind(paragraph_id="p1", start=0, end=3, anchor_id="n2", kind="span"),
            Replace(paragraph_id="p1", start=0, end=3, text="Look", expect="See"),
        ),
        output=tmp_path / "out.docx",
    )
    assert result.applied == 1
    assert result.refused == [
        "bind p1[6:9]: range cuts a field result, which is regenerated on refresh",
        "bind p1[26:30]: anchor 'claim-1' already exists",
        "bind p1[0:3]: range overlaps an edit in this transaction",
    ]
    out = adapter.open(result.output)
    assert {a.id for a in out.anchors} == {"claim-1"}
    assert out.paragraphs[1].text.startswith("Look Table 2")
    assert _in_field_result(_root(result.output)) == []


def test_bind_over_a_whole_field_result_wraps_the_field(
    adapter: DocxAdapter, tmp_path: Path
):
    view = adapter.open(fields_and_controls(tmp_path))
    result = adapter.apply(
        view,
        _txn(Bind(paragraph_id="p1", start=4, end=11, anchor_id="ref-1", kind="span")),
        output=tmp_path / "out.docx",
    )
    assert result.applied == 1 and result.refused == []
    out = adapter.open(result.output)
    ref = next(a for a in out.anchors if a.id == "ref-1")
    assert out.paragraphs[1].text[ref.start : ref.end] == "Table 2"
    assert _in_field_result(_root(result.output)) == []
    crossrefs = [
        f for f in out.fields if f.kind == "crossref" and f.paragraph_id == "p1"
    ]
    assert [f.display for f in crossrefs][:1] == ["Table 2"]


# --------------------------------------------------------- verification and I/O


def test_untouched_paragraphs_are_byte_identical_after_apply(
    adapter: DocxAdapter, tmp_path: Path
):
    source = table_and_list(tmp_path)
    view = adapter.open(source)
    start, end = _span(view, "p0", "Intake")
    result = adapter.apply(
        view,
        _txn(
            Replace(
                paragraph_id="p0", start=start, end=end, text="Weekly", expect="Intake"
            )
        ),
        output=tmp_path / "out.docx",
    )
    assert result.applied == 1 and result.untouched.ok
    before = _body_paragraphs(source)
    after = _body_paragraphs(result.output)
    assert len(before) == len(after) == 8
    changed = [
        i
        for i, (x, y) in enumerate(zip(before, after))
        if etree.tostring(x) != etree.tostring(y)
    ]
    assert changed == [0]
    assert _root(result.output).findall(f".//{w('tbl')}") != []
    assert _root(source).findall(f".//{w('tbl')}") != []


def test_untouched_package_parts_are_copied_byte_for_byte(
    adapter: DocxAdapter, tmp_path: Path
):
    source = hyperlink_and_footnote(tmp_path)
    view = adapter.open(source)
    start, end = _span(view, "p1", "Second")
    result = adapter.apply(
        view,
        _txn(
            Replace(
                paragraph_id="p1", start=start, end=end, text="Third", expect="Second"
            )
        ),
        output=tmp_path / "out.docx",
    )
    assert result.applied == 1
    before, after = _parts(source), _parts(result.output)
    assert sorted(before) == sorted(after)
    assert [n for n in before if before[n] != after[n]] == ["word/document.xml"]


def test_comment_rewrites_only_the_comment_related_parts(
    adapter: DocxAdapter, tmp_path: Path
):
    source = hyperlink_and_footnote(tmp_path)
    view = adapter.open(source)
    start, end = _span(view, "p1", "Second")
    result = adapter.apply(
        view,
        _txn(CommentOp(paragraph_id="p1", start=start, end=end, text="Expand?")),
        output=tmp_path / "out.docx",
    )
    assert result.applied == 1
    before, after = _parts(source), _parts(result.output)
    changed = sorted(
        n for n in set(before) | set(after) if before.get(n) != after.get(n)
    )
    assert "word/footnotes.xml" not in changed
    assert "word/styles.xml" not in changed
    assert "word/document.xml" in changed and "word/comments.xml" in changed


def test_output_is_removed_when_verification_fails(
    adapter: DocxAdapter, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
):
    """Surgery claims the edit applied but changed nothing: verification must fail."""
    real = writer_module.apply_edits

    def lying(textmap, edits, **kwargs):
        report = real(textmap, [], **kwargs)
        report.applied_edits = list(edits)
        return report

    monkeypatch.setattr(writer_module, "apply_edits", lying)
    view = adapter.open(plain_two_paragraphs(tmp_path))
    start, end = _span(view, "p0", "Jane")
    output = tmp_path / "out.docx"
    with pytest.raises(DocxApplyError, match="do not read as intended"):
        adapter.apply(
            view,
            _txn(
                Replace(
                    paragraph_id="p0", start=start, end=end, text="Joan", expect="Jane"
                )
            ),
            output=output,
        )
    assert not output.exists()


def test_output_is_removed_when_an_untargeted_paragraph_changes(
    adapter: DocxAdapter, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
):
    real = writer_module._Writer.apply_edits

    def also_change_p1(self, origin):
        real(self, origin)
        element = self.elements[1]  # p1 is not targeted by the transaction
        next(element.iter(w("t"))).text += "!"

    monkeypatch.setattr(writer_module._Writer, "apply_edits", also_change_p1)
    view = adapter.open(plain_two_paragraphs(tmp_path))
    start, end = _span(view, "p0", "Jane")
    output = tmp_path / "out.docx"
    with pytest.raises(
        DocxApplyError, match="paragraphs outside the operations changed"
    ):
        adapter.apply(
            view,
            _txn(
                Replace(
                    paragraph_id="p0", start=start, end=end, text="Joan", expect="Jane"
                )
            ),
            output=output,
        )
    assert not output.exists()


def test_in_place_output_and_stale_views_are_refused_before_any_write(
    adapter: DocxAdapter, tmp_path: Path
):
    source = plain_two_paragraphs(tmp_path)
    before = source.read_bytes()
    view = adapter.open(source)
    with pytest.raises(ValueError, match="never edits in place"):
        adapter.apply(view, _txn(), output=source)
    assert source.read_bytes() == before

    changed = tmp_path / "changed.docx"
    start, end = _span(view, "p0", "Jane")
    adapter.apply(
        view,
        _txn(
            Replace(paragraph_id="p0", start=start, end=end, text="Joan", expect="Jane")
        ),
        output=changed,
    )
    stale = view.model_copy(update={"path": changed})
    with pytest.raises(ValueError, match="stale"):
        adapter.apply(stale, _txn(), output=tmp_path / "x.docx")
    assert not (tmp_path / "x.docx").exists()


def test_output_opens_in_python_docx_and_keeps_the_existing_structures(
    adapter: DocxAdapter, tmp_path: Path
):
    view = adapter.open(table_and_list(tmp_path))
    result = adapter.apply(
        view,
        _txn(
            CommentOp(paragraph_id="p5", start=0, end=7, text="Which?"),
            InsertParagraph(after_paragraph_id="p7", text="Added item."),
        ),
        output=tmp_path / "out.docx",
    )
    assert result.applied == 2
    document = docx.Document(str(result.output))
    assert len(document.tables) == 1
    # python-docx reads text in direct runs only, so a tracked insertion reads as
    # an empty paragraph; the reader sees the inserted text.
    assert [p.text for p in document.paragraphs][-1] == ""
    assert adapter.open(result.output).paragraphs[-1].text == "Added item."


# ----------------------------------------------------------------- citations


def test_citation_in_a_plain_document_is_text_and_adds_a_bibliography_entry(
    adapter: DocxAdapter, tmp_path: Path
):
    view = adapter.open(headed_sections(tmp_path))
    at = len(_text(view, "p2")) - 1  # before the full stop
    result = adapter.apply(
        view,
        _txn(
            InsertCitation(paragraph_id="p2", at=at, items=[CitationItem(csl=SMITH)]),
            UpdateBibliography(),
        ),
        output=tmp_path / "out.docx",
    )
    assert result.applied == 2 and result.refused == []
    out = adapter.open(result.output)
    assert _text(out, "p2") == (
        "Sodium intake exceeds guidelines in the invented cohort(Smith and Jones 2021)."
    )
    assert not [f for f in out.fields if f.manager == "zotero"]
    texts = [p.text for p in out.paragraphs]
    bibliography = texts.index("Bibliography")
    assert texts[bibliography + 1].startswith("Doe, J. (2020).")
    assert texts[-1].startswith("Smith, J. Q., & Jones, K. (2021).")
    assert texts[-1].endswith("https://doi.org/10.5555/saltlog.smith")


def test_a_work_already_in_the_bibliography_is_not_added_again(
    adapter: DocxAdapter, tmp_path: Path
):
    view = adapter.open(headed_sections(tmp_path))
    at = len(_text(view, "p2")) - 1
    result = adapter.apply(
        view,
        _txn(
            InsertCitation(paragraph_id="p2", at=at, items=[CitationItem(csl=DOE)]),
            UpdateBibliography(),
        ),
        output=tmp_path / "out.docx",
    )
    assert result.applied == 2 and result.refused == []
    assert len(adapter.open(result.output).paragraphs) == len(view.paragraphs)
    assert _text(adapter.open(result.output), "p2").endswith("(Doe 2020).")


def test_citation_in_a_zotero_document_is_a_zotero_field_with_the_display_text(
    adapter: DocxAdapter, tmp_path: Path
):
    source = build_docx(
        tmp_path / "zotero.docx",
        paragraph(
            run("Sodium rose ")
            + complex_field("ADDIN ZOTERO_ITEM CSL_CITATION {}", "(Doe 2020)")
            + run(" overall.")
        ),
    )
    view = adapter.open(source)
    at = len(_text(view, "p0")) - 1
    result = adapter.apply(
        view,
        _txn(InsertCitation(paragraph_id="p0", at=at, items=[CitationItem(csl=SMITH)])),
        output=tmp_path / "out.docx",
    )
    assert result.applied == 1 and result.refused == []
    out = adapter.open(result.output)
    zotero = [f for f in out.fields if f.kind == "citation" and f.manager == "zotero"]
    assert len(zotero) == 2
    assert zotero[1].display == "(Smith and Jones 2021)"
    assert _in_field_result(_root(result.output)) == []
