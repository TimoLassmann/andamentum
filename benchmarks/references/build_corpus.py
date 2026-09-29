"""Build a test set of Word documents, one per real citation style.

Real styles come from the Citation Style Language repository; real reference
metadata from Crossref (CSL-JSON by DOI). ``citeproc-py`` renders each work's
in-text citation and bibliography entry in each style, and every document
carries the same planted errors and the same distractors, with the answers
written beside it:

- a work cited in the text but missing from the list;
- a work in the list that is never cited;
- in author-year styles, one citation whose year is off by one;
- distractors a reader must not mistake for citations: affiliation marks in
  the author block, "10⁵ cells", "3 m²", "Fig. 2", "(n = 12)" and a numbered
  step list "(1) ... (2) ...".

Each superscript style is also written a second time as ``NAME.flat.docx``
with every raised character set as plain text — what a converted PDF gives:
"patients1." with nothing to say the 1 was raised. That variant is where the
model earns its place.

Run (citeproc-py is needed only here, not by the package)::

    uv run --with citeproc-py python benchmarks/references/build_corpus.py OUT_DIR

Downloads are cached in OUT_DIR/cache, so a rerun is offline.
"""

from __future__ import annotations

import json
import re
import sys
import time
import urllib.request
import warnings
from pathlib import Path

warnings.filterwarnings("ignore")

STYLES = [
    "apa",
    "harvard-cite-them-right",
    "chicago-author-date",
    "elsevier-harvard",
    "springer-basic-author-date",
    "nature",
    "cell",
    "ieee",
    "american-medical-association",
    "nucleic-acids-research",
    "plos",
    "bmj",
    "modern-language-association",
]
DOIS = [
    "10.1093/nar/gky1016",
    "10.1002/humu.24309",
    "10.1038/s41586-020-2434-2",
    "10.1038/nature14539",
    "10.1038/s41586-021-03819-2",
    "10.1093/bioinformatics/btp324",
    "10.1038/nmeth.1923",
    "10.1186/s13059-014-0550-8",
    "10.1038/nbt.1754",
    "10.1101/gr.107524.110",
]
MISSING = 4  # cited, but left out of the list (0-based index into DOIS)
UNCITED = 9  # in the list, never cited
WRONG_YEAR = 2  # author-year styles: this citation's year is off by one
SENTENCES = [
    "Rare disease diagnosis still fails for most patients",
    "Population databases shape how variants are read",
    "Genome sequencing in a national health system raised the yield",
    "Deep learning has changed many fields",
    "Protein structure can now be predicted",
    "Short reads are aligned with fast index methods",
    "Gapped alignment scales to large genomes",
    "Count data need their own statistical models",
    "Visual inspection of alignments remains useful",
]


def fetch(url: str, path: Path, *, accept: str | None = None) -> bytes:
    if path.exists():
        return path.read_bytes()
    request = urllib.request.Request(url, headers={"Accept": accept} if accept else {})
    with urllib.request.urlopen(request, timeout=60) as response:
        data = response.read()
    path.write_bytes(data)
    time.sleep(0.3)
    return data


def load_items(cache: Path) -> list[dict]:
    items = []
    for index, doi in enumerate(DOIS):
        raw = fetch(
            f"https://api.crossref.org/works/{doi}/transform",
            cache / f"item-{index}.json",
            accept="application/vnd.citationstyles.csl+json",
        )
        item = json.loads(raw)
        for key in ("reference", "assertion", "license", "link", "relation"):
            item.pop(key, None)
        item["id"] = f"r{index}"
        items.append(item)
    return items


def flat(rendered) -> str:
    if isinstance(rendered, str):
        return rendered
    return "".join(flat(part) for part in rendered)


def render(
    style_path: Path, items: list[dict]
) -> tuple[list[str], dict[str, str], str]:
    from citeproc import (
        Citation,
        CitationItem,
        CitationStylesBibliography,
        CitationStylesStyle,
        formatter,
    )
    from citeproc.source.json import CiteProcJSON

    style = CitationStylesStyle(str(style_path), validate=False)
    bibliography = CitationStylesBibliography(
        style, CiteProcJSON(items), formatter.plain
    )
    cited = [i for i in range(len(items)) if i != UNCITED]
    citations = [Citation([CitationItem(f"r{i}")]) for i in cited]
    for citation in citations:
        bibliography.register(citation)
    # Register the uncited work too, so it is in the list.
    bibliography.register(Citation([CitationItem(f"r{UNCITED}")]))
    shown = {
        f"r{i}": flat(bibliography.cite(c, lambda _: None)).strip()
        for i, c in zip(cited, citations)
    }
    entries = [flat(entry).strip() for entry in bibliography.bibliography()]
    ids = [str(key) for key in bibliography.keys]
    fmt = "numeric"
    text = style_path.read_text()
    if 'citation-format="author-date"' in text:
        fmt = "author_year"
    elif 'citation-format="author"' in text:
        fmt = "author_only"
    superscript = (
        'vertical-align="sup"'
        in text.split("<citation", 1)[-1].split("</citation>", 1)[0]
    )
    return (
        ids,
        shown | {"__entries__": json.dumps(entries)},
        fmt + ("_superscript" if superscript else ""),
    )


FOOTNOTE_CONTENT_TYPE = (
    "application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml"
)
FOOTNOTE_REL = (
    "http://schemas.openxmlformats.org/officeDocument/2006/relationships/footnotes"
)
W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def _short_title(title: str) -> str:
    words = re.split(r"\s+", re.split(r"[:?]", title)[0].strip())
    return " ".join(words[:4])


def _add_footnotes(path: Path, notes: list[str]) -> None:
    """Turn the ⟦FNn⟧ placeholder runs into real Word footnotes."""
    import zipfile
    from io import BytesIO
    from xml.sax.saxutils import escape

    data = path.read_bytes()
    out = BytesIO()
    with (
        zipfile.ZipFile(BytesIO(data)) as src,
        zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst,
    ):
        for item in src.infolist():
            body = src.read(item.filename)
            if item.filename == "word/document.xml":
                text = body.decode()
                for n in range(1, len(notes) + 1):
                    text = re.sub(
                        rf"<w:r>(?:<w:rPr>.*?</w:rPr>)?<w:t[^>]*>⟦FN{n}⟧</w:t></w:r>",
                        f'<w:r><w:rPr><w:vertAlign w:val="superscript"/></w:rPr><w:footnoteReference w:id="{n}"/></w:r>',
                        text,
                    )
                body = text.encode()
            elif item.filename == "word/_rels/document.xml.rels":
                body = body.replace(
                    b"</Relationships>",
                    f'<Relationship Id="rIdFootnotes" Type="{FOOTNOTE_REL}" Target="footnotes.xml"/></Relationships>'.encode(),
                )
            elif item.filename == "[Content_Types].xml":
                body = body.replace(
                    b"</Types>",
                    f'<Override PartName="/word/footnotes.xml" ContentType="{FOOTNOTE_CONTENT_TYPE}"/></Types>'.encode(),
                )
            dst.writestr(item, body)
        footnotes = [
            '<w:footnote w:type="separator" w:id="-1"><w:p><w:r><w:separator/></w:r></w:p></w:footnote>',
            '<w:footnote w:type="continuationSeparator" w:id="0"><w:p><w:r><w:continuationSeparator/></w:r></w:p></w:footnote>',
        ] + [
            f'<w:footnote w:id="{n}"><w:p><w:r><w:t xml:space="preserve">{escape(note)}</w:t></w:r></w:p></w:footnote>'
            for n, note in enumerate(notes, start=1)
        ]
        dst.writestr(
            "word/footnotes.xml",
            f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?><w:footnotes xmlns:w="{W}">{"".join(footnotes)}</w:footnotes>',
        )
    path.write_bytes(out.getvalue())


def _build_notes(
    bibliography: bool, *, cache: Path, items: list[dict], out: Path
) -> dict | None:
    """A Chicago notes-and-bibliography document with real Word footnotes.

    Notes, in order: full r0, full r1, short r0, Ibid. (fine), full r2 + r3
    in one note, Ibid. (planted: unclear — two sources before it), a note of
    commentary (distractor), full r4 (planted with a bibliography: r4 is not
    in it), short r5 BEFORE full r5 (planted without a bibliography), full
    r5, full r6, r7, r8. With a bibliography it also lists r9, never cited
    (planted).
    """
    from citeproc import (
        Citation,
        CitationItem,
        CitationStylesBibliography,
        CitationStylesStyle,
        formatter,
    )
    from citeproc.source.json import CiteProcJSON
    from docx import Document

    name = "chicago-notes-bibliography"
    style_path = cache / f"{name}.csl"
    fetch(
        f"https://raw.githubusercontent.com/citation-style-language/styles/master/{name}.csl",
        style_path,
    )
    style = CitationStylesStyle(str(style_path), validate=False)
    bib = CitationStylesBibliography(style, CiteProcJSON(items), formatter.plain)

    def full(i: int) -> str:
        citation = Citation([CitationItem(f"r{i}")])
        bib.register(citation)
        text = flat(bib.cite(citation, lambda _: None)).strip()
        return (
            re.sub(r";\s*originally published as\s*\[\],?", ",", text).rstrip(".") + "."
        )

    def family(i: int) -> str:
        return items[i]["author"][0]["family"]

    def short(i: int, page: int) -> str:
        etal = " et al." if len(items[i]["author"]) > 3 else ""
        return f"{family(i)}{etal}, “{_short_title(items[i]['title'] if isinstance(items[i]['title'], str) else items[i]['title'][0])},” {page}."

    notes: list[tuple[str, str | None]] = [  # (text, planted kind or None)
        (full(0), None),
        (full(1), None),
        (short(0, 12), None),
        ("Ibid., 14.", None),
        (full(2).rstrip(".") + "; " + full(3), None),
        ("Ibid., 3.", "ibid_unclear"),
        ("This point is discussed further in the next section.", "comment"),
        (full(4), "cited_not_listed" if bibliography else None),
        (short(5, 22), None if bibliography else "short_before_full"),
        (full(5), None),
        (full(6), None),
        (full(7), None),
        (full(8), None),
        # Notes the rules cannot read: an interview is not a published source,
        # a report by an organisation with a dotted name is.
        ("Interview with the author, March 2019.", "comment"),
        ("U.S. Census Bureau, “American Community Survey,” 2019.", None),
    ]
    for i in range(10):
        bib.register(Citation([CitationItem(f"r{i}")]))
    entries = {str(k): flat(e).strip() for k, e in zip(bib.keys, bib.bibliography())}

    doc = Document()
    paragraphs: list[str] = []
    truth: list[tuple[int, int]] = []
    planted: list[dict] = []

    def offset() -> int:
        return sum(len(p) + 1 for p in paragraphs)

    doc.add_heading("A test of citation reading", level=1)
    paragraphs.append("A test of citation reading")
    doc.add_heading("Introduction", level=1)
    paragraphs.append("Introduction")
    sentences = SENTENCES + [
        "Further work followed",
        "The field keeps moving",
        "Methods keep improving",
        "Tools keep changing",
        "Clinicians agreed",
        "Population data were used",
    ]
    for n, (note, kind) in enumerate(notes, start=1):
        sentence = sentences[n - 1] + "."
        p = doc.add_paragraph()
        p.add_run(sentence)
        p.add_run(f"⟦FN{n}⟧")
        word_start = offset() + sentence.rfind(" ") + 1
        span = (word_start, offset() + len(sentence))
        paragraphs.append(sentence)
        if kind != "comment":
            truth.append(span)
        if kind and kind != "comment":
            planted.append({"kind": kind, "start": span[0], "end": span[1]})
    if bibliography:
        doc.add_heading("Bibliography", level=1)
        paragraphs.append("Bibliography")
        for key in sorted(entries, key=lambda k: entries[k]):
            if key == "r4":
                continue
            if key == "r9":
                planted.append(
                    {"kind": "listed_not_cited", "start": offset(), "end": offset() + 8}
                )
            doc.add_paragraph(entries[key])
            paragraphs.append(entries[key])
    variant = "chicago-notes" + ("" if bibliography else ".nobib")
    path = out / f"{variant}.docx"
    doc.save(str(path))
    _add_footnotes(path, [n for n, _ in notes])
    print(f"built {variant}: notes  e.g. {notes[0][0][:70]!r}")
    return {
        "style": variant,
        "file": path.name,
        "format": "note",
        "superscript": False,
        "expected_style": "notes",
        "text": "\n".join(paragraphs),
        "citations": truth,
        "planted": planted,
        "sample_citation": notes[0][0],
        "sample_entry": "",
    }


def build(out: Path) -> None:
    from docx import Document

    cache = out / "cache"
    cache.mkdir(parents=True, exist_ok=True)
    items = load_items(cache)
    manifest = []
    for name in STYLES:
        for flat_text in (False, True):
            case = _build_one(
                name,
                flat_text,
                cache=cache,
                items=items,
                out=out,
                document_cls=Document,
            )
            if case is not None:
                manifest.append(case)
    for with_bibliography in (True, False):
        case = _build_notes(with_bibliography, cache=cache, items=items, out=out)
        if case is not None:
            manifest.append(case)
    (out / "manifest.json").write_text(json.dumps(manifest, indent=1))


def _build_one(name, flat_text, *, cache, items, out, document_cls):
    Document = document_cls
    if True:
        style_path = cache / f"{name}.csl"
        try:
            fetch(
                f"https://raw.githubusercontent.com/citation-style-language/styles/master/{name}.csl",
                style_path,
            )
            ids, shown, kind = render(style_path, items)
        except Exception as exc:  # a style citeproc-py cannot render
            print(f"skip {name}: {type(exc).__name__}: {exc}")
            return None
        entries = json.loads(shown.pop("__entries__"))
        superscript = kind.endswith("_superscript")
        if flat_text and not superscript:
            return None
        fmt = kind.removesuffix("_superscript")
        numbered_list = fmt == "numeric"

        doc = Document()
        paragraphs: list[str] = []  # the projection, paragraph by paragraph
        truth_citations: list[tuple[int, int]] = []
        planted: list[dict] = []  # the problems a correct reader reports

        def para(parts: list[tuple[str, bool]], *, heading: bool = False) -> None:
            if heading:
                p = doc.add_heading("", level=1)
            else:
                p = doc.add_paragraph()
            for text, raised in parts:
                run = p.add_run(text)
                run.font.superscript = raised and not flat_text
            paragraphs.append("".join(t for t, _ in parts))

        def offset() -> int:
            return sum(len(p) + 1 for p in paragraphs)

        para([("A test of citation reading", False)], heading=True)
        para(
            [
                ("Ada Lovelace", False),
                ("1", True),
                (", Alan Turing", False),
                ("2", True),
            ]
        )
        para([("Introduction", False)], heading=True)
        for position, index in enumerate(i for i in range(len(items)) if i != UNCITED):
            key = f"r{index}"
            cite = shown[key]
            wrong_year = fmt == "author_year" and index == WRONG_YEAR
            if wrong_year:
                year = str(items[index]["issued"]["date-parts"][0][0])
                cite = cite.replace(year, str(int(year) - 1))
            sentence = SENTENCES[position]
            if superscript:
                start = offset() + len(sentence)
                para([(sentence, False), (cite, True), (".", False)])
            else:
                start = offset() + len(sentence) + 1
                para([(f"{sentence} {cite}.", False)])
            truth_citations.append((start, start + len(cite)))
            if index == MISSING:
                planted.append(
                    {
                        "kind": "cited_not_listed",
                        "start": start,
                        "end": start + len(cite),
                    }
                )
            if wrong_year:
                planted.append(
                    {"kind": "year_mismatch", "start": start, "end": start + len(cite)}
                )
        para(
            [
                ("We grew 10", False),
                ("5", True),
                (" cells on 3 m", False),
                ("2", True),
                (" plates (n = 12), as in Fig. 2.", False),
            ]
        )
        # Harder: gene, exon and day labels that look like glued citations in
        # a text that lost its formatting, and a repeated citation, which the
        # citation-order rule cannot settle.
        para([("Reads on chr11 and exon2 were counted at day3.", False)])
        if superscript:
            sentence = "This agrees with earlier work"
            start = offset() + len(sentence)
            para([(sentence, False), ("2", True), (".", False)])
            truth_citations.append((start, start + 1))
        para([("(1) prepare the samples.", False)])
        para([("(2) sequence them.", False)])
        para([("References", False)], heading=True)
        list_entries = []
        for key, entry in zip(ids, entries):
            if key == f"r{MISSING}":
                continue
            list_entries.append((key, entry))
        for number, (key, entry) in enumerate(list_entries, start=1):
            if numbered_list and not re.match(r"\[?\d{1,3}[\].)]?", entry):
                entry = f"{number}. {entry}"
            if key == f"r{UNCITED}":
                planted.append(
                    {"kind": "listed_not_cited", "start": offset(), "end": offset() + 8}
                )
            para([(entry, False)])

        path = out / f"{name}{'.flat' if flat_text else ''}.docx"
        doc.save(str(path))
        case = {
            "style": name + (".flat" if flat_text else ""),
            "file": path.name,
            "format": fmt,
            "superscript": superscript,
            "expected_style": (
                "author_year"
                if fmt == "author_year"
                else "numeric_superscript"
                if superscript
                else "numeric_brackets"
                if "[" in shown["r0"]
                else "numeric_parentheses"
                if "(" in shown["r0"]
                else "unsupported"
            )
            if fmt != "author_only"
            else "unsupported",
            "text": "\n".join(paragraphs),
            "citations": truth_citations,
            "planted": planted,
            "sample_citation": shown["r0"],
            "sample_entry": list_entries[0][1][:120],
        }
        print(f"built {case['style']}: {case['expected_style']}  e.g. {shown['r0']!r}")
        return case


if __name__ == "__main__":
    build(Path(sys.argv[1]))
